"""第四问：真实运动约束下的双圆弧优化、result4 与结果/检验图。

python 问题4代码.py --quick
python 问题4代码.py --clearance-margin 0.01 --improvement-rel 0.001
默认输出基目录为 ../结果/运行记录/第四问，每次运行创建独立时间戳目录。
依赖：numpy scipy matplotlib openpyxl，以及同目录 问题1代码.py、问题2代码.py。
分段弧长路径、刚性链和全队调头检查已合并在本脚本中，供第五问复用。
圆弧长度由几何公式计算；Matérn 高斯过程只学习完整运动检查的采样最小间隙。
本脚本保存“数值采样可行候选”，不宣称全局最优或连续时间不漏检认证。
默认保留 0.01 m 的 SAT 分离裕量，连续 8 次未获得超过 0.1% 的弧长改进即停止优化。
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import time
from dataclasses import asdict
from pathlib import Path

import numpy as np
from scipy.linalg import cho_factor, cho_solve
from scipy.optimize import brentq, minimize
from scipy.special import ndtr
from scipy.stats import qmc

from 问题1代码 import (
    Parameters, board_vertices, configure_plotting, draw_chain, handle_label,
    prepare_output_dir, save_figure, save_json, scientific_tex,
)
from 问题2代码 import sat_clearance

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
SELECTED = np.array([0, 1, 51, 101, 151, 201, 223])
EPS_X = 1e-8
SCOPE = "完整运动区间的数值采样约束；未证明连续无碰撞或全局最短。"
MOTION_CHECK_VERSION = 2


# 第四、五问共用的分段路径与刚性链运动
def _F(theta):
    return .5 * (theta * math.hypot(1., theta) + math.asinh(theta))


def _J(point):
    return np.array([-point[1], point[0]])


def _spiral(theta, b):
    c, s = math.cos(theta), math.sin(theta)
    return np.array([b * theta * c, b * theta * s]), np.array([b * (c-theta*s), b * (s+theta*c)])


def _angle(point):
    return math.atan2(point[1], point[0])


def _deadline(deadline):
    if deadline is not None and time.perf_counter() >= deadline:
        raise TimeoutError("调头运动计算预算已用尽")


class TurnPath:
    """s=0 在第一段圆弧起点；s 沿行进方向递增，单位 m。"""

    def __init__(self, x, params, radius, rho):
        self.x = np.asarray(x, dtype=float)
        self.params, self.radius = params, float(radius)
        self.theta_a, self.theta_c = self.x * self.radius / params.b
        self.rho = float(rho)
        self.A, da = _spiral(self.theta_a, params.b)
        pb, db = _spiral(self.theta_c, params.b)
        self.B = -pb
        self.TA, self.TB = -da / np.linalg.norm(da), -db / np.linalg.norm(db)
        self.O1 = self.A - 2 * rho * _J(self.TA)
        self.O2 = self.B + rho * _J(self.TB)
        self.C = (self.O1 + 2 * self.O2) / 3
        self.phi1, self.phi2 = _angle(self.A-self.O1), _angle(self.C-self.O2)
        self.alpha1 = (self.phi1 - _angle(self.C-self.O1)) % (2*math.pi)
        self.alpha2 = (_angle(self.B-self.O2) - self.phi2) % (2*math.pi)
        self.ell1, self.ell2 = 2*rho*self.alpha1, rho*self.alpha2
        self.length = self.ell1 + self.ell2
        self._events = None
        self._domain = None

    def _theta(self, s, outgoing=False):
        start = self.theta_c if outgoing else self.theta_a
        distance = s-self.length if outgoing else -s
        target = _F(start) + distance/self.params.b
        # 弧长原函数单调，Newton 仅在保证正值时更新，最终核对弧长残差。
        theta = math.sqrt(max(start*start + 2*distance/self.params.b, 0.))
        for _ in range(12):
            delta = (_F(theta)-target) / math.hypot(1., theta)
            new = theta-delta
            if new < 0:
                break
            theta = new
            if abs(delta) < 2e-13 * max(1., theta):
                return theta
        upper = max(start + distance/self.params.b, start+1.)
        return brentq(lambda a: _F(a)-target, start, upper, xtol=self.params.root_tol)

    def _s_theta(self, theta, outgoing=False):
        if outgoing:
            return self.length + self.params.b * (_F(theta)-_F(self.theta_c))
        return self.params.b * (_F(self.theta_a)-_F(theta))

    def point_tangent(self, s):
        s = float(s)
        if not math.isfinite(s):
            raise ValueError("路径弧长必须有限")
        if s < 0:
            point, derivative = _spiral(self._theta(s), self.params.b)
            return point, -derivative/np.linalg.norm(derivative)
        if s <= self.ell1:
            phi = self.phi1-s/(2*self.rho)
            radial = np.array([math.cos(phi), math.sin(phi)])
            return self.O1+2*self.rho*radial, -_J(radial)
        if s <= self.length:
            phi = self.phi2+(s-self.ell1)/self.rho
            radial = np.array([math.cos(phi), math.sin(phi)])
            return self.O2+self.rho*radial, _J(radial)
        point, derivative = _spiral(self._theta(s, True), self.params.b)
        return -point, -derivative/np.linalg.norm(derivative)

    @staticmethod
    def _arc_max_radius(center, radius, phi, angle, direction):
        phases = [phi, phi+direction*angle]
        for stationary in (_angle(center), _angle(center)+math.pi):
            travel = ((stationary-phi)*direction) % (2*math.pi)
            if travel <= angle+1e-12:
                phases.append(phi+direction*travel)
        return max(math.hypot(*(center+radius*np.array([math.cos(a), math.sin(a)]))) for a in phases)

    def geometry_report(self):
        p0, t0 = self.point_tangent(0.)
        p1, t1 = self.point_tangent(self.ell1)
        p2, t2 = self.point_tangent(self.length)
        second_tangent = _J((self.C-self.O2)/self.rho)
        position_error = max(np.linalg.norm(p0-self.A), np.linalg.norm(p1-self.C), np.linalg.norm(p2-self.B))
        tangent_error = max(np.linalg.norm(t0-self.TA), np.linalg.norm(t1-second_tangent), np.linalg.norm(t2-self.TB))
        circle_error = abs(np.linalg.norm(self.O2-self.O1)-3*self.rho)
        radii = [self._arc_max_radius(self.O1, 2*self.rho, self.phi1, self.alpha1, -1),
                 self._arc_max_radius(self.O2, self.rho, self.phi2, self.alpha2, 1)]
        passed = (min(self.alpha1, self.alpha2, self.rho) > 1e-9 and
                  max(position_error, tangent_error, circle_error) < 1e-8 and
                  max(radii) <= self.radius+1e-9)
        return {"passed": bool(passed), "position_residual_m": float(position_error),
                "tangent_residual": float(tangent_error), "external_tangency_residual_m": float(circle_error),
                "radius_ratio": 2., "arc_max_radii_m": radii, "disk_excess_m": max(0., max(radii)-self.radius)}

    def to_dict(self):
        return {"schema_version": 1, "x": self.x.tolist(), "parameters": asdict(self.params),
                "radius_m": self.radius, "theta_a": float(self.theta_a), "theta_c": float(self.theta_c),
                "rho": self.rho, "A": self.A.tolist(), "B": self.B.tolist(), "C": self.C.tolist(),
                "O1": self.O1.tolist(), "O2": self.O2.tolist(), "alpha1": self.alpha1, "alpha2": self.alpha2,
                "ell1": self.ell1, "ell2": self.ell2, "length": self.length,
                "s_zero": "第一段圆弧起点A", "geometry": self.geometry_report()}

    @classmethod
    def from_dict(cls, data):
        if data.get("schema_version") != 1:
            raise ValueError("未知的路径参数版本")
        path = build_path(data["x"], Parameters(**data["parameters"]), data["radius_m"])
        for key in ("theta_a", "theta_c", "rho", "alpha1", "alpha2", "ell1", "ell2", "length"):
            if key in data and not math.isclose(float(data[key]), float(getattr(path, key)), rel_tol=1e-10, abs_tol=1e-9):
                raise ValueError(f"保存路径的 {key} 与重建结果不一致")
        for key in ("A", "B", "C", "O1", "O2"):
            if key in data and not np.allclose(data[key], getattr(path, key), rtol=0., atol=1e-9):
                raise ValueError(f"保存路径的 {key} 与重建结果不一致")
        return path


def build_path(x=(1., 1.), params=None, radius=4.5):
    params = params or Parameters(pitch=1.7)
    x = np.asarray(x, dtype=float)
    if x.shape != (2,) or not np.isfinite(x).all() or np.any(x <= 0) or np.any(x > 1) or not math.isfinite(radius) or radius <= 0:
        raise ValueError("接入角比例须在 (0,1]，调头区域半径须为有限正数")
    a, da = _spiral(x[0]*radius/params.b, params.b)
    bpoint, db = _spiral(x[1]*radius/params.b, params.b)
    d = -bpoint-a
    q = -2*_J(da/np.linalg.norm(da))-_J(db/np.linalg.norm(db))
    na, nb = -_J(da/np.linalg.norm(da)), -_J(db/np.linalg.norm(db))
    aa, bb, cc = float(-2*np.sum((na-nb)**2)), float(2*d@q), float(d@d)
    if abs(aa) < 1e-11:
        roots = [-cc/bb] if abs(bb) > 1e-14 else []
    else:
        discriminant = bb*bb-4*aa*cc
        if discriminant < 0:
            raise ValueError("无实数外切圆半径")
        stable = -.5*(bb+math.copysign(math.sqrt(discriminant), bb))
        roots = [stable/aa, cc/stable] if stable else []
    paths = [TurnPath(x, params, radius, r) for r in roots if math.isfinite(r) and r > 1e-9]
    paths = [p for p in paths if p.geometry_report()["passed"]]
    if not paths:
        raise ValueError("接入角不满足双圆弧方向、相切或完整圆域约束")
    return min(paths, key=lambda p: p.length)


def _spiral_root(path, anchor, distance, lo, hi, increasing):
    """在有界角度区间按指定顺序寻找首根；同号区间不能直接略过。"""
    b = path.params.b
    if hi < lo-1e-12:
        return None
    def f(theta):
        point, _ = _spiral(theta, b)
        return float(np.linalg.norm(point-anchor)-distance)
    stack = [(lo, hi)]
    count = 0
    while stack:
        left, right = stack.pop()
        count += 1
        if count > 20000:
            raise ArithmeticError("螺线首交点的区间排除未收敛")
        middle, half = .5*(left+right), .5*(right-left)
        fl, fr = f(left), f(right)
        first, last = (left, right) if increasing else (right, left)
        ff, lf = (fl, fr) if increasing else (fr, fl)
        if abs(ff) < 2e-11:
            return first
        pm, dm = _spiral(middle, b)
        dist_mid = float(np.linalg.norm(pm-anchor))
        reach = b*math.hypot(1., right)*half
        if dist_mid+reach < distance-2e-12 or dist_mid-reach > distance+2e-12:
            continue
        # 平方距离导数的区间界；单调时括区间内至多一个交点。
        derivative = float(2*(pm-anchor)@dm)
        second_bound = 2*b*b*(1+right*right)+2*(dist_mid+reach)*b*math.sqrt(4+right*right)
        monotone = abs(derivative) > second_bound*half
        if monotone:
            if fl*fr < 0:
                return brentq(f, left, right, xtol=path.params.root_tol)
            if abs(lf) < 2e-11:
                return last
            continue
        if right-left < 1e-11:
            raise ArithmeticError("检测到近切孔距根，无法确认连续前进分支")
        if increasing:
            stack.extend([(middle, right), (left, middle)])
        else:
            stack.extend([(left, middle), (middle, right)])
    return None


def _neighbor(path, s_anchor, distance, direction):
    """向前(direction=1)或向后(-1)求最近固定孔距交点。"""
    anchor, _ = path.point_tangent(s_anchor)
    b, length = path.params.b, path.length
    region = 0 if s_anchor < 0 else 1 if s_anchor <= path.ell1 else 2 if s_anchor <= length else 3
    segments = range(region, 4) if direction > 0 else range(region, -1, -1)
    for segment in segments:
        if segment in (0, 3):
            outgoing = segment == 3
            theta_min = path.theta_c if outgoing else path.theta_a
            same = (segment == region and ((outgoing and s_anchor > length) or (not outgoing and s_anchor < 0)))
            theta_direction = direction if outgoing else -direction
            theta_anchor = path._theta(s_anchor, outgoing) if same else None
            # 同段螺线的单调近根，避免对每个远端把手进行通用区间细分。
            if same:
                r = b*theta_anchor
                def delta_f(delta):
                    return b*b*delta*delta+4*r*(r+theta_direction*b*delta)*math.sin(delta/2)**2-distance*distance
                if theta_direction > 0 and r > distance/math.sqrt(2):
                    cap = 2*math.asin(min(1., distance/(2*r)))
                    delta = brentq(delta_f, 0., cap, xtol=path.params.root_tol)
                    return path._s_theta(theta_anchor+delta, outgoing)
                if theta_direction < 0 and theta_min >= 1.:
                    cap = min(math.pi/2, theta_anchor-theta_min)
                    if cap > 0 and delta_f(cap) >= 0:
                        delta = brentq(delta_f, 0., cap, xtol=path.params.root_tol)
                        return path._s_theta(theta_anchor-delta, outgoing)
                    if cap == theta_anchor-theta_min:
                        continue
            adjusted = -anchor if outgoing else anchor
            radius_anchor = float(np.linalg.norm(adjusted))
            lo, hi = max(theta_min, (radius_anchor-distance)/b, 0.), (radius_anchor+distance)/b
            if same:
                if theta_direction > 0:
                    lo = max(lo, theta_anchor)
                else:
                    hi = min(hi, theta_anchor)
            theta = _spiral_root(path, adjusted, distance, lo, hi, theta_direction > 0)
            if theta is not None:
                return path._s_theta(theta, outgoing)
        else:
            first_arc = segment == 1
            center, radius = (path.O1, 2*path.rho) if first_arc else (path.O2, path.rho)
            start = 0. if first_arc else path.ell1
            end = path.ell1 if first_arc else length
            phi = path.phi1 if first_arc else path.phi2
            sense = -1 if first_arc else 1
            lower, upper = (max(start, s_anchor), end) if direction > 0 else (start, min(end, s_anchor))
            delta, norm = anchor-center, float(np.linalg.norm(anchor-center))
            if norm < 1e-12:
                if abs(radius-distance) < 1e-9:
                    raise ArithmeticError("圆弧孔距根不唯一")
                continue
            along = (radius*radius-distance*distance+norm*norm)/(2*norm)
            height2 = radius*radius-along*along
            if height2 < -1e-10:
                continue
            normal = _J(delta/norm)
            base = center+along*delta/norm
            candidates = []
            for sign in (-1., 1.):
                point = base+sign*math.sqrt(max(0., height2))*normal
                angle = ((_angle(point-center)-phi)*sense) % (2*math.pi)
                if abs(angle-2*math.pi) < 1e-10:
                    angle = 0.
                s = start+radius*angle
                if lower-1e-10 <= s <= upper+1e-10 and direction*(s-s_anchor) > 1e-9:
                    candidates.append(min(max(s, lower), upper))
            if candidates:
                return min(candidates) if direction > 0 else max(candidates)
    raise ArithmeticError("指定方向不存在可确认的最近孔距根")


def chain_state(path, s_head, speed=1.):
    if not math.isfinite(speed) or speed <= 0:
        raise ValueError("龙头速度须为有限正数")
    coordinates = np.empty(224)
    positions, tangents = np.empty((224, 2)), np.empty((224, 2))
    coefficients = np.ones(224)
    coordinates[0] = s_head
    positions[0], tangents[0] = path.point_tangent(s_head)
    minimum = 1.
    distances = path.params.distances
    for i, distance in enumerate(distances, 1):
        coordinates[i] = _neighbor(path, coordinates[i-1], float(distance), -1)
        if coordinates[i] >= coordinates[i-1]:
            raise ArithmeticError("把手弧长顺序不正确")
        positions[i], tangents[i] = path.point_tangent(coordinates[i])
        chord = positions[i]-positions[i-1]
        numerator, denominator = float(chord@tangents[i-1]), float(chord@tangents[i])
        projection = abs(denominator)/distance
        minimum = min(minimum, projection)
        if projection < 1e-8:
            raise ArithmeticError("速度递推分母接近零")
        coefficients[i] = coefficients[i-1]*numerator/denominator
        if coefficients[i] <= 0 or not math.isfinite(coefficients[i]):
            raise ArithmeticError("连接分支不支持连续向前运动")
    residual = float(np.max(np.abs(np.linalg.norm(np.diff(positions, axis=0), axis=1)-distances)))
    if residual > 1e-8:
        raise ArithmeticError(f"固定孔距残差超限：{residual}")
    velocities = speed*coefficients[:, None]*tangents
    return {"s": coordinates, "positions": positions, "tangents": tangents,
            "coefficients": coefficients, "velocities": velocities, "speeds": speed*coefficients,
            "max_distance_residual": residual, "min_projection_ratio": float(minimum)}


def _forward_inverse(path, s, distance):
    next_s = _neighbor(path, s, distance, 1)
    recovered = _neighbor(path, next_s, distance, -1)
    if abs(recovered-s) > 1e-7:
        raise ArithmeticError("前向孔距交点不是原后继映射的逆；事件分支无法确认")
    return next_s


def crossing_events(path, deadline=None):
    if path._events is not None:
        return path._events
    events = []
    first, body = path.params.distances[:2]
    for junction, target in enumerate((0., path.ell1, path.length)):
        events.append({"handle": 0, "junction": junction, "path_s": target, "head_s": target})
        body_s = target
        for handle in range(1, 224):
            _deadline(deadline)
            head = _forward_inverse(path, body_s, float(first))
            events.append({"handle": handle, "junction": junction, "path_s": target, "head_s": head})
            if handle < 223:
                body_s = _forward_inverse(path, body_s, float(body))
    path._events = sorted(events, key=lambda event: (event["head_s"], event["handle"], event["junction"]))
    return path._events


def motion_domain(path, deadline=None):
    if path._domain is not None:
        return path._domain
    theta_r = path.radius/path.params.b
    lower = path._s_theta(theta_r)
    exit_s = path._s_theta(theta_r, True)
    current = exit_s
    first, body = path.params.distances[:2]
    for _ in range(222):
        _deadline(deadline)
        current = _forward_inverse(path, current, float(body))
    upper = _forward_inverse(path, current, float(first))
    # 反向构造必须仍对应正向最近分支，不能仅用总板长推断尾端已出界。
    check = chain_state(path, upper)
    if abs(check["s"][-1]-exit_s) > 1e-6:
        raise ArithmeticError("龙尾出界事件未通过正向构型回代")
    path._domain = (min(-100., lower), max(100., upper))
    return path._domain


def evaluate_motion(path, step=2., deadline=None, refine=True):
    """完整域网格、全部跨段事件与危险区间加密；失败/超时不当作可行。"""
    if not math.isfinite(step) or step <= 0:
        raise ValueError("扫描步长须为有限正数")
    records = {}
    domain = None
    max_residual, min_projection = 0., 1.
    def sample(s):
        nonlocal max_residual, min_projection
        _deadline(deadline)
        key = round(float(s), 11)
        if key in records:
            return records[key]
        state = chain_state(path, float(s))
        gap = sat_clearance(state["positions"], path.params)
        max_residual = max(max_residual, state["max_distance_residual"])
        min_projection = min(min_projection, state["min_projection_ratio"])
        row = {"head_s": float(s), "gap": gap["gap"], "pair_0based": gap["pair_0based"],
               "max_coefficient": float(np.max(state["coefficients"]))}
        records[key] = row
        return row
    complete, error = False, None
    try:
        _deadline(deadline)
        domain = motion_domain(path, deadline)
        # 粗网格也须确认跨段事件的逆映射，防止跳支候选进入可行集。
        events = crossing_events(path, deadline)
        nodes = list(np.arange(domain[0], domain[1], step)) + [domain[1]]
        if refine:
            for event in events:
                if domain[0] <= event["head_s"] <= domain[1]:
                    nodes.extend([max(domain[0], event["head_s"]-1e-5), event["head_s"],
                                  min(domain[1], event["head_s"]+1e-5)])
        nodes = sorted(set(nodes))
        for node in nodes:
            sample(node)
        if refine:
            for _ in range(2):
                ordered = sorted(records.values(), key=lambda row: row["head_s"])
                threshold = min(row["gap"] for row in ordered) + .02
                extra = [.5*(a["head_s"]+b["head_s"]) for a, b in zip(ordered, ordered[1:])
                         if b["head_s"]-a["head_s"] > step/8 and
                         (min(a["gap"], b["gap"]) < max(.02, threshold) or a["pair_0based"] != b["pair_0based"])]
                for node in extra:
                    sample(node)
        complete = True
    except (TimeoutError, ArithmeticError, ValueError) as exc:
        error = str(exc)
    ordered = sorted(records.values(), key=lambda row: row["head_s"])
    best = min(ordered, key=lambda row: row["gap"]) if ordered else None
    collision = best is not None and best["gap"] < -1e-9
    status = "COLLISION" if complete and collision else "FEASIBLE_SAMPLED" if complete else "UNRESOLVED"
    return {"status": status, "complete": complete, "min_gap": best["gap"] if best else None,
            "pair_0based": best["pair_0based"] if best else None,
            "critical_head_s": best["head_s"] if best else None, "sample_count": len(ordered),
            "domain": list(domain) if domain else None, "max_distance_residual": max_residual,
            "min_projection_ratio": min_projection, "records": ordered, "error": error,
            "continuous_safety_certified": False, "step_m": step, "crossing_refinement": bool(refine)}


# 第四问的路径优化、检验和结果导出
def _hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _remaining(deadline):
    return deadline is None or time.perf_counter() < deadline


def _csv(path, headers, rows):
    with Path(path).open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(rows)


def _point_key(x):
    return tuple(np.round(np.asarray(x), 12))


def _is_feasible(row):
    return (row.get("status") == "FEASIBLE_SAMPLED" and row.get("complete") is True
            and row.get("min_gap") is not None and row["min_gap"] >= row.get("required_gap_m", .01))


def _best(rows):
    return min((r for r in rows if _is_feasible(r)), key=lambda r: r["length"], default=None)


def _practical_progress(records, state, settings):
    """相对上次显著改进累计判断；初始样本未完成或尚无带裕量候选时不停。"""
    best = _best(records)
    if best is None:
        state.update(last_significant_length=None, no_significant_improvement=0)
        return False, False
    anchor = state.get("last_significant_length")
    improved = anchor is None or best["length"] < anchor*(1-settings["improvement_rel"])
    if improved:
        state.update(last_significant_length=best["length"], no_significant_improvement=0)
    else:
        state["no_significant_improvement"] = state.get("no_significant_improvement", 0)+1
    stop = (len(records) >= settings["initial_samples"] and
            state["no_significant_improvement"] >= settings["patience"])
    return improved, stop


def _max_arc_radius(center, start, turn, angle, radius):
    phi = math.atan2(start[1]-center[1], start[0]-center[0])
    values = [0.0, float(angle)]
    stationary = math.atan2(center[1], center[0])
    for candidate in (stationary, stationary+math.pi):
        delta = (turn*(candidate-phi)) % (2*math.pi)
        if delta <= angle+1e-12:
            values.append(delta)
    points = [center + radius*np.array([math.cos(phi+turn*u), math.sin(phi+turn*u)]) for u in values]
    return max(float(np.linalg.norm(p)) for p in points)


def geometry_audit(path):
    """独立回代相切、端点与完整圆弧径向驻点；不靠绘图网格判圆域。"""
    b, rho = path.params.b, path.rho
    r = lambda theta: b*theta*np.array([math.cos(theta), math.sin(theta)])
    derivative = lambda theta: b*np.array([math.cos(theta)-theta*math.sin(theta),
                                           math.sin(theta)+theta*math.cos(theta)])
    ta = -derivative(path.theta_a); ta /= np.linalg.norm(ta)
    tb = -derivative(path.theta_c); tb /= np.linalg.norm(tb)
    rotate = lambda v: np.array([-v[1], v[0]])
    tangent_a = -rotate((path.A-path.O1)/(2*rho))
    tangent_c1 = -rotate((path.C-path.O1)/(2*rho))
    tangent_c2 = rotate((path.C-path.O2)/rho)
    tangent_b = rotate((path.B-path.O2)/rho)
    first_max = _max_arc_radius(path.O1, path.A, -1, path.alpha1, 2*rho)
    second_max = _max_arc_radius(path.O2, path.C, 1, path.alpha2, rho)
    residuals = {
        "entry_position_m": float(np.linalg.norm(path.A-r(path.theta_a))),
        "exit_position_m": float(np.linalg.norm(path.B+r(path.theta_c))),
        "external_tangency_m": abs(float(np.linalg.norm(path.O2-path.O1))-3*rho),
        "connection_point_m": float(np.linalg.norm(path.C-(path.O1+2*path.O2)/3)),
        "radius_ratio": abs(float(np.linalg.norm(path.A-path.O1)/np.linalg.norm(path.B-path.O2))-2),
        "entry_tangent": float(np.linalg.norm(tangent_a-ta)),
        "middle_tangent": float(np.linalg.norm(tangent_c1-tangent_c2)),
        "exit_tangent": float(np.linalg.norm(tangent_b-tb)),
        "first_arc_end_radius_m": abs(float(np.linalg.norm(path.C-path.O1))-2*rho),
        "second_arc_start_radius_m": abs(float(np.linalg.norm(path.C-path.O2))-rho),
    }
    bounds = (np.all(np.asarray(path.x) > 0) and np.all(np.asarray(path.x) <= 1+1e-12)
              and rho > 0 and 0 < path.alpha1 < 2*math.pi and 0 < path.alpha2 < 2*math.pi)
    passed = bool(bounds and max(residuals.values()) <= 1e-7
                  and max(first_max, second_max) <= path.radius+1e-8)
    return {"passed": passed, "maximum_absolute_residuals": residuals,
            "max_arc_radii_m": [first_max, second_max], "circle_radius_m": path.radius,
            "residual_tolerance": 1e-7, "circle_tolerance_m": 1e-8,
            "circle_check_method": "端点及完整圆弧中的径向驻点"}


class ConstraintGP:
    """对完整采样 H 的零均值标准化 Matérn-5/2 GP，极大似然拟合长度尺度。"""

    @staticmethod
    def _kernel(first, second, scales):
        distance = np.linalg.norm((first[:, None, :]-second[None, :, :])/scales, axis=2)
        root5 = math.sqrt(5)*distance
        return (1+root5+5*distance*distance/3)*np.exp(-root5)

    def fit(self, x, y):
        self.x = np.asarray(x, dtype=float)
        y = np.asarray(y, dtype=float)
        self.mean = float(y.mean())
        self.scale = max(float(y.std()), 1e-5)
        z = (y-self.mean)/self.scale
        noise = 1e-6

        def likelihood(log_scales):
            scales = np.exp(log_scales)
            matrix = self._kernel(self.x, self.x, scales)+noise*np.eye(len(self.x))
            try:
                factor = cho_factor(matrix, lower=True, check_finite=False)
                alpha = cho_solve(factor, z, check_finite=False)
                return .5*float(z@alpha)+float(np.log(np.diag(factor[0])).sum())
            except np.linalg.LinAlgError:
                return 1e20

        result = minimize(likelihood, np.log([.25, .25]), method="L-BFGS-B",
                          bounds=[(math.log(.015), math.log(2))]*2,
                          options={"maxiter": 30, "ftol": 1e-8})
        self.scales = np.exp(result.x)
        self.factor = cho_factor(self._kernel(self.x, self.x, self.scales)+noise*np.eye(len(self.x)),
                                 lower=True, check_finite=False)
        self.alpha = cho_solve(self.factor, z, check_finite=False)
        self.fit_report = {"length_scales": self.scales.tolist(), "mean": self.mean,
                           "scale": self.scale, "training_count": len(y),
                           "optimizer_success": bool(result.success), "noise_variance": noise}
        return self

    def predict(self, candidates):
        cross = self._kernel(np.asarray(candidates), self.x, self.scales)
        mean = self.mean+self.scale*(cross@self.alpha)
        variance = np.maximum(1-np.einsum("ij,ji->i", cross, cho_solve(self.factor, cross.T)), 1e-12)
        return mean, self.scale*np.sqrt(variance)


def candidate_paths(count, seed, params, radius, center=None, side=2.0, deadline=None):
    sampler = qmc.Sobol(d=2, scramble=True, seed=int(seed))
    points = sampler.random_base2(int(math.ceil(math.log2(max(2,count)))))[:count]
    low = np.full(2, EPS_X) if center is None else np.maximum(EPS_X, np.asarray(center)-side/2)
    high = np.ones(2) if center is None else np.minimum(1, np.asarray(center)+side/2)
    paths = []
    for index, point in enumerate(low+(high-low)*points):
        if index % 32 == 0 and not _remaining(deadline):
            break
        try:
            path = build_path(point, params=params, radius=radius)
            if path.geometry_report()["passed"]:
                paths.append(path)
        except (ValueError, ArithmeticError, RuntimeError):
            pass
    return paths


def shared_initial_points(count, seed, params, radius, deadline):
    points = [np.ones(2)]
    if count <= 1:
        return points
    # 从基准路径附近加入对称内移候选，避免初始样本只覆盖过紧的调头弧。
    points.append(build_path((.95, .95), params, radius).x.copy())
    if len(points) >= count:
        return points
    for round_index in range(8):
        paths = candidate_paths(max(64, count*8), seed+round_index, params, radius, deadline=deadline)
        for path in paths:
            if _point_key(path.x) not in {_point_key(p) for p in points}:
                points.append(np.asarray(path.x).copy())
                if len(points) >= count:
                    return points
        if not _remaining(deadline):
            break
    return points


def _checkpoint(out, method, records, settings, algorithm_state, status):
    save_json(out/f"{method}_检查点.json", {"method": method, "settings": settings,
              "status": status, "records": records, "algorithm_state": algorithm_state,
              "scope": SCOPE})


def _motion_trial(path, step, deadline, method, index, out, clearance_margin=.01):
    started = time.perf_counter()
    try:
        motion = evaluate_motion(path, step=step, deadline=deadline, refine=False)
    except (ValueError, ArithmeticError, RuntimeError) as error:
        motion = {"status": "UNRESOLVED", "complete": False, "min_gap": None,
                  "sample_count": 0, "reason": f"{type(error).__name__}: {error}"}
    motion["required_gap_m"] = clearance_margin
    if motion.get("status") == "FEASIBLE_SAMPLED" and motion["min_gap"] < clearance_margin:
        motion["status"] = "CLEARANCE_MARGIN_NOT_MET"
    motion_file = f"{method}_{index:04d}_运动记录.json"
    save_json(out/motion_file, motion)
    row = {"evaluation": index, "method": method, "x": np.asarray(path.x).tolist(),
           "length": float(path.length), "elapsed_seconds": time.perf_counter()-started,
           "motion_record": motion_file, "path": path.to_dict(),
           "motion_refine": False, "motion_step": step}
    row.update({key: value for key, value in motion.items() if key != "records"})
    row["complete"] = bool(row.get("complete", False))
    if row.get("min_gap") is not None and not math.isfinite(float(row["min_gap"])):
        row.update(status="UNRESOLVED", complete=False, min_gap=None,
                   reason="运动检查返回非有限间隙，未作为可行数据使用")
    print(f"[{method}] 真实运动计算 {index}：{row['status']}，L={path.length:.6f} m，"
          f"用时 {row['elapsed_seconds']:.1f} s", flush=True)
    return row


def optimize_cbo(initial, params, radius, settings, deadline, out, resumed=None):
    records = list(resumed.get("records", [])) if resumed else []
    tr = dict(resumed.get("algorithm_state", {})) if resumed else {}
    tr = {"side": .8, "successes": 0, "failures": 0, "restarts": 0, **tr}
    limit, seed = settings["evaluations"], settings["seed"]
    center = np.asarray(tr.get("center", [1.,1.]))
    status = "EVALUATION_BUDGET_EXHAUSTED"
    while len(records) < limit:
        if not _remaining(deadline):
            status = "TIME_BUDGET_EXHAUSTED"; break
        used = {_point_key(row["x"]) for row in records}
        initial_left = [x for x in initial if _point_key(x) not in used]
        if initial_left:
            path = build_path(initial_left[0], params=params, radius=radius)
        else:
            best = _best(records)
            complete = [r for r in records if r.get("complete") and r.get("min_gap") is not None]
            if tr["side"] < 2**-7:
                distant = candidate_paths(settings["candidates"], seed+50000+tr["restarts"],
                                          params, radius, deadline=deadline)
                if distant:
                    distances = [min(np.linalg.norm(np.asarray(p.x)-np.asarray(r["x"])) for r in records) for p in distant]
                    center = np.asarray(distant[int(np.argmax(distances))].x)
                tr.update(side=.8, successes=0, failures=0, restarts=tr["restarts"]+1)
            paths = candidate_paths(settings["candidates"], seed+1000+len(records), params, radius,
                                    center=center, side=tr["side"], deadline=deadline)
            paths = [p for p in paths if _point_key(p.x) not in used]
            if not paths:
                if not _remaining(deadline):
                    status = "TIME_BUDGET_EXHAUSTED"; break
                tr["side"] /= 2
                if tr["restarts"] >= 10:
                    status = "CANDIDATE_SEARCH_EXHAUSTED"; break
                continue
            if len(complete) >= 2:
                gp = ConstraintGP().fit([r["x"] for r in complete], [r["min_gap"] for r in complete])
                mean, sigma = gp.predict([p.x for p in paths])
                probability = ndtr((mean-settings["clearance_margin"])/np.maximum(sigma,1e-12))
                reference_length = build_path(params=params, radius=radius).length
                improvement = (np.maximum((best["length"]-np.array([p.length for p in paths]))/reference_length,0)*probability
                               if best else probability)
                boundary = sigma/gp.scale*4*probability*(1-probability)
                weight = .1*(1-len(records)/limit)
                score = improvement+weight*boundary
                path = paths[int(np.argmax(score))]
                tr["last_gp"] = gp.fit_report
            else:
                path = max(paths, key=lambda p:min(np.linalg.norm(np.asarray(p.x)-r["x"]) for r in records))
        previous_best = _best(records)
        previous_complete = [r["min_gap"] for r in records if r.get("complete") and r.get("min_gap") is not None]
        row = _motion_trial(path, settings["motion_step"], deadline, "CBO", len(records)+1, out,
                            settings["clearance_margin"])
        records.append(row)
        improved, practical_stop = _practical_progress(records, tr, settings)
        if previous_best is None and not improved and row.get("complete") and row.get("min_gap") is not None:
            improved = not previous_complete or row["min_gap"] > max(previous_complete)+1e-8
        if improved:
            tr["successes"] += 1; tr["failures"] = 0; center = np.asarray(row["x"])
        else:
            tr["failures"] += 1; tr["successes"] = 0
        if tr["successes"] >= 3:
            tr.update(side=min(2*tr["side"],1.6),successes=0,failures=0)
        elif tr["failures"] >= 4:
            tr.update(side=tr["side"]/2,successes=0,failures=0)
        tr["center"] = center.tolist()
        _checkpoint(out,"CBO",records,settings,tr,"RUNNING")
        if practical_stop:
            status="PRACTICAL_TOLERANCE_REACHED"; break
    _checkpoint(out,"CBO",records,settings,tr,status)
    return {"method":"CBO","status":status,"records":records,"algorithm_state":tr}


def _merit(row):
    if _is_feasible(row):
        return (0, row["length"])
    if row.get("complete") and row.get("min_gap") is not None:
        return (1, -row["min_gap"])
    return (2, float("inf"))


def optimize_mads(initial, params, radius, settings, deadline, out, resumed=None):
    """多起点网格自适应直接搜索：整数网格正生成轮询方向、机会式接受与网格细化。"""
    records = list(resumed.get("records", [])) if resumed else []
    saved = resumed.get("algorithm_state", {}) if resumed else {}
    mesh = float(saved.get("mesh", .0625))
    poll = int(saved.get("poll", 0)); restarts = int(saved.get("restarts", 0))
    local_index = saved.get("local_index")
    limit, seed = settings["evaluations"], settings["seed"]
    status = "EVALUATION_BUDGET_EXHAUSTED"
    while len(records) < limit:
        if not _remaining(deadline):
            status="TIME_BUDGET_EXHAUSTED"; break
        used = {_point_key(r["x"]) for r in records}
        restart_target = False
        initial_left = [x for x in initial if _point_key(x) not in used]
        if initial_left:
            paths = [build_path(initial_left[0],params=params,radius=radius)]
            is_poll = False
        else:
            is_poll = True
            if local_index is None:
                local_index = min(range(len(records)), key=lambda i:_merit(records[i]))
            incumbent = records[local_index]
            center = np.asarray(incumbent["x"])
            angle = (seed*.137+poll*math.pi*(3-math.sqrt(5))) % (2*math.pi)
            extent = max(1,int(math.ceil(1/math.sqrt(mesh))))
            direction = np.rint(extent*np.array([math.cos(angle),math.sin(angle)])).astype(int)
            if not np.any(direction):
                direction[0]=1
            orthogonal = np.array([-direction[1],direction[0]])
            directions = (direction,orthogonal,-direction,-orthogonal)
            paths = []
            for d in directions:
                trial = center+mesh*d
                if np.any(trial<EPS_X) or np.any(trial>1) or _point_key(trial) in used:
                    continue
                try:
                    paths.append(build_path(trial,params=params,radius=radius))
                except (ValueError,ArithmeticError,RuntimeError):
                    pass
            if mesh < 2**-16 or (not paths and poll%8==7):
                global_paths = candidate_paths(settings["candidates"],seed+90000+restarts,params,radius,deadline=deadline)
                global_paths = [p for p in global_paths if _point_key(p.x) not in used]
                paths = [max(global_paths,key=lambda p:min(np.linalg.norm(np.asarray(p.x)-r["x"]) for r in records))] if global_paths else []
                mesh=.0625; restarts+=1; is_poll=False; restart_target=True
        improved = False
        for path in paths:
            if len(records)>=limit or not _remaining(deadline):
                break
            previous = records[local_index] if local_index is not None else None
            row = _motion_trial(path,settings["motion_step"],deadline,"MADS",len(records)+1,out,
                                settings["clearance_margin"])
            records.append(row)
            _, practical_stop = _practical_progress(records, saved, settings)
            improved = previous is None or _merit(row)<_merit(previous)
            if improved or restart_target:
                local_index = len(records)-1
            saved = {**saved,"mesh":mesh,"poll":poll,"restarts":restarts,"local_index":local_index}
            _checkpoint(out,"MADS",records,settings,saved,"RUNNING")
            if practical_stop:
                status="PRACTICAL_TOLERANCE_REACHED"; break
            if is_poll and improved:
                break
        if is_poll:
            mesh = min(.25,mesh*4) if improved else mesh/4
            poll += 1
        if status=="PRACTICAL_TOLERANCE_REACHED":
            break
        if poll>max(100,limit*50):
            status="POLL_SEARCH_EXHAUSTED"; break
    saved={**saved,"mesh":mesh,"poll":poll,"restarts":restarts,"local_index":local_index}
    _checkpoint(out,"MADS",records,settings,saved,status)
    return {"method":"MADS","status":status,"records":records,"algorithm_state":saved}


def export_q4_workbook(path, positions, speeds):
    """只填题目模板的数据单元格，随后逐数值及原标签重读核对。"""
    from openpyxl import load_workbook
    positions,speeds = np.asarray(positions),np.asarray(speeds)
    if positions.shape!=(201,224,2) or speeds.shape!=(201,224) or not np.isfinite(positions).all() or not np.isfinite(speeds).all():
        raise ValueError("result4 要求完整有限的 201×224 状态")
    source=PROJECT/"题目/附件/result4.xlsx"
    if Path(path).resolve()==source.resolve():
        raise ValueError("输出不能覆盖题目模板")
    workbook=load_workbook(source)
    if workbook.sheetnames!=["位置","速度"]:
        raise ValueError("result4 工作表标签不匹配")
    normalize=lambda text:str(text or "").replace(" ","").replace("　","")
    label=lambda i:"龙头" if i==0 else "龙尾" if i==222 else "龙尾（后）" if i==223 else f"第{i}节龙身"
    arrays={"位置":positions.reshape(201,448).T,"速度":speeds.T}
    original={}
    for name,values in arrays.items():
        sheet=workbook[name]
        if (sheet.max_row,sheet.max_column)!=(len(values)+1,202):
            raise ValueError("result4 模板尺寸不匹配")
        original[name]=([sheet.cell(r,1).value for r in range(1,sheet.max_row+1)],
                        [sheet.cell(1,c).value for c in range(1,203)])
        for index,t in enumerate(range(-100,101),start=2):
            if normalize(sheet.cell(1,index).value)!=f"{t}s":
                raise ValueError("result4 时间标签不匹配")
        for i in range(224):
            expected=[label(i)+"x(m)",label(i)+"y(m)"] if name=="位置" else [label(i)+"(m/s)"]
            for k,wanted in enumerate(expected):
                if normalize(sheet.cell(2+i*len(expected)+k,1).value)!=wanted:
                    raise ValueError("result4 把手标签不匹配")
        for r,row in enumerate(values,start=2):
            for c,value in enumerate(row,start=2):
                cell=sheet.cell(r,c,float(value)); cell.number_format="0.000000"
    workbook.save(path); workbook.close()
    check=load_workbook(path,read_only=True,data_only=True)
    count=0; error=0.0
    try:
        for name,values in arrays.items():
            sheet=check[name]
            label_rows=sheet.iter_rows(values_only=True)
            header=next(label_rows)
            column=[header[0]]+[row[0] for row in label_rows]
            if original[name]!=(column,list(header)):
                raise ArithmeticError("result4 原始标签改变")
            for i,row in enumerate(sheet.iter_rows(min_row=2,max_row=len(values)+1,min_col=2,max_col=202,values_only=True)):
                actual=np.asarray(row,dtype=float)
                if not np.isfinite(actual).all():
                    raise ArithmeticError("result4 重读出现非有限数值")
                error=max(error,float(np.max(np.abs(actual-values[i])))); count+=len(row)
    finally:
        check.close()
    if error>1e-10:
        raise ArithmeticError("result4 导出数值核对失败")
    return {"passed":True,"checked_values":count,"max_absolute_export_error":error,
            "original_labels_preserved":True,"template_sha256":_hash(source),"output_sha256":_hash(path)}


def state_residuals(path, states):
    positions=np.stack([s["positions"] for s in states])
    velocities=np.stack([s["velocities"] for s in states])
    tangents=np.stack([s["tangents"] for s in states])
    coordinates=np.stack([s["s"] for s in states])
    displacement=np.diff(positions,axis=1)
    lengths=np.linalg.norm(displacement,axis=2)
    rates=np.sum(displacement*np.diff(velocities,axis=1),axis=2)/lengths
    transverse=tangents[:,:,0]*velocities[:,:,1]-tangents[:,:,1]*velocities[:,:,0]
    residuals={"hole_distance_m":float(np.max(np.abs(lengths-path.params.distances))),
               "rigid_length_rate_m_s":float(np.max(np.abs(rates))),
               "transverse_speed_m_s":float(np.max(np.abs(transverse))),
               "head_speed_m_s":float(np.max(np.abs(np.linalg.norm(velocities[:,0],axis=1)-1)))}
    finite=all(np.isfinite(a).all() for a in (positions,velocities,tangents,coordinates))
    ordered=bool(np.all(np.diff(coordinates,axis=1)<0))
    return {"passed":bool(finite and ordered and max(residuals.values())<=1e-7),
            "finite":bool(finite),"ordered":ordered,"sample_count":len(states),
            "maximum_absolute_residuals":residuals,"residual_tolerance":1e-7}


def event_continuity_check(path, deadline, out, quick=False):
    started=time.perf_counter()
    result={"status":"RUNNING","passed":None,"scope":"逐把手、逐连接点两侧连续性数值检验"}
    try:
        events=crossing_events(path,deadline=deadline)
    except (ValueError,ArithmeticError,RuntimeError,TimeoutError) as error:
        return {**result,"status":"UNRESOLVED","reason":str(error)}
    chosen=[e for e in events if e["handle"] in (0,1,223)] if quick else events
    rows=[]; initial_h=1e-5
    for event in chosen:
        if not _remaining(deadline):
            result["status"]="BUDGET_EXHAUSTED"; break
        try:
            handle=int(event["handle"])
            refinements=[]
            for level in range(4):
                _deadline(deadline)
                h=initial_h*10.**(-level)
                left=chain_state(path,event["head_s"]-h)
                right=chain_state(path,event["head_s"]+h)
                position_error=float(np.max(np.linalg.norm(
                    left["positions"]+h*left["velocities"]-right["positions"]+h*right["velocities"],axis=1)))
                velocity_difference=float(np.max(np.linalg.norm(left["velocities"]-right["velocities"],axis=1)))
                root_error=abs(float((left["s"][handle]+right["s"][handle])/2-event["path_s"]))
                ordered=bool(np.all(np.diff(left["s"])<0) and np.all(np.diff(right["s"])<0))
                passed=bool(position_error<=1e-6 and velocity_difference<=1e-3 and root_error<=1e-6 and ordered)
                refinements.append({"delta_s":h,"position_extrapolation_error_m":position_error,
                                    "velocity_side_difference_m_s":velocity_difference,"event_root_error_m":root_error})
                if passed:
                    break
            rows.append({**event,"position_extrapolation_error_m":position_error,
                         "velocity_side_difference_m_s":velocity_difference,
                         "event_root_error_m":root_error,"ordered":ordered,"passed":passed,
                         "delta_s":h,"step_refinement":refinements})
        except TimeoutError:
            result["status"]="BUDGET_EXHAUSTED"; break
        except (ValueError,ArithmeticError,RuntimeError) as error:
            rows.append({**event,"passed":False,"error":str(error)})
    else:
        result["status"]="PILOT_ONLY" if quick else "PASSED"
    any_failure=any(r["passed"] is False for r in rows)
    if any_failure:
        result.update(status="FAILED",passed=False)
    elif len(rows)==len(events):
        result.update(status="PASSED",passed=True)
    result.update(all_event_count=len(events),requested_event_count=len(chosen),completed_event_count=len(rows),
                  delta_s=initial_h,minimum_delta_s=initial_h*1e-3,
                  adaptively_refined_event_count=sum(len(r.get("step_refinement",[]))>1 for r in rows),
                  position_tolerance_m=1e-6,velocity_difference_tolerance_m_s=1e-3,
                  event_root_tolerance_m=1e-6,elapsed_seconds=time.perf_counter()-started,events=rows)
    save_json(out/"跨段事件连续性检查.json",result)
    _csv(out/"跨段事件作图与检验数据.csv",
         ["handle_0based","junction","head_s_m","path_s_m","position_error_m","velocity_side_difference_m_s","event_root_error_m","passed"],
         [[r["handle"],r["junction"],r["head_s"],r["path_s"],r.get("position_extrapolation_error_m"),
           r.get("velocity_side_difference_m_s"),r.get("event_root_error_m"),r["passed"]] for r in rows])
    return result


def _path_samples(path):
    primitive=lambda theta:.5*(theta*math.sqrt(1+theta*theta)+math.asinh(theta))
    theta_r=path.radius/path.params.b
    entry=-path.params.b*(primitive(theta_r)-primitive(path.theta_a))
    exit_s=path.length+path.params.b*(primitive(theta_r)-primitive(path.theta_c))
    segments=[np.linspace(min(entry,-2),0,200),np.linspace(0,path.ell1,180),
              np.linspace(path.ell1,path.length,120),np.linspace(path.length,max(exit_s,path.length+2),200)]
    return [(s,np.stack([path.point_tangent(float(value))[0] for value in s])) for s in segments]


def plot_path_comparison(out,reference,path):
    plt=configure_plotting()
    from matplotlib.patches import Circle
    fig,axes=plt.subplots(1,2,figsize=(11,5.4))
    data=[]
    for ax,current,title in zip(axes,(reference,path),("(a) 边界接入基准","(b) 已检查的可行候选")):
        ax.add_patch(Circle((0,0),current.radius,fill=False,color=".55",linestyle="--",linewidth=.9))
        for segment,(s,points) in enumerate(_path_samples(current)):
            color=("#84929D","#2563A6","#E88B26","#84929D")[segment]
            ax.plot(points[:,0],points[:,1],color=color,lw=1.6 if segment in(1,2) else 1)
            for value,point in zip(s,points):
                data.append([title,segment,float(value),*point])
        for label,point in (("A",current.A),("B",current.B),("C",current.C),(r"O_1",current.O1),(r"O_2",current.O2)):
            ax.plot(*point,"o",ms=3,color="#303840")
            ax.annotate(f"${label}$",point,xytext=(4,4),textcoords="offset points",fontsize=9)
        ax.plot([current.O1[0],current.A[0]],[current.O1[1],current.A[1]],"--",color=".7",lw=.7)
        ax.plot([current.O2[0],current.B[0]],[current.O2[1],current.B[1]],"--",color=".7",lw=.7)
        text=(rf"$\ell_1={scientific_tex(current.ell1)}\,\mathrm{{m}}$"+"\n"
              +rf"$\ell_2={scientific_tex(current.ell2)}\,\mathrm{{m}}$"+"\n"
              +rf"$L_{{\mathrm{{arc}}}}={scientific_tex(current.length)}\,\mathrm{{m}}$"+"\n"
              +rf"$R_1={scientific_tex(2*current.rho)}\,\mathrm{{m}},\quad R_2={scientific_tex(current.rho)}\,\mathrm{{m}}$")
        ax.text(.02,.02,text,transform=ax.transAxes,fontsize=8,bbox={"facecolor":"white","edgecolor":"none","alpha":.85})
        ax.set_title(title); ax.set_aspect("equal")
        ax.set_xlim(-reference.radius-.7,reference.radius+.7); ax.set_ylim(-reference.radius-.7,reference.radius+.7)
        ax.set_xlabel(r"$x\,/\,\mathrm{m}$"); ax.set_ylabel(r"$y\,/\,\mathrm{m}$")
    reduction=(reference.length-path.length)/reference.length
    fig.suptitle(rf"圆弧长度相对变化：$\Delta L/L_{{\mathrm{{ref}}}}={scientific_tex(reduction)}$",fontsize=12)
    fig.tight_layout()
    save_figure(fig,out,"R5_双圆弧路径调整前后对照")
    plt.close(fig)
    _csv(out/"R5_路径作图数据.csv",["path","segment","s_m","x_m","y_m"],data)


def plot_key_configurations(out,path):
    plt=configure_plotting()
    from matplotlib.patches import Circle
    times=[0.,path.ell1,path.length]
    states=[chain_state(path,t) for t in times]
    fig,axes=plt.subplots(1,3,figsize=(12,4.5))
    for ax,current,t,title in zip(axes,states,times,("(a) 龙头进入大圆弧","(b) 龙头到达两弧接点","(c) 龙头离开小圆弧")):
        draw_chain(ax,current["positions"],path.params,highlight=(0,))
        ax.add_patch(Circle((0,0),path.radius,fill=False,linestyle="--",color=".5",linewidth=.8))
        ax.set_xlim(-path.radius-.7,path.radius+.7); ax.set_ylim(-path.radius-.7,path.radius+.7)
        ax.set_title(title+"\n"+rf"$t={scientific_tex(t)}\,\mathrm{{s}}$",fontsize=9)
        ax.set_xlabel(r"$x\,/\,\mathrm{m}$"); ax.set_ylabel(r"$y\,/\,\mathrm{m}$")
    fig.suptitle("调头区域局部队形（矩形采用实际板长和板宽）",fontsize=12)
    fig.tight_layout(); save_figure(fig,out,"R6_调头关键构型"); plt.close(fig)
    np.savez_compressed(out/"R6_关键构型作图数据.npz",times=times,
                        positions=np.stack([s["positions"] for s in states]),
                        velocities=np.stack([s["velocities"] for s in states]),
                        s=np.stack([s["s"] for s in states]))


def plot_optimization(out,results,reference_length):
    plt=configure_plotting(); fig,ax=plt.subplots(figsize=(8,4.8))
    from matplotlib.ticker import MaxNLocator
    csv_rows=[]; total_feasible=0
    for result,color in zip(results,("#2563A6","#E88B26")):
        running=None; curve=[]
        for row in result["records"]:
            if _is_feasible(row):
                running=row["length"] if running is None else min(running,row["length"])
            curve.append(np.nan if running is None else running)
            csv_rows.append([result["method"],row["evaluation"],row["status"],row["length"],row.get("min_gap"),running])
        if any(math.isfinite(v) for v in curve):
            total_feasible+=1
            ax.step(np.arange(1,len(curve)+1),curve,where="post",label=result["method"],color=color)
    ax.axhline(reference_length,color=".6",linestyle="--",label="边界接入的几何弧长")
    refined_file=out/"候选加密检查.json"
    if refined_file.is_file():
        rejected={_point_key(item["x"]) for item in json.loads(refined_file.read_text(encoding="utf-8"))
                  if not _is_feasible(item["motion"])}
        rejected_rows=[row for result in results for row in result["records"]
                       if _is_feasible(row) and _point_key(row["x"]) in rejected]
        if rejected_rows:
            ax.scatter([row["evaluation"] for row in rejected_rows],[row["length"] for row in rejected_rows],
                       marker="x",color="#B34C35",zorder=5,label="加密复核未通过或未完成")
    if not total_feasible:
        ax.text(.5,.5,"尚无完成运动检查的可行候选",ha="center",transform=ax.transAxes)
    ax.set_xlabel("实际运动检查次数（含失败及未完成的调用）")
    ax.set_ylabel(r"$L_{\mathrm{best,sampled}}\,/\,\mathrm{m}$")
    ax.set_title("V2 优化阶段的采样改进记录")
    ax.xaxis.set_major_locator(MaxNLocator(integer=True))
    ax.text(.02,.04,"曲线记录优化网格上的采样候选；最终路径以加密复核为准",transform=ax.transAxes,
            fontsize=8,bbox={"facecolor":"white","edgecolor":"none","alpha":.85})
    ax.legend(fontsize=9,loc="upper right"); fig.tight_layout()
    save_figure(fig,out,"V2_路径优化改进过程"); plt.close(fig)
    _csv(out/"V2_优化改进作图数据.csv",["method","evaluation","status","arc_length_m","min_gap_m","best_feasible_length_m"],csv_rows)


def budget_prefix_report(results,levels=(100,200,400)):
    rows=[]
    for result in results:
        for level in levels:
            reached=len(result["records"])>=level
            best=_best(result["records"][:level]) if reached else None
            rows.append({"method":result["method"],"budget":level,
                         "status":"OBSERVED_PREFIX" if reached else "NOT_EXECUTED",
                         "best_feasible_length":best["length"] if best else None,
                         "best_min_gap":best["min_gap"] if best else None})
    return {"scope":"同一实际运行轨迹的嵌套预算前缀；不是独立的 100/200/400 次配对重跑，也不证明收敛。",
            "prefixes":rows}


def _load_resume(directory,method,settings):
    if directory is None:
        return None
    file=Path(directory)/f"{method}_检查点.json"
    if not file.is_file():
        return None
    checkpoint=json.loads(file.read_text(encoding="utf-8-sig"))
    for key in ("seed","motion_step","pitch","radius","candidates","initial_samples",
                "clearance_margin","improvement_rel","patience","motion_check_version"):
        if checkpoint["settings"].get(key)!=settings[key]:
            raise ValueError(f"恢复检查点参数不一致：{key}")
    if len(checkpoint["records"])>settings["evaluations"]:
        raise ValueError("新预算小于检查点已有计算次数")
    for row in checkpoint["records"]:
        if row.get("motion_record") and not Path(row["motion_record"]).is_absolute():
            row["motion_record"]=str((Path(directory)/row["motion_record"]).resolve())
    return checkpoint


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir",type=Path,default=HERE.parent/"结果"/"运行记录"/"第四问",help="输出基目录，运行时新建时间戳子目录")
    parser.add_argument("--quick",action="store_true",help="缩小优化预算，跨段检验仅作三把手试验")
    parser.add_argument("--evaluations",type=int,default=None,help="每种算法的真实运动调用上限；默认40，quick默认4，可按实用精度提前停止")
    parser.add_argument("--clearance-margin",type=float,default=.01,help="全队检查要求的最小 SAT 分离量，默认0.01 m")
    parser.add_argument("--improvement-rel",type=float,default=.001,help="显著弧长改进的相对阈值，默认0.001，即0.1%%")
    parser.add_argument("--patience",type=int,default=8,help="有带裕量候选后，连续无显著改进的次数上限")
    parser.add_argument("--initial-samples",type=int,default=None)
    parser.add_argument("--candidates",type=int,default=None)
    parser.add_argument("--motion-step",type=float,default=None,help="优化阶段运动网格步长 m，默认2，quick默认8")
    parser.add_argument("--validation-step",type=float,default=None,help="候选加密运动网格，默认0.5，quick默认2")
    parser.add_argument("--time-budget",type=float,default=None,help="数值任务墙钟软上限秒数，最终导出不强制中断；默认600，quick默认180")
    parser.add_argument("--seed",type=int,default=2024)
    parser.add_argument("--skip-validation",action="store_true",help="跳过MADS及跨段独立检验；候选加密运动检查仍执行")
    parser.add_argument("--resume-dir",type=Path,help="从已有时间戳目录的检查点继续，允许增加计算预算")
    args=parser.parse_args(argv)
    settings={"evaluations":args.evaluations if args.evaluations is not None else (4 if args.quick else 40),
              "initial_samples":args.initial_samples if args.initial_samples is not None else (2 if args.quick else 8),
              "candidates":args.candidates if args.candidates is not None else (128 if args.quick else 512),
              "motion_step":args.motion_step if args.motion_step is not None else (8. if args.quick else 2.),
              "seed":args.seed,"pitch":1.7,"radius":4.5,
              "motion_check_version":MOTION_CHECK_VERSION,
              "clearance_margin":args.clearance_margin,"improvement_rel":args.improvement_rel,"patience":args.patience}
    seconds=args.time_budget if args.time_budget is not None else (180. if args.quick else 600.)
    validation_step=args.validation_step if args.validation_step is not None else (2. if args.quick else .5)
    if min(settings["evaluations"],settings["initial_samples"],settings["candidates"],settings["motion_step"],validation_step,seconds)<=0:
        parser.error("预算、样本数和步长必须为正")
    if not math.isfinite(args.clearance_margin) or args.clearance_margin<0:
        parser.error("--clearance-margin 必须是有限非负数")
    if not math.isfinite(args.improvement_rel) or not 0<args.improvement_rel<1 or args.patience<1:
        parser.error("--improvement-rel 须在(0,1)内，--patience 须为正整数")
    settings["initial_samples"]=min(settings["initial_samples"],settings["evaluations"])
    out=prepare_output_dir(args.output_dir)
    started=time.perf_counter(); deadline=started+seconds
    params=Parameters(pitch=1.7)
    reference=build_path(params=params,radius=4.5)
    initial=shared_initial_points(settings["initial_samples"],args.seed,params,4.5,deadline)
    if args.quick and len(initial)>=2:
        initial[1]=np.array([.8,.8])
    save_json(out/"运行配置.json",{"settings":settings,"initial_points":initial,"quick":args.quick,
              "validation_step":validation_step,"time_budget":seconds,"scope":SCOPE,
              "skip_validation":args.skip_validation,"resume_directory":str(args.resume_dir) if args.resume_dir else None,
              "optimizer_motion_refine":False,"finalist_motion_refine":True,
              "source_sha256":{name:_hash(HERE/name) for name in ("问题1代码.py","问题2代码.py","问题4代码.py")},
              "input_sha256":{"A题.pdf":_hash(PROJECT/"题目/A题.pdf"),"result4.xlsx":_hash(PROJECT/"题目/附件/result4.xlsx")}})
    # 同预算、同初始点、同网格；分别分配相同墙钟上限，不重用对方的真实运动结果。
    optimization_share=.60*seconds
    per_method=optimization_share if args.skip_validation else optimization_share/2
    print(f"[第四问] 每算法预算 {settings['evaluations']} 次，保留加密检查及表格输出时间。",flush=True)
    cbo=optimize_cbo(initial,params,4.5,settings,min(deadline,time.perf_counter()+per_method),out,
                     _load_resume(args.resume_dir,"CBO",settings))
    results=[cbo]
    if not args.skip_validation:
        mads=optimize_mads(initial,params,4.5,settings,min(deadline,time.perf_counter()+per_method),out,
                           _load_resume(args.resume_dir,"MADS",settings))
        results.append(mads)
    save_json(out/"实际预算前缀比较.json",budget_prefix_report(results))
    plot_optimization(out,results,reference.length)
    summary={"status":"NO_FEASIBLE_CANDIDATE","scope":SCOPE,"settings":settings,
             "reference_length_m":reference.length,"reference_path":reference.to_dict(),
             "algorithms":[{k:v for k,v in r.items() if k!="records"} for r in results],
             "fair_comparison":{"same_evaluation_budget":True,"same_initial_points":True,
                                "same_motion_grid":True,"same_wall_clock_cap":True,
                                "same_practical_stopping_rule":True,
                                "scope":"两算法上限和停止规则一致；按实际调用数报告，不将不同调用数视作等量完成比较。",
                                "actual_counts":{r["method"]:len(r["records"]) for r in results},
                                "both_reached_evaluation_budget":len(results)==2 and all(len(r["records"])==settings["evaluations"] for r in results)},
             "continuous_safety_certified":False,"global_optimum_certified":False,
             "output_directory":str(out)}
    feasible=sorted((row for result in results for row in result["records"] if _is_feasible(row)),key=lambda row:row["length"])
    unique=[]; keys=set()
    for row in feasible:
        if _point_key(row["x"]) not in keys:
            unique.append(row); keys.add(_point_key(row["x"]))
    selected=None; refined=None; attempts=[]
    for row in unique:
        if not _remaining(deadline):
            break
        path=TurnPath.from_dict(row["path"])
        print(f"[第四问] 加密检查候选 L={path.length:.6f} m。",flush=True)
        try:
            motion=evaluate_motion(path,step=validation_step,deadline=deadline,refine=True)
        except (ValueError,ArithmeticError,RuntimeError) as error:
            motion={"status":"UNRESOLVED","complete":False,"min_gap":None,"reason":str(error)}
        motion["required_gap_m"] = settings["clearance_margin"]
        if motion.get("status")=="FEASIBLE_SAMPLED" and motion["min_gap"]<settings["clearance_margin"]:
            motion["status"]="CLEARANCE_MARGIN_NOT_MET"
        attempts.append({"x":row["x"],"length":row["length"],"motion":motion})
        save_json(out/"候选加密检查.json",attempts)
        if _is_feasible(motion) and geometry_audit(path)["passed"]:
            selected=row; refined=motion; break
    plot_optimization(out,results,reference.length)
    if selected is None:
        summary.update(status="BUDGET_EXHAUSTED" if not _remaining(deadline) else "NO_REFINED_FEASIBLE_CANDIDATE",
                       optimization_feasible_count=len(feasible),refinement_attempts=len(attempts),
                       message="未取得完成加密运动检查的可行路径；不写 result4、选定路径或完成标志。",
                       elapsed_seconds=time.perf_counter()-started)
        save_json(out/"第四问结果摘要.json",summary)
        print(f"[第四问] {summary['status']}；部分计算记录已保存：{out}",flush=True)
        return 2
    path=TurnPath.from_dict(selected["path"])
    summary.update(status="FEASIBLE_SAMPLED",selected_method=selected["method"],selected_x=selected["x"],
                   arc_length_m=path.length,shortening_m=reference.length-path.length,
                   shortening_fraction=(reference.length-path.length)/reference.length,
                   shortening_confirmed=bool(path.length < reference.length-1e-8),
                   motion_check={k:v for k,v in refined.items() if k!="records"},
                   geometry_check=geometry_audit(path),source="本次实际完成的运动约束检查")
    times=np.arange(-100,101,dtype=float); states=[]
    for t in times:
        if not _remaining(deadline):
            break
        try:
            states.append(chain_state(path,float(t),speed=1.0))
        except (ValueError,ArithmeticError,RuntimeError) as error:
            summary["table_error"]=f"{type(error).__name__}: {error}"; break
    if len(states)!=201:
        if states:
            np.savez_compressed(out/"未完成_位置速度样本.npz",times=times[:len(states)],
                                positions=np.stack([s["positions"] for s in states]),
                                velocities=np.stack([s["velocities"] for s in states]))
        summary.update(status="TABLE_INCOMPLETE",completed_times=len(states),elapsed_seconds=time.perf_counter()-started)
        save_json(out/"第四问结果摘要.json",summary)
        return 2
    positions=np.stack([s["positions"] for s in states]); velocities=np.stack([s["velocities"] for s in states])
    speeds=np.linalg.norm(velocities,axis=2); coordinates=np.stack([s["s"] for s in states])
    residuals=state_residuals(path,states)
    if not residuals["passed"]:
        summary.update(status="TABLE_CONSTRAINT_FAILURE",table_residuals=residuals)
        save_json(out/"第四问结果摘要.json",summary)
        return 2
    np.savez_compressed(out/"第四问完整状态.npz",times=times,s=coordinates,positions=positions,velocities=velocities,
                        speeds=speeds,coefficients=np.stack([s["coefficients"] for s in states]),
                        tangents=np.stack([s["tangents"] for s in states]))
    workbook_check=export_q4_workbook(out/"result4.xlsx",positions,speeds)
    _csv(out/"七把手规定时刻表.csv",["t_s","handle_0based","名称","x_m","y_m","speed_m_s"],
         [[float(t),int(i),handle_label(int(i)),*positions[int(t)+100,i],speeds[int(t)+100,i]]
          for t in (-100,-50,0,50,100) for i in SELECTED])
    _csv(out/"全部把手位置速度.csv",["t_s","handle_0based","s_m","x_m","y_m","vx_m_s","vy_m_s","speed_m_s"],
         [[float(t),i,coordinates[k,i],*positions[k,i],*velocities[k,i],speeds[k,i]] for k,t in enumerate(times) for i in range(224)])
    continuity=({"status":"SKIPPED","passed":None,"reason":"--skip-validation"} if args.skip_validation
                else event_continuity_check(path,deadline,out,args.quick))
    completed_searches=(len(results)==2 and all(r["status"] in
                        ("EVALUATION_BUDGET_EXHAUSTED","PRACTICAL_TOLERANCE_REACHED") for r in results))
    comparison={"status":"SKIPPED" if args.skip_validation else "COMPLETED" if completed_searches else "BUDGET_INCOMPLETE",
                "comparison_scope":"相同预算上限和实用停止规则；分别列出粗网格最好值和已经加密确认的值，不声明全局收敛。",
                "results":[{"method":r["method"],"actual_evaluations":len(r["records"]),
                            "best_sampled_length_m":_best(r["records"])["length"] if _best(r["records"]) else None,
                            "best_refined_length_m":min((a["length"] for a in attempts
                                if _is_feasible(a["motion"]) and any(_point_key(a["x"]) == _point_key(row["x"])
                                    for row in r["records"])), default=None)} for r in results]}
    validation_status=("FAILED" if continuity.get("passed") is False else "PASSED" if continuity.get("passed") is True and comparison["status"]=="COMPLETED" else "PARTIAL")
    validation={"status":validation_status,"constraint_substitution":{"geometry":summary["geometry_check"],
                "refined_motion":summary["motion_check"],"table_residuals":residuals,"event_continuity":continuity},
                "independent_algorithm":comparison,"convergence":budget_prefix_report(results),"scope":SCOPE}
    save_json(out/"三类检验报告.json",validation)
    if validation_status=="FAILED":
        summary.update(status="CONSTRAINT_VALIDATION_FAILED",validation_status=validation_status,
                       workbook_verification=workbook_check,table_residuals=residuals,
                       message="事件连续性回代未通过；保留候选数据，不发布第五问路径或完成标志。")
        save_json(out/"第四问结果摘要.json",summary)
        return 2
    frozen={"status":"FEASIBLE_SAMPLED","path":path.to_dict(),"source":"第四问实际求得并加密检查的可行候选",
            "clearance_margin_m":settings["clearance_margin"],"optimization_improvement_relative_tolerance":settings["improvement_rel"],
            "selected_method":selected["method"],"validation_status":validation_status,"motion_check":summary["motion_check"],
            "continuous_safety_certified":False,"global_optimum_certified":False,
            "arc_coordinate_zero":"第一段圆弧起点 A；龙头速度为 1 m/s 时 s_head=t",
            "source_sha256":{name:_hash(HERE/name) for name in ("问题1代码.py","问题2代码.py","问题4代码.py")}}
    save_json(out/"选定路径.json",frozen)
    plot_path_comparison(out,reference,path); plot_key_configurations(out,path)
    summary.update(validation_status=validation_status,workbook_verification=workbook_check,
                   table_residuals=residuals,elapsed_seconds=time.perf_counter()-started,
                   notes="正式表格完整；路径为数值可行候选。完整边界到边界路径长度与圆弧总长分开记录。")
    primitive=lambda theta:.5*(theta*math.sqrt(1+theta*theta)+math.asinh(theta))
    theta_r=path.radius/path.params.b
    summary["boundary_to_boundary_length_m"]=(path.length+path.params.b*(2*primitive(theta_r)-primitive(path.theta_a)-primitive(path.theta_c)))
    save_json(out/"第四问结果摘要.json",summary)
    save_json(out/"第四问完成.json",{"status":"TABLE_AND_PATH_COMPLETE","validation_status":validation_status,
              "path_file":"选定路径.json","path_sha256":_hash(out/"选定路径.json"),
              "workbook":"result4.xlsx","workbook_sha256":_hash(out/"result4.xlsx"),
              "times":201,"handles":224,"scope":SCOPE})
    print(f"[第四问完成] L={path.length:.6f} m，检验 {validation_status}，输出：{out}",flush=True)
    return 0 if validation_status=="PASSED" else 2


if __name__=="__main__":
    raise SystemExit(main())
