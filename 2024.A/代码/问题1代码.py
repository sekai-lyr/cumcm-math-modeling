"""2024A 板凳龙第一问：运动求解、附件导出、结果图和三类检验。

安装依赖：python -m pip install numpy scipy matplotlib openpyxl
运行：python 问题1代码.py
输出默认写入项目“结果/运行记录/第一问”；--output-dir 可更改。
第二、三问直接导入本文件的公共计算函数，不要求先运行第一问。
本脚本不生成流程图或建模示意图，不修改题目附件。
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import time
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
from scipy.integrate import solve_ivp
from scipy.optimize import brentq

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
SELECTED = np.array([0, 1, 51, 101, 151, 201, 223])
KEY_TIMES = np.arange(0, 301, 60, dtype=float)


@dataclass(frozen=True)
class Parameters:
    """长度单位 m，时间单位 s；板凳 0 基编号，把手 0…223。"""

    pitch: float = 0.55
    width: float = 0.30
    hole_offset: float = 0.275
    head_speed: float = 1.0
    start_turns: float = 16.0
    root_tol: float = 1e-13

    def __post_init__(self):
        values = tuple(asdict(self).values())
        if not all(math.isfinite(v) and v > 0 for v in values):
            raise ValueError("模型参数必须是有限正数")
        if self.hole_offset >= 1.10:
            raise ValueError("孔端距离必须小于最短板凳长度的一半")

    @property
    def b(self):
        return self.pitch / (2 * math.pi)

    @property
    def theta_start(self):
        return 2 * math.pi * self.start_turns

    @property
    def lengths(self):
        return np.r_[3.41, np.full(222, 2.20)]

    @property
    def distances(self):
        return self.lengths - 2 * self.hole_offset

    @property
    def end_time(self):
        """仅指龙头沿螺线到原点的弧长时间，不保证此前全链可行。"""
        return float(self.b * primitive(self.theta_start) / self.head_speed)


DEFAULT = Parameters()


def prepare_output_dir(base):
    """每次运行独立保存，预算中断或跳过检验不会混入旧图表。"""
    directory = Path(base).resolve() / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    directory.mkdir(parents=True, exist_ok=False)
    return directory


def primitive(theta):
    """无量纲螺线弧长原函数；实际弧长须乘 b。"""
    theta = np.asarray(theta)
    return 0.5 * (theta * np.sqrt(1 + theta * theta) + np.arcsinh(theta))


def positions_from_angles(angles, params=DEFAULT):
    angles = np.asarray(angles, dtype=float)
    return params.b * angles[..., None] * np.stack(
        (np.cos(angles), np.sin(angles)), axis=-1
    )


def configuration(t, params=DEFAULT):
    """弧长反演后取外侧最近连接根，返回 224 个角度和坐标。"""
    t = float(t)
    if not math.isfinite(t) or not 0 <= t <= params.end_time:
        raise ValueError(f"时刻须在 [0, {params.end_time:g}] 内")
    if t == 0:
        head = params.theta_start
    else:
        target = params.head_speed * t
        start_arc = params.b * float(primitive(params.theta_start))
        head = brentq(
            lambda theta: start_arc - params.b * float(primitive(theta)) - target,
            0.0, params.theta_start, xtol=params.root_tol,
        )
    angles = np.empty(224)
    angles[0] = head
    b = params.b
    for i, d in enumerate(params.distances, start=1):
        radius = b * angles[i - 1]

        def residual(delta):
            return ((b * delta) ** 2
                    + 4 * radius * (radius + b * delta) * math.sin(delta / 2) ** 2
                    - d * d)

        # 在 (0, pi) 上残差严格递增，保证所取的是最近正根。
        upper = (2 * math.asin(min(1.0, d / (2 * radius)))
                 if radius > d / 2 else math.pi)
        if residual(upper) < 0:
            upper = math.pi
        if residual(upper) < 0:
            raise ArithmeticError("当前构型在最近半圈分支内无孔距根，停止而不跳圈")
        angles[i] = angles[i - 1] + brentq(
            residual, 0.0, upper, xtol=params.root_tol
        )
    return angles, positions_from_angles(angles, params)


def angular_rates_from_angles(angles, params=DEFAULT):
    """刚性孔距约束的解析微分；同时供独立 RK 积分调用。"""
    angles = np.asarray(angles, dtype=float)
    positions = positions_from_angles(angles, params)
    derivative = params.b * np.column_stack(
        (np.cos(angles) - angles * np.sin(angles),
         np.sin(angles) + angles * np.cos(angles))
    )
    chord = np.diff(positions, axis=0)
    numerator = np.einsum("ij,ij->i", chord, derivative[:-1])
    denominator = np.einsum("ij,ij->i", chord, derivative[1:])
    normalized = denominator / (
        np.linalg.norm(chord, axis=1) * np.linalg.norm(derivative[1:], axis=1)
    )
    if np.any(normalized <= 1e-9) or np.any(numerator <= 0):
        raise ArithmeticError("速度递推接近奇异或发生非前进分支")
    rates = np.empty(224)
    rates[0] = -params.head_speed / (params.b * math.hypot(1.0, angles[0]))
    rates[1:] = rates[0] * np.cumprod(numerator / denominator)
    if not np.isfinite(rates).all():
        raise ArithmeticError("速度递推产生非有限数")
    return rates, derivative, float(normalized.min())


def state(t, params=DEFAULT):
    angles, positions = configuration(t, params)
    rates, derivative, minimum = angular_rates_from_angles(angles, params)
    return {
        "t": float(t), "angles": angles, "positions": positions,
        "velocities": derivative * rates[:, None], "angular_rates": rates,
        "minimum_normalized_denominator": minimum,
    }


def board_vertices(positions, params=DEFAULT):
    positions = np.asarray(positions)
    chord = np.diff(positions, axis=0)
    direction = chord / np.linalg.norm(chord, axis=1)[:, None]
    normal = np.column_stack((-direction[:, 1], direction[:, 0]))
    center = (positions[:-1] + positions[1:]) / 2
    along = direction * params.lengths[:, None] / 2
    across = normal * params.width / 2
    return np.stack(
        (center - along - across, center + along - across,
         center + along + across, center - along + across), axis=1
    )


def constraint_checks(times, angles, positions, velocities, params=DEFAULT):
    """按量纲分别回代；这些残差不能代替独立算法或连续碰撞检验。"""
    times = np.atleast_1d(np.asarray(times, dtype=float))
    angles = np.asarray(angles, dtype=float).reshape(-1, 224)
    positions = np.asarray(positions, dtype=float).reshape(-1, 224, 2)
    velocities = np.asarray(velocities, dtype=float).reshape(-1, 224, 2)
    if not len(times) == len(angles) == len(positions) == len(velocities):
        raise ValueError("时刻与状态数组长度不一致")
    finite = all(np.isfinite(a).all() for a in (times, angles, positions, velocities))
    if not finite:
        return {"passed": False, "all_finite": False, "maximum_absolute_residuals": {}}
    chord = np.diff(positions, axis=1)
    length = np.linalg.norm(chord, axis=-1)
    tangent = np.stack(
        (np.cos(angles) - angles * np.sin(angles),
         np.sin(angles) + angles * np.cos(angles)), axis=-1
    )
    tangent /= np.linalg.norm(tangent, axis=-1)[..., None]
    residuals = {
        "spiral_position_m": float(np.max(np.linalg.norm(
            positions - positions_from_angles(angles, params), axis=-1))),
        "hole_distance_m": float(np.max(np.abs(length - params.distances))),
        "head_arc_length_m": float(np.max(np.abs(
            params.b * (primitive(params.theta_start) - primitive(angles[:, 0]))
            - params.head_speed * times))),
        "rigid_length_rate_m_per_s": float(np.max(np.abs(np.sum(
            chord * np.diff(velocities, axis=1), axis=-1) / length))),
        "transverse_speed_m_per_s": float(np.max(np.abs(
            tangent[..., 0] * velocities[..., 1] - tangent[..., 1] * velocities[..., 0]))),
        "head_speed_m_per_s": float(np.max(np.abs(
            np.linalg.norm(velocities[:, 0], axis=-1) - params.head_speed))),
    }
    ordered = bool(np.all((np.diff(angles, axis=1) > 0)
                          & (np.diff(angles, axis=1) <= math.pi)))
    return {
        "passed": bool(ordered and max(residuals.values()) <= 1e-8),
        "all_finite": True, "outward_nearest_half_turn_branch": ordered,
        "sample_count": len(times),
        "maximum_absolute_residuals": residuals,
        "distance_tolerance_m": 1e-8, "velocity_tolerance_m_per_s": 1e-8,
    }


def save_json(path, obj):
    def convert(value):
        if isinstance(value, np.ndarray):
            return convert(value.tolist())
        if isinstance(value, np.generic):
            return convert(value.item())
        if isinstance(value, dict):
            return {str(k): convert(v) for k, v in value.items()}
        if isinstance(value, (tuple, list)):
            return [convert(v) for v in value]
        if isinstance(value, Path):
            return str(value)
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("结果含非有限数；须明确记录失败状态而非输出 NaN/Infinity")
        return value

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(convert(obj), ensure_ascii=False, indent=2,
                               allow_nan=False) + "\n", encoding="utf-8")


def handle_label(i):
    if i == 0:
        return "龙头前把手"
    if i == 222:
        return "龙尾前把手"
    if i == 223:
        return "龙尾后把手"
    return f"第{i}节龙身前把手"


def _template_name(i):
    return "龙头" if i == 0 else "龙尾" if i == 222 else "龙尾（后）" if i == 223 else f"第{i}节龙身"


def _normalized(text):
    return str(text or "").replace(" ", "").replace("　", "")


def _load_template(name, destination):
    from openpyxl import load_workbook

    source = PROJECT / "题目" / "附件" / name
    if Path(destination).resolve() == source.resolve():
        raise ValueError("输出不能覆盖原始附件")
    if not source.is_file():
        raise FileNotFoundError(f"缺少题目模板：{source}")
    return load_workbook(source)


def _save_and_verify(workbook, path, expected):
    """保留模板；重读每个填入值，避免编号/转置/精度错误。"""
    from openpyxl import load_workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    labels = {s.title: (tuple(s.cell(r, 1).value for r in range(1, s.max_row + 1)),
                        tuple(s.cell(1, c).value for c in range(1, s.max_column + 1)))
              for s in workbook.worksheets}
    workbook.save(path)
    workbook.close()
    check = load_workbook(path, read_only=True, data_only=True)
    count, max_error = 0, 0.0
    try:
        for sheet, (start_row, start_col, values) in expected.items():
            ws = check[sheet]
            for ri, row in enumerate(ws.iter_rows(
                    min_row=start_row, max_row=start_row + len(values) - 1,
                    min_col=start_col, max_col=start_col + len(values[0]) - 1)):
                for ci, cell in enumerate(row):
                    if not isinstance(cell.value, (int, float)):
                        raise ArithmeticError(f"表格出现非数值：{sheet}!{cell.coordinate}")
                    max_error = max(max_error, abs(float(cell.value) - float(values[ri, ci])))
                    if cell.number_format != "0.000000":
                        raise ArithmeticError("表格显示格式未保留六位小数")
                    count += 1
            # read_only 模式下避免逐单元格反复解析大表。
            rows = ws.iter_rows(values_only=True)
            header = next(rows)
            column = [header[0]] + [r[0] for r in rows]
            if tuple(column) != labels[sheet][0] or tuple(header) != labels[sheet][1]:
                raise ArithmeticError("导出后模板标签发生变化")
        if max_error > 5e-13:
            raise ArithmeticError(f"工作簿与原始数据不一致：{max_error:g}")
    finally:
        check.close()
    return {"path": str(path), "verified_cells": count, "max_export_error": max_error}


def export_q1_workbook(path, positions, speeds):
    positions, speeds = np.asarray(positions), np.asarray(speeds)
    if positions.shape != (301, 224, 2) or speeds.shape != (301, 224):
        raise ValueError("result1 要求 301 个时刻、224 个把手")
    wb = _load_template("result1.xlsx", path)
    if wb.sheetnames != ["位置", "速度"]:
        raise ValueError("result1 模板工作表名称不匹配")
    expected = {}
    for sheet, values in (
        ("位置", positions.reshape(301, 448).T), ("速度", speeds.T)
    ):
        ws = wb[sheet]
        if (ws.max_row, ws.max_column) != (len(values) + 1, 302):
            raise ValueError("result1 模板尺寸不匹配")
        for t in range(301):
            if _normalized(ws.cell(1, t + 2).value) != f"{t}s":
                raise ValueError("result1 时间标签不匹配")
        for i in range(224):
            if sheet == "位置":
                for axis, offset in (("x", 0), ("y", 1)):
                    if _normalized(ws.cell(2 + 2 * i + offset, 1).value) != _template_name(i) + axis + "(m)":
                        raise ValueError("result1 位置把手标签不匹配")
            elif _normalized(ws.cell(2 + i, 1).value) != _template_name(i) + "(m/s)":
                raise ValueError("result1 速度把手标签不匹配")
        for r, row in enumerate(values, start=2):
            for c, value in enumerate(row, start=2):
                cell = ws.cell(r, c, float(value))
                cell.number_format = "0.000000"
        expected[sheet] = (2, 2, values)
    return _save_and_verify(wb, path, expected)


def export_q2_workbook(path, positions, speeds):
    values = np.column_stack((positions, speeds))
    if values.shape != (224, 3) or not np.isfinite(values).all():
        raise ValueError("result2 要求 224 个把手的有限坐标与速率")
    wb = _load_template("result2.xlsx", path)
    ws = wb.worksheets[0]
    if (ws.max_row, ws.max_column) != (225, 4):
        raise ValueError("result2 模板尺寸不匹配")
    if [_normalized(ws.cell(1, c).value) for c in (2, 3, 4)] != ["横坐标x(m)", "纵坐标y(m)", "速度(m/s)"]:
        raise ValueError("result2 列标签不匹配")
    for i, row in enumerate(values):
        if _normalized(ws.cell(i + 2, 1).value) != _template_name(i):
            raise ValueError("result2 把手标签不匹配")
        for c, value in enumerate(row, start=2):
            cell = ws.cell(i + 2, c, float(value))
            cell.number_format = "0.000000"
    return _save_and_verify(wb, path, {ws.title: (2, 2, values)})


def configure_plotting():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "DejaVu Sans"],
        "axes.unicode_minus": True, "font.size": 9,
        "mathtext.fontset": "stix", "mathtext.default": "it",
        "axes.formatter.use_mathtext": True,
        "axes.formatter.useoffset": False, "axes.formatter.limits": (0, 0),
        "axes.spines.top": False, "axes.spines.right": False,
        "axes.edgecolor": "#646B73", "text.color": "#222222",
        "axes.labelcolor": "#222222", "svg.fonttype": "path",
        "savefig.facecolor": "white",
    })
    return plt


def scientific_tex(value, decimals=3):
    """返回不含 $ 分隔符的数学科学计数法；编号和分类标签不调用它。"""
    value = float(value)
    if not math.isfinite(value):
        raise ValueError("图内数值标注必须有限")
    if value == 0:
        return "0"
    coefficient, exponent = f"{value:.{decimals}e}".split("e")
    coefficient = coefficient.rstrip("0").rstrip(".") if "." in coefficient else coefficient
    return rf"{coefficient}\times 10^{{{int(exponent)}}}"


def format_figure_numbers(fig):
    """线性轴采用公共 10 的幂倍率；保留对数轴和分类轴各自的定位规则。"""
    from matplotlib.ticker import LogFormatter, LogFormatterSciNotation, ScalarFormatter
    for ax in fig.axes:
        for axis in (ax.xaxis, ax.yaxis):
            formatter = axis.get_major_formatter()
            if isinstance(formatter, ScalarFormatter) and axis.get_scale() == "linear":
                formatter.set_useMathText(True)
                formatter.set_useOffset(False)
                formatter.set_scientific(True)
                formatter.set_powerlimits((0, 0))
            elif isinstance(formatter, LogFormatter) and axis.get_scale() == "log":
                axis.set_major_formatter(LogFormatterSciNotation(base=10))


def save_figure(fig, outdir, stem):
    format_figure_numbers(fig)
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    paths = []
    for extension in ("png", "svg"):
        path = outdir / f"{stem}.{extension}"
        fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.12)
        paths.append(str(path))
    return paths


def draw_chain(ax, positions, params=DEFAULT, highlight=()):
    from matplotlib.collections import PolyCollection
    polygons = board_vertices(positions, params)
    colors = ["#E5EBEF"] * len(polygons)
    edges = ["#80909C"] * len(polygons)
    for i in highlight:
        colors[int(i)] = "#F9DDBF"
        edges[int(i)] = "#B66B2C"
    ax.add_collection(PolyCollection(polygons, facecolors=colors, edgecolors=edges,
                                     linewidths=0.35))
    ax.scatter(*np.asarray(positions)[0], s=15, color="#255C87", zorder=4)
    ax.autoscale_view()
    ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel(r"$x/\mathrm{m}$")
    ax.set_ylabel(r"$y/\mathrm{m}$")
    return polygons


def run_validation(times, angles, positions, velocities, quick=False):
    constraints = constraint_checks(times, angles, positions, velocities)
    integration = solve_ivp(
        lambda _, theta: angular_rates_from_angles(theta)[0],
        (0.0, 300.0), angles[0], method="DOP853", t_eval=times,
        rtol=1e-10, atol=1e-11, max_step=2.0,
    )
    rk = {"method": "DOP853 Runge-Kutta", "rtol": 1e-10, "atol": 1e-11,
          "success": bool(integration.success), "function_calls": integration.nfev}
    position_errors = np.array([])
    if integration.success and integration.y.shape == (224, len(times)):
        position_errors = np.linalg.norm(
            positions_from_angles(integration.y.T) - positions, axis=-1
        )
        rk.update(max_position_error_m=float(position_errors.max()),
                  rms_position_error_m=float(np.sqrt(np.mean(position_errors ** 2))),
                  tolerance_m=1e-6, passed=bool(position_errors.max() <= 1e-6))
    else:
        rk.update(passed=False, message=integration.message)
    checkpoints = (0, 150, 300) if quick else tuple(KEY_TIMES.astype(int))
    records = []
    for t in checkpoints:
        runs = []
        center, analytical = positions[t], velocities[t]
        for h in (0.01, 0.005, 0.0025):
            if t == 0:
                numerical = (-3 * center + 4 * configuration(h)[1]
                             - configuration(2 * h)[1]) / (2 * h)
            elif t == 300:
                numerical = (3 * center - 4 * configuration(t - h)[1]
                             + configuration(t - 2 * h)[1]) / (2 * h)
            else:
                numerical = (configuration(t + h)[1] - configuration(t - h)[1]) / (2 * h)
            error = np.linalg.norm(numerical - analytical, axis=1)
            runs.append({"h_s": h, "max_error_m_per_s": float(error.max()),
                         "rms_error_m_per_s": float(np.sqrt(np.mean(error ** 2)))})
        records.append({"time_s": t, "runs": runs,
                        "passed": runs[-1]["max_error_m_per_s"] <= 1e-6})
    fd = {"tolerance_m_per_s": 1e-6, "records": records,
          "passed": all(r["passed"] for r in records)}
    return {
        "status": "COMPLETED",
        "constraints": constraints, "independent_runge_kutta": rk,
        "finite_difference": fd,
        "all_passed": bool(constraints["passed"] and rk["passed"] and fd["passed"]),
    }, position_errors


def make_figures(outdir, times, positions, velocities, validation, rk_errors):
    plt = configure_plotting()
    fig, axes = plt.subplots(2, 3, figsize=(10, 6.5), layout="constrained")
    vertices = np.concatenate([board_vertices(positions[int(t)]).reshape(-1, 2)
                               for t in KEY_TIMES])
    lower, upper = vertices.min(axis=0) - .7, vertices.max(axis=0) + .7
    for ax, t in zip(axes.flat, KEY_TIMES.astype(int)):
        draw_chain(ax, positions[t], highlight=(0, 222))
        ax.annotate("", xy=positions[t, 0] + velocities[t, 0] * .8,
                    xytext=positions[t, 0],
                    arrowprops={"arrowstyle": "->", "color": "#255C87"})
        ax.set(xlim=(lower[0], upper[0]), ylim=(lower[1], upper[1]),
               title=rf"$t={scientific_tex(t, 2)}\,\mathrm{{s}}$")
    fig.suptitle("第一问：盘入过程关键时刻队形")
    save_figure(fig, outdir, "R1_盘入关键时刻队形")
    plt.close(fig)

    fig, axes = plt.subplots(2, 1, figsize=(9, 6), layout="constrained")
    speeds = np.linalg.norm(velocities, axis=-1)
    colors = ["#555555", "#255C87", "#6E9BBE", "#A5C1D7", "#B46628", "#D49762", "#E6BC98"]
    for n, i in enumerate(SELECTED):
        style = ("-", "--", "-.", ":")[n % 4]
        for ax, values in zip(axes, (speeds[:, i], speeds[:, i] - speeds[:, 0])):
            ax.plot(times, values, color=colors[n], linestyle=style,
                    label=handle_label(int(i)), linewidth=1.2)
    axes[0].set(ylabel=r"速率 $u_i/(\mathrm{m\,s^{-1}})$", title="第一问：代表把手速率")
    axes[0].legend(ncol=3, fontsize=8)
    axes[1].set(xlabel=r"时间 $t/\mathrm{s}$",
                ylabel=r"相对龙头速率差 $(u_i-u_0)/(\mathrm{m\,s^{-1}})$")
    axes[1].ticklabel_format(axis="y", style="sci", scilimits=(-2, 2))
    for ax in axes:
        ax.grid(alpha=.2)
    save_figure(fig, outdir, "R2_代表把手速率")
    plt.close(fig)

    if validation["status"] == "COMPLETED":
        fig, axes = plt.subplots(1, 2, figsize=(10, 3.8), layout="constrained")
        zero_fd_times = []
        for n, record in enumerate(validation["finite_difference"]["records"]):
            errors = np.asarray([r["max_error_m_per_s"] for r in record["runs"]])
            if np.any(errors == 0):
                zero_fd_times.append(record["time_s"])
            axes[0].loglog(
                [r["h_s"] for r in record["runs"]],
                np.ma.masked_less_equal(errors, 0),
                marker=("o", "s", "^")[n % 3], linestyle=("-", "--")[n % 2],
                color=colors[n + 1], markersize=3,
                label=rf'$t={scientific_tex(record["time_s"], 2)}\,\mathrm{{s}}$',
            )
        axes[0].axhline(1e-6, color="#555555", linestyle="--", label="检验阈值")
        axes[0].set(xlabel=r"差分步长 $h/\mathrm{s}$",
                    ylabel=r"$\max_i\|\mathbf{v}_i^{\mathrm{FD}}-\mathbf{v}_i\|/(\mathrm{m\,s^{-1}})$",
                    title="位置差分与解析速度")
        if zero_fd_times:
            axes[0].text(.03, .97, "部分计算误差为 $0$，对数轴中不绘制零值",
                         transform=axes[0].transAxes, va="top", fontsize=7)
        axes[0].legend(fontsize=7, ncol=2)
        if rk_errors.size:
            maxima = rk_errors.max(axis=1)
            positive = maxima > 0
            axes[1].semilogy(times[positive], maxima[positive], color="#255C87")
            axes[1].axhline(1e-6, color="#555555", linestyle="--")
            note = "最大差 $" + scientific_tex(maxima.max()) + r"\,\mathrm{m}$"
            if not np.any(positive):
                note += "\n全部计算误差为 $0$"
            axes[1].text(.03, .94, note,
                         transform=axes[1].transAxes, va="top", fontsize=8)
        else:
            axes[1].text(.5, .5, "独立积分未完成", ha="center", transform=axes[1].transAxes)
        axes[1].set(xlabel=r"时间 $t/\mathrm{s}$",
                    ylabel=r"$\max_i\|\mathbf{P}_i^{\mathrm{RK}}-\mathbf{P}_i\|/\mathrm{m}$",
                    title="独立积分与逐时刻求根")
        save_figure(fig, outdir, "V1a_第一问数值检验")
        plt.close(fig)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=HERE.parent / "结果" / "运行记录" / "第一问",
                        help="输出基目录；每次运行在其下新建时间戳目录")
    parser.add_argument("--quick", action="store_true", help="完整求解；差分检验取三个代表时刻")
    parser.add_argument("--skip-validation", action="store_true", help="显式跳过独立检验")
    parser.add_argument("--no-plots", action="store_true", help="只输出数值与表格")
    args = parser.parse_args(argv)
    outdir = prepare_output_dir(args.output_dir)
    started = time.perf_counter()
    print("第一问：计算 0～300 s 全部 224 个把手。", flush=True)
    times = np.arange(301, dtype=float)
    states = []
    for t in times:
        states.append(state(t))
        if int(t) % 60 == 0:
            print(f"已计算 t={t:.0f} s", flush=True)
    angles = np.stack([s["angles"] for s in states])
    positions = np.stack([s["positions"] for s in states])
    velocities = np.stack([s["velocities"] for s in states])
    speeds = np.linalg.norm(velocities, axis=-1)
    np.savez_compressed(outdir / "第一问完整状态.npz", times=times, angles=angles,
                        positions=positions, velocities=velocities, speeds=speeds)
    with (outdir / "第一问论文表.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间(s)", "把手编号(0基)", "把手", "x(m)", "y(m)", "速率(m/s)"])
        for t in KEY_TIMES.astype(int):
            for i in SELECTED:
                writer.writerow([t, int(i), handle_label(int(i)),
                                 f"{positions[t,i,0]:.6f}", f"{positions[t,i,1]:.6f}",
                                 f"{speeds[t,i]:.6f}"])
    print("第一问：独立积分、约束回代与速度差分检验。", flush=True)
    validation, rk_errors = (
        ({"status": "NOT_RUN", "reason": "--skip-validation", "all_passed": None}, np.array([]))
        if args.skip_validation else run_validation(times, angles, positions, velocities, args.quick)
    )
    save_json(outdir / "第一问检验.json", validation)
    if rk_errors.size:
        np.savez_compressed(outdir / "第一问积分检验作图数据.npz",
                            times=times, position_errors=rk_errors)
    export = export_q1_workbook(outdir / "result1.xlsx", positions, speeds)
    if not args.no_plots:
        print("第一问：输出结果图和检验图。", flush=True)
        make_figures(outdir, times, positions, velocities, validation, rk_errors)
    summary = {
        "status": "NUMERICAL_SOLUTION_COMPLETE", "validation_passed": validation["all_passed"],
        "parameters": asdict(DEFAULT), "time_range_s": [0, 300],
        "time_count": 301, "handle_count": 224, "selected_handles": SELECTED,
        "maximum_sampled_speed_m_per_s": float(speeds.max()),
        "workbook": export, "elapsed_seconds": time.perf_counter() - started,
        "source_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "template_sha256": hashlib.sha256((PROJECT / "题目/附件/result1.xlsx").read_bytes()).hexdigest(),
        "scope": "规定的301个输出时刻；第一问不据此认证连续时间无碰撞",
    }
    save_json(outdir / "第一问结果摘要.json", summary)
    print(f"第一问完成：{outdir}；检验通过={validation['all_passed']}", flush=True)
    return 1 if validation["all_passed"] is False else 0


if __name__ == "__main__":
    raise SystemExit(main())
