"""2022 年高教社杯全国大学生数学建模竞赛 A 题——问题 3。

模型采用四个独立的二阶自由度：
    1. 浮子垂荡位移 x_f；
    2. 振子垂荡位移 x_z；
    3. 浮子纵摇角 theta_f；
    4. 振子纵摇角 theta_z。

垂荡方向和纵摇方向均使用相同的 PTO 相对运动结构：
    - 直线/旋转阻尼分别与相对速度/相对角速度成正比；
    - 线性/扭转弹簧分别与相对位移/相对角位移成正比。

程序从附件 3、附件 4 读取参数，将运动方程写成状态空间形式，并使用增广矩阵
指数法从静水平衡初值推进前 40 个波浪周期，以 0.2 s 为间隔将结果写入
data/result3.xlsx。

主结果之外，程序还输出数值算法交叉验证、方程回代与能量守恒、关键参数的
±10% 灵敏度分析，以及安装高度 a 扫描、G0/G1 重力模型对比和小角度适用性检验。
"""

from __future__ import annotations

import math
from copy import copy
from dataclasses import dataclass, replace
from pathlib import Path
from shutil import copyfile
from typing import Any

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border
from scipy.integrate import solve_ivp
from scipy.linalg import expm


# 题目给定的第三问设置
WAVE_FREQUENCY = 1.7152
LINEAR_DAMPING = 10_000.0       # N·s/m
ROTARY_DAMPING = 1_000.0        # N·m·s
OUTPUT_STEP = 0.2               # s
ENERGY_CHECK_STEP = 0.01        # s
PAPER_TIMES = (10.0, 20.0, 40.0, 60.0, 100.0)
NUMERICAL_RELATIVE_TOLERANCE = 1e-6
ENERGY_RELATIVE_TOLERANCE = 1e-4
SENSITIVITY_PERTURBATION = 0.1
INSTALL_HEIGHT_SAMPLES = (0.0, 0.125, 0.25, 0.5)  # m；探索性扫描
SMALL_ANGLE_LIMIT_DEG = 5.0       # 小角度适用性审核阈值
GRAVITY_MODELS = ("G0", "G1")

# 最终修正版模型参数
# 浮子按均匀薄壳模型绕隔层中心铰轴 O 计算得到的纵摇转动惯量。
FLOAT_PITCH_INERTIA = 18_044.4204  # kg·m²
INSTALL_HEIGHT_A = 0.0              # m；基准等效工况，文档规定的 a=0


@dataclass(frozen=True)
class Parameters:
    """第三问所需的物理参数。"""

    float_mass: float
    oscillator_mass: float
    added_mass: float
    added_pitch_inertia: float
    heave_damping: float
    pitch_damping: float
    wave_force: float
    wave_moment: float
    wave_frequency: float
    heave_spring: float
    pitch_spring: float
    hydrostatic_heave_stiffness: float
    hydrostatic_pitch_stiffness: float
    oscillator_pitch_inertia: float
    float_radius: float = 1.0
    cylinder_height: float = 3.0
    cone_height: float = 0.8
    oscillator_radius: float = 0.5
    oscillator_height: float = 0.5
    water_density: float = 1025.0
    gravity: float = 9.8
    spring_natural_length: float = 0.5
    install_height_a: float = INSTALL_HEIGHT_A
    gravity_model: str = "G1"
    linear_damping: float = LINEAR_DAMPING
    rotary_damping: float = ROTARY_DAMPING
    float_pitch_inertia: float = FLOAT_PITCH_INERTIA

    @property
    def effective_float_mass(self) -> float:
        return self.float_mass + self.added_mass

    @property
    def effective_float_pitch_inertia(self) -> float:
        return self.float_pitch_inertia + self.added_pitch_inertia

    @property
    def equilibrium_spring_length(self) -> float:
        """基准图示拓扑下，静平衡时的弹簧长度。"""
        return self.spring_natural_length - (
            self.oscillator_mass * self.gravity / self.heave_spring
        )

    @property
    def oscillator_orbit_radius(self) -> float:
        """振子质心到隔层中心铰轴 O 的距离 r_0(a)。"""
        return (
            self.install_height_a
            + self.equilibrium_spring_length
            + self.oscillator_height / 2.0
        )

    @property
    def effective_oscillator_pitch_inertia(self) -> float:
        """振子关于铰轴 O 的转动惯量 J_2O(a)。"""
        return self.oscillator_pitch_inertia + self.oscillator_mass * (
            self.oscillator_orbit_radius**2
        )

    @property
    def gravity_pitch_stiffness(self) -> float:
        """G1 相对重力修正 K_g(a)=m_2*g*r_0(a)。"""
        return self.oscillator_mass * self.gravity * self.oscillator_orbit_radius

    @property
    def period(self) -> float:
        return 2.0 * math.pi / self.wave_frequency


def project_root() -> Path:
    """返回项目根目录：本文件位于 项目根/代码/第三问代码/。"""
    return Path(__file__).resolve().parents[2]


def as_float(value: Any, name: str) -> float:
    """将 Excel 单元格转换为浮点数。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"附件中缺少参数：{name}")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"参数 {name} 不是数值：{value!r}") from exc


def find_value_by_label(
    worksheet: Any, keywords: tuple[str, ...], name: str
) -> float:
    """在附件 4 的前两列中按关键词查找参数，兼容标签中的空格。"""
    normalized_keywords = tuple(keyword.replace(" ", "") for keyword in keywords)
    for row in worksheet.iter_rows(min_col=1, max_col=2, values_only=True):
        label = "" if row[0] is None else str(row[0]).replace(" ", "")
        if all(keyword in label for keyword in normalized_keywords):
            return as_float(row[1], name)
    raise ValueError(f"附件 4 中找不到参数：{name}（关键词：{keywords}）")


def find_wave_row(worksheet: Any, frequency: float) -> tuple[Any, ...]:
    """在附件 3 中查找指定波浪频率对应的参数行。"""
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        try:
            if math.isclose(float(row[1]), frequency, abs_tol=1e-8):
                return tuple(row)
        except (IndexError, TypeError, ValueError):
            continue
    raise ValueError(f"附件 3 中找不到波浪频率 {frequency:g} 的参数行")


def read_parameters() -> Parameters:
    """读取第三问的附件参数。"""
    data_dir = project_root() / "data"
    attachment3 = data_dir / "附件3.xlsx"
    attachment4 = data_dir / "附件4.xlsx"
    if not attachment3.exists() or not attachment4.exists():
        raise FileNotFoundError(
            f"找不到输入附件，请确认文件存在：{attachment3}、{attachment4}"
        )

    workbook3 = load_workbook(attachment3, data_only=True, read_only=True)
    try:
        wave_row = find_wave_row(workbook3.active, WAVE_FREQUENCY)
    finally:
        workbook3.close()

    workbook4 = load_workbook(attachment4, data_only=True, read_only=True)
    try:
        worksheet4 = workbook4.active
        float_mass = find_value_by_label(
            worksheet4, ("浮子质量",), "浮子质量"
        )
        oscillator_mass = find_value_by_label(
            worksheet4, ("振子质量",), "振子质量"
        )
        float_radius = find_value_by_label(
            worksheet4, ("浮子底半径",), "浮子底半径"
        )
        oscillator_radius = find_value_by_label(
            worksheet4, ("振子半径",), "振子半径"
        )
        oscillator_height = find_value_by_label(
            worksheet4, ("振子高度",), "振子高度"
        )
        water_density = find_value_by_label(
            worksheet4, ("海水", "密度"), "海水密度"
        )
        gravity = find_value_by_label(
            worksheet4, ("重力加速度",), "重力加速度"
        )
        heave_spring = find_value_by_label(
            worksheet4, ("弹簧刚度",), "弹簧刚度"
        )
        pitch_spring = find_value_by_label(
            worksheet4, ("扭转弹簧刚度",), "扭转弹簧刚度"
        )
        pitch_stiffness = find_value_by_label(
            worksheet4, ("静水恢复力矩系数",), "静水恢复力矩系数"
        )
        cylinder_height = find_value_by_label(
            worksheet4, ("浮子", "圆柱部分高度"), "浮子圆柱部分高度"
        )
        cone_height = find_value_by_label(
            worksheet4, ("浮子", "圆锥部分高度"), "浮子圆锥部分高度"
        )
    finally:
        workbook4.close()

    # 附件 3 的列顺序：频率、附加质量、附加转动惯量、
    # 垂荡兴波阻尼、纵摇兴波阻尼、激励力、激励力矩。
    oscillator_pitch_inertia = (
        oscillator_mass * (3.0 * oscillator_radius**2 + oscillator_height**2)
        / 12.0
    )
    hydrostatic_heave_stiffness = water_density * gravity * math.pi * float_radius**2
    return Parameters(
        float_mass=float_mass,
        oscillator_mass=oscillator_mass,
        added_mass=as_float(wave_row[2], "附加质量"),
        added_pitch_inertia=as_float(wave_row[3], "附加转动惯量"),
        heave_damping=as_float(wave_row[4], "垂荡兴波阻尼系数"),
        pitch_damping=as_float(wave_row[5], "纵摇兴波阻尼系数"),
        wave_force=as_float(wave_row[6], "波浪激励力幅值"),
        wave_moment=as_float(wave_row[7], "波浪激励力矩幅值"),
        wave_frequency=as_float(wave_row[1], "波浪频率"),
        heave_spring=heave_spring,
        pitch_spring=pitch_spring,
        hydrostatic_heave_stiffness=hydrostatic_heave_stiffness,
        hydrostatic_pitch_stiffness=pitch_stiffness,
        oscillator_pitch_inertia=oscillator_pitch_inertia,
        float_radius=float_radius,
        cylinder_height=cylinder_height,
        cone_height=cone_height,
        oscillator_radius=oscillator_radius,
        oscillator_height=oscillator_height,
        water_density=water_density,
        gravity=gravity,
        linear_damping=LINEAR_DAMPING,
        rotary_damping=ROTARY_DAMPING,
        float_pitch_inertia=FLOAT_PITCH_INERTIA,
    )


def gravity_pitch_terms(params: Parameters) -> tuple[float, float]:
    """返回浮子、振子纵摇方程中的重力刚度项。"""
    model = params.gravity_model.upper()
    if model == "G0":
        return 0.0, 0.0
    if model == "G1":
        return params.gravity_pitch_stiffness, params.gravity_pitch_stiffness
    raise ValueError(
        f"不支持的重力解释模型 {params.gravity_model!r}；"
        f"可选模型：{GRAVITY_MODELS}"
    )


def state_space_matrices(
    params: Parameters,
) -> tuple[np.ndarray, np.ndarray]:
    """建立状态空间矩阵 A、输入向量 b，使 x_dot=A*x+b*cos(omega*t)。"""
    m_float = params.effective_float_mass
    m_oscillator = params.oscillator_mass
    j_float = params.effective_float_pitch_inertia
    j_oscillator = params.effective_oscillator_pitch_inertia
    float_gravity_stiffness, oscillator_gravity_stiffness = gravity_pitch_terms(
        params
    )

    matrix = np.zeros((8, 8), dtype=float)
    input_vector = np.zeros(8, dtype=float)

    # 垂荡状态：[x_f, v_f, x_z, v_z]
    matrix[0, 1] = 1.0
    matrix[1, 0] = (
        -params.heave_spring - params.hydrostatic_heave_stiffness
    ) / m_float
    matrix[1, 1] = -(params.heave_damping + params.linear_damping) / m_float
    matrix[1, 2] = params.heave_spring / m_float
    matrix[1, 3] = params.linear_damping / m_float
    input_vector[1] = params.wave_force / m_float

    matrix[2, 3] = 1.0
    matrix[3, 0] = params.heave_spring / m_oscillator
    matrix[3, 1] = params.linear_damping / m_oscillator
    matrix[3, 2] = -params.heave_spring / m_oscillator
    matrix[3, 3] = -params.linear_damping / m_oscillator

    # 纵摇状态：[theta_f, omega_f, theta_z, omega_z]
    matrix[4, 5] = 1.0
    matrix[5, 4] = (
        -params.pitch_spring
        - params.hydrostatic_pitch_stiffness
        - float_gravity_stiffness
    ) / j_float
    matrix[5, 5] = -(params.pitch_damping + params.rotary_damping) / j_float
    matrix[5, 6] = params.pitch_spring / j_float
    matrix[5, 7] = params.rotary_damping / j_float
    input_vector[5] = params.wave_moment / j_float

    matrix[6, 7] = 1.0
    matrix[7, 4] = params.pitch_spring / j_oscillator
    matrix[7, 5] = params.rotary_damping / j_oscillator
    matrix[7, 6] = (
        oscillator_gravity_stiffness - params.pitch_spring
    ) / j_oscillator
    matrix[7, 7] = -params.rotary_damping / j_oscillator

    return matrix, input_vector


def augmented_transition_matrix(params: Parameters, step: float) -> np.ndarray:
    """构造增广矩阵指数法的单步转移矩阵。

    增广状态为 [x, cos(omega*t), sin(omega*t)]。由于正弦激励本身也满足
    一阶线性微分方程，增广后可用一个矩阵指数精确推进每个固定时间步。
    """
    if step <= 0.0:
        raise ValueError("矩阵指数法的时间步长必须为正数")

    matrix, input_vector = state_space_matrices(params)
    augmented = np.zeros((10, 10), dtype=float)
    augmented[:8, :8] = matrix
    augmented[:8, 8] = input_vector
    augmented[8, 9] = -params.wave_frequency
    augmented[9, 8] = params.wave_frequency
    return expm(augmented * step)


def time_grid(params: Parameters, step: float = OUTPUT_STEP) -> np.ndarray:
    """生成不超过前 40 个波浪周期、间隔为 step 的输出网格。"""
    if step <= 0.0:
        raise ValueError("输出时间间隔必须为正数")
    end_time = 40.0 * params.period
    number_of_steps = int(math.floor(end_time / step + 1e-12))
    return np.arange(number_of_steps + 1, dtype=float) * step


def solve_model(
    params: Parameters, step: float = OUTPUT_STEP
) -> tuple[np.ndarray, np.ndarray]:
    """用增广矩阵指数法从静水平衡位置求解第三问模型。"""
    time = time_grid(params, step)
    step = float(time[1] - time[0])
    transition = augmented_transition_matrix(params, step)

    # 初始时刻静止且处于平衡位置，cos(0)=1、sin(0)=0。
    augmented_state = np.zeros(10, dtype=float)
    augmented_state[8] = 1.0
    states = np.empty((len(time), 8), dtype=float)
    states[0] = augmented_state[:8]
    for index in range(1, len(time)):
        augmented_state = transition @ augmented_state
        states[index] = augmented_state[:8]
    return time, states


def solve_model_ivp(
    params: Parameters, output_step: float = OUTPUT_STEP
) -> tuple[np.ndarray, np.ndarray]:
    """用 DOP853 独立求解同一组八阶一阶状态方程。"""
    time = time_grid(params, output_step)
    matrix, input_vector = state_space_matrices(params)

    def rhs(current_time: float, state: np.ndarray) -> np.ndarray:
        return matrix @ state + input_vector * math.cos(
            params.wave_frequency * current_time
        )

    solution = solve_ivp(
        rhs,
        (float(time[0]), float(time[-1])),
        np.zeros(8, dtype=float),
        method="DOP853",
        t_eval=time,
        rtol=1e-13,
        atol=1e-15,
    )
    if not solution.success:
        raise RuntimeError(f"DOP853 求解失败：{solution.message}")
    return time, solution.y.T


def state_error(
    reference: np.ndarray, candidate: np.ndarray
) -> tuple[float, float]:
    """返回最大绝对误差和按各状态幅值归一化的最大误差。"""
    difference = np.abs(reference - candidate)
    scale = np.maximum(np.max(np.abs(reference), axis=0), 1e-12)
    return float(np.max(difference)), float(np.max(difference / scale))


def finite_difference_derivative(
    values: np.ndarray, time: np.ndarray
) -> np.ndarray:
    """计算时间导数；均匀细网格内部使用七点中心差分。"""
    derivative = np.gradient(values, time, axis=0, edge_order=2)
    if len(time) >= 7:
        step = float(time[1] - time[0])
        if np.allclose(np.diff(time), step, rtol=0.0, atol=1e-12):
            derivative[3:-3] = (
                -values[:-6]
                + 9.0 * values[1:-5]
                - 45.0 * values[2:-4]
                + 45.0 * values[4:-2]
                - 9.0 * values[5:-1]
                + values[6:]
            ) / (60.0 * step)
    return derivative


def common_grid_states(
    coarse_time: np.ndarray,
    fine_time: np.ndarray,
    fine_state: np.ndarray,
) -> np.ndarray:
    """取细网格在粗网格共同时间点上的状态。"""
    indices = np.searchsorted(fine_time, coarse_time)
    if np.any(indices >= len(fine_time)) or not np.allclose(
        fine_time[indices], coarse_time, rtol=0.0, atol=1e-10
    ):
        raise ValueError("两个时间网格没有完全对齐的共同时间点")
    return fine_state[indices]


def run_numerical_cross_validation(
    params: Parameters,
) -> tuple[bool, list[tuple[str, float, str]]]:
    """交叉比较矩阵指数法、DOP853 以及不同输出步长的结果。"""
    results: dict[float, tuple[np.ndarray, np.ndarray]] = {}
    report_rows: list[tuple[str, float, str]] = []
    passed = True
    print("\n=== 检验一：数值算法交叉验证 ===")
    for step in (0.2, 0.1, 0.05):
        matrix_time, matrix_state = solve_model(params, step=step)
        _, ivp_state = solve_model_ivp(params, output_step=step)
        _, relative_error = state_error(matrix_state, ivp_state)
        results[step] = (matrix_time, matrix_state)
        case_passed = relative_error < NUMERICAL_RELATIVE_TOLERANCE
        passed = passed and case_passed
        status = "通过" if case_passed else "需关注"
        report_rows.append((f"矩阵指数法 vs DOP853，{step:.2f} s", relative_error, status))
        print(
            f"步长 {step:.2f} s：矩阵指数法 vs DOP853 最大相对误差="
            f"{relative_error:.3e}，{status}"
        )

    base_time, base_state = results[0.2]
    for step, label in ((0.1, "0.2/0.1"), (0.05, "0.2/0.05")):
        current_time, current_state = results[step]
        common_state = common_grid_states(base_time, current_time, current_state)
        _, relative_error = state_error(base_state, common_state)
        case_passed = relative_error < NUMERICAL_RELATIVE_TOLERANCE
        passed = passed and case_passed
        status = "通过" if case_passed else "需关注"
        report_rows.append((f"共同时间点 {label} s", relative_error, status))
        print(
            f"共同时间点 {label} s 最大相对差异={relative_error:.3e}，"
            f"{status}"
        )
    status = "通过" if passed else "需关注"
    print(
        f"数值交叉验证总判定：{status}"
        f"（判据：最大相对误差 < {NUMERICAL_RELATIVE_TOLERANCE:.0e}）"
    )
    return passed, report_rows


def equation_residual_summary(
    params: Parameters, time: np.ndarray, state: np.ndarray
) -> tuple[float, float]:
    """用有限差分加速度计算四条动力学方程的归一化回代残差。"""
    derivative = finite_difference_derivative(state, time)
    x_float, v_float, x_oscillator, v_oscillator = state.T[:4]
    theta_float, omega_float, theta_oscillator, omega_oscillator = state.T[4:]
    acceleration_float = derivative[:, 1]
    acceleration_oscillator = derivative[:, 3]
    angular_acceleration_float = derivative[:, 5]
    angular_acceleration_oscillator = derivative[:, 7]
    relative_velocity = v_float - v_oscillator
    relative_angular_velocity = omega_float - omega_oscillator
    relative_displacement = x_float - x_oscillator
    relative_angle = theta_float - theta_oscillator
    float_gravity_stiffness, oscillator_gravity_stiffness = gravity_pitch_terms(
        params
    )
    wave_factor = np.cos(params.wave_frequency * time)
    heave_float_force = (
        params.wave_force * wave_factor
        - params.heave_damping * v_float
        - params.linear_damping * relative_velocity
        - params.hydrostatic_heave_stiffness * x_float
        - params.heave_spring * relative_displacement
    )
    heave_oscillator_force = (
        params.heave_spring * relative_displacement
        + params.linear_damping * relative_velocity
    )
    pitch_float_moment = (
        params.wave_moment * wave_factor
        - params.pitch_damping * omega_float
        - params.rotary_damping * relative_angular_velocity
        - (params.hydrostatic_pitch_stiffness + float_gravity_stiffness)
        * theta_float
        - params.pitch_spring * relative_angle
    )
    pitch_oscillator_moment = (
        params.pitch_spring * relative_angle
        + params.rotary_damping * relative_angular_velocity
        + oscillator_gravity_stiffness * theta_oscillator
    )

    residual = np.column_stack(
        (
            params.effective_float_mass * acceleration_float
            + (params.heave_damping + params.linear_damping) * v_float
            + (params.heave_spring + params.hydrostatic_heave_stiffness) * x_float
            - params.heave_spring * x_oscillator
            - params.linear_damping * v_oscillator
            - params.wave_force * wave_factor,
            params.oscillator_mass * acceleration_oscillator
            - params.heave_spring * relative_displacement
            - params.linear_damping * relative_velocity,
            params.effective_float_pitch_inertia * angular_acceleration_float
            + (params.pitch_damping + params.rotary_damping) * omega_float
            + (
                params.pitch_spring
                + params.hydrostatic_pitch_stiffness
                + float_gravity_stiffness
            )
            * theta_float
            - params.pitch_spring * theta_oscillator
            - params.rotary_damping * omega_oscillator
            - params.wave_moment * wave_factor,
            params.effective_oscillator_pitch_inertia
            * angular_acceleration_oscillator
            - params.pitch_spring * relative_angle
            - params.rotary_damping * relative_angular_velocity
            - oscillator_gravity_stiffness * theta_oscillator,
        )
    )
    interior = slice(3, -3)
    physical_terms = np.column_stack(
        (
            heave_float_force,
            heave_oscillator_force,
            pitch_float_moment,
            pitch_oscillator_moment,
            params.effective_float_mass * acceleration_float,
            params.oscillator_mass * acceleration_oscillator,
            params.effective_float_pitch_inertia * angular_acceleration_float,
            params.effective_oscillator_pitch_inertia
            * angular_acceleration_oscillator,
        )
    )
    scale = np.maximum(np.max(np.abs(physical_terms[interior]), axis=0), 1.0)
    scale = np.maximum(scale[:4], scale[4:])
    normalized = np.abs(residual[interior]) / scale
    return float(np.max(normalized)), float(np.sqrt(np.mean(normalized**2)))


def energy_constraint_summary(
    params: Parameters, time: np.ndarray, state: np.ndarray
) -> tuple[float, float, float, float, bool, float]:
    """检查能量收支、两类 PTO 非负耗散和浮子水线范围。"""
    x_float, v_float, x_oscillator, v_oscillator = state.T[:4]
    theta_float, omega_float, theta_oscillator, omega_oscillator = state.T[4:]
    relative_velocity = v_float - v_oscillator
    relative_angular_velocity = omega_float - omega_oscillator
    float_gravity_stiffness, oscillator_gravity_stiffness = gravity_pitch_terms(
        params
    )
    energy = (
        0.5 * params.effective_float_mass * v_float**2
        + 0.5 * params.oscillator_mass * v_oscillator**2
        + 0.5 * params.effective_float_pitch_inertia * omega_float**2
        + 0.5 * params.effective_oscillator_pitch_inertia * omega_oscillator**2
        + 0.5 * params.hydrostatic_heave_stiffness * x_float**2
        + 0.5 * params.hydrostatic_pitch_stiffness * theta_float**2
        + 0.5 * params.heave_spring * (x_float - x_oscillator) ** 2
        + 0.5 * params.pitch_spring * (theta_float - theta_oscillator) ** 2
        + 0.5 * float_gravity_stiffness * theta_float**2
        - 0.5 * oscillator_gravity_stiffness * theta_oscillator**2
    )
    wave_power = (
        params.wave_force * wave_factor(time, params) * v_float
        + params.wave_moment * wave_factor(time, params) * omega_float
    )
    radiation_power = params.heave_damping * v_float**2 + params.pitch_damping * omega_float**2
    linear_pto_power = params.linear_damping * relative_velocity**2
    rotary_pto_power = params.rotary_damping * relative_angular_velocity**2
    net_power = wave_power - radiation_power - linear_pto_power - rotary_pto_power
    work = np.zeros_like(time)
    work[1:] = np.cumsum(
        0.5 * (net_power[1:] + net_power[:-1]) * np.diff(time)
    )
    integrated_residual = energy - energy[0] - work
    energy_scale = max(float(np.max(np.abs(energy))), 1.0)
    energy_relative = float(np.max(np.abs(integrated_residual)) / energy_scale)

    power_derivative_residual = (
        finite_difference_derivative(energy, time) - net_power
    )
    interior = slice(3, -3)
    power_scale = max(float(np.max(np.abs(net_power[interior]))), 1.0)
    differential_relative = float(
        np.max(np.abs(power_derivative_residual[interior])) / power_scale
    )
    lower_bound, upper_bound = waterline_bounds(params)
    maximum_displacement = float(np.max(np.abs(x_float)))
    waterline_passed = bool(
        np.max(x_float) < upper_bound and np.min(x_float) > lower_bound
    )
    return (
        energy_relative,
        float(np.min(linear_pto_power)),
        float(np.min(rotary_pto_power)),
        maximum_displacement,
        waterline_passed,
        differential_relative,
    )


def wave_factor(time: np.ndarray, params: Parameters) -> np.ndarray:
    """返回波浪激励的余弦时间因子。"""
    return np.cos(params.wave_frequency * time)


def waterline_bounds(params: Parameters) -> tuple[float, float]:
    """返回静水平衡圆柱段对应的浮子位移允许区间。"""
    waterline_area = math.pi * params.float_radius**2
    total_displaced_volume = (
        params.float_mass + params.oscillator_mass
    ) / params.water_density
    cone_volume = waterline_area * params.cone_height / 3.0
    submerged_cylinder_height = (
        total_displaced_volume - cone_volume
    ) / waterline_area
    return (
        submerged_cylinder_height - params.cylinder_height,
        submerged_cylinder_height,
    )


def run_equation_and_energy_check(
    params: Parameters,
) -> tuple[bool, dict[str, float | bool]]:
    """在 0.01 s 网格上执行方程回代、能量和物理约束检验。"""
    time, state = solve_model(params, step=ENERGY_CHECK_STEP)
    equation_max, equation_rms = equation_residual_summary(params, time, state)
    (
        energy_relative,
        linear_pto_min,
        rotary_pto_min,
        maximum_displacement,
        waterline_passed,
        differential_relative,
    ) = energy_constraint_summary(params, time, state)
    lower_bound, upper_bound = waterline_bounds(params)
    passed = (
        energy_relative < ENERGY_RELATIVE_TOLERANCE
        and equation_max < ENERGY_RELATIVE_TOLERANCE
        and equation_rms < ENERGY_RELATIVE_TOLERANCE
        and differential_relative < ENERGY_RELATIVE_TOLERANCE
        and linear_pto_min >= -1e-10
        and rotary_pto_min >= -1e-10
        and bool(waterline_passed)
    )
    print("\n=== 检验二：方程回代与能量守恒 ===")
    print(
        f"动力学方程归一化残差：最大={equation_max:.3e}，"
        f"RMS={equation_rms:.3e}"
    )
    print(
        f"能量相对残差={energy_relative:.3e}，"
        f"dE/dt 差分残差={differential_relative:.3e}，"
        f"判据 < {ENERGY_RELATIVE_TOLERANCE:.0e}"
    )
    print(
        f"直线 PTO 最小耗散功率={linear_pto_min:.3e} W，"
        f"旋转 PTO 最小耗散功率={rotary_pto_min:.3e} W"
    )
    print(
        f"浮子最大垂荡位移={maximum_displacement:.6f} m，"
        f"水线允许区间=({lower_bound:.6f}, {upper_bound:.6f}) m，"
        f"水线约束={'通过' if waterline_passed else '不通过'}"
    )
    print(f"方程/能量检验总判定：{'通过' if passed else '需关注'}")
    report = {
        "equation_max": equation_max,
        "equation_rms": equation_rms,
        "energy_relative": energy_relative,
        "differential_relative": differential_relative,
        "linear_pto_min": linear_pto_min,
        "rotary_pto_min": rotary_pto_min,
        "maximum_displacement": maximum_displacement,
        "waterline_lower": lower_bound,
        "waterline_upper": upper_bound,
        "waterline_passed": waterline_passed,
    }
    return passed, report


# 波浪频率不放进普通 ±10% 单因素灵敏度：附件 3 中的附加质量、附加转动惯量、
# 兴波阻尼、激励力和激励力矩都随频率变化，单独改变频率会破坏参数的一致性。
SENSITIVITY_PARAMETERS = (
    ("浮子质量", "float_mass"),
    ("振子质量", "oscillator_mass"),
    ("附加质量", "added_mass"),
    ("附加转动惯量", "added_pitch_inertia"),
    ("直线阻尼系数", "linear_damping"),
    ("旋转阻尼系数", "rotary_damping"),
    ("直线弹簧刚度", "heave_spring"),
    ("扭转弹簧刚度", "pitch_spring"),
    ("波浪激励力", "wave_force"),
    ("波浪激励力矩", "wave_moment"),
    ("浮子纵摇转动惯量", "float_pitch_inertia"),
)


def sensitivity_outputs(
    time: np.ndarray, state: np.ndarray
) -> dict[str, float]:
    """提取稳健的幅值/RMS 指标，避免瞬时值接近零放大灵敏度。"""
    _ = time
    relative_heave = state[:, 0] - state[:, 2]
    relative_heave_velocity = state[:, 1] - state[:, 3]
    relative_pitch = state[:, 4] - state[:, 6]
    relative_pitch_velocity = state[:, 5] - state[:, 7]
    return {
        "浮子垂荡最大振幅": float(np.max(np.abs(state[:, 0]))),
        "振子垂荡最大振幅": float(np.max(np.abs(state[:, 2]))),
        "浮子垂荡速度 RMS": float(np.sqrt(np.mean(state[:, 1] ** 2))),
        "振子垂荡速度 RMS": float(np.sqrt(np.mean(state[:, 3] ** 2))),
        "相对垂荡位移 RMS": float(np.sqrt(np.mean(relative_heave**2))),
        "相对垂荡速度 RMS": float(np.sqrt(np.mean(relative_heave_velocity**2))),
        "浮子纵摇最大振幅": float(np.max(np.abs(state[:, 4]))),
        "振子纵摇最大振幅": float(np.max(np.abs(state[:, 6]))),
        "浮子纵摇角速度 RMS": float(np.sqrt(np.mean(state[:, 5] ** 2))),
        "振子纵摇角速度 RMS": float(np.sqrt(np.mean(state[:, 7] ** 2))),
        "相对纵摇角位移 RMS": float(np.sqrt(np.mean(relative_pitch**2))),
        "相对纵摇角速度 RMS": float(np.sqrt(np.mean(relative_pitch_velocity**2))),
    }


def run_parameter_sensitivity(
    params: Parameters,
) -> tuple[bool, dict[str, Any]]:
    """对关键参数做 ±10% 单因素灵敏度分析并打印摘要。"""
    base_time, base_state = solve_model(params, step=OUTPUT_STEP)
    base_outputs = sensitivity_outputs(base_time, base_state)
    sensitivity_by_parameter: dict[str, dict[str, float]] = {}

    for label, attribute in SENSITIVITY_PARAMETERS:
        base_value = float(getattr(params, attribute))
        low_params = replace(
            params,
            **{attribute: base_value * (1.0 - SENSITIVITY_PERTURBATION)},
        )
        high_params = replace(
            params,
            **{attribute: base_value * (1.0 + SENSITIVITY_PERTURBATION)},
        )
        low_time, low_state = solve_model(low_params, step=OUTPUT_STEP)
        high_time, high_state = solve_model(high_params, step=OUTPUT_STEP)
        low_outputs = sensitivity_outputs(low_time, low_state)
        high_outputs = sensitivity_outputs(high_time, high_state)
        sensitivity_by_parameter[label] = {
            name: (
                (high_outputs[name] - low_outputs[name])
                / (2.0 * SENSITIVITY_PERTURBATION * base_outputs[name])
                if abs(base_outputs[name]) > 1e-12
                else float("nan")
            )
            for name in base_outputs
        }

    heave_names = tuple(name for name in base_outputs if "垂荡" in name)
    pitch_names = tuple(name for name in base_outputs if "纵摇" in name)
    summaries = []
    for label, _ in SENSITIVITY_PARAMETERS:
        values = sensitivity_by_parameter[label]
        heave_score = max(abs(values[name]) for name in heave_names)
        pitch_score = max(abs(values[name]) for name in pitch_names)
        summaries.append((label, heave_score, pitch_score))

    print("\n=== 检验三：关键参数 ±10% 灵敏度分析 ===")
    print("灵敏度系数定义：S=(Y(1.1p)-Y(0.9p))/(0.2Y(p))")
    print("参数 | 垂荡运动 max|S| | 纵摇运动 max|S|")
    for label, heave_score, pitch_score in summaries:
        print(f"{label} | {heave_score:.3e} | {pitch_score:.3e}")

    heave_rank = sorted(summaries, key=lambda item: item[1], reverse=True)
    pitch_rank = sorted(summaries, key=lambda item: item[2], reverse=True)
    print(
        "垂荡运动敏感参数（按 max|S|）："
        + "、".join(item[0] for item in heave_rank[:4])
    )
    print(
        "纵摇运动敏感参数（按 max|S|）："
        + "、".join(item[0] for item in pitch_rank[:4])
    )

    _, base_state_fine = solve_model(params, step=ENERGY_CHECK_STEP)
    relative_velocity = base_state_fine[:, 1] - base_state_fine[:, 3]
    relative_angular_velocity = base_state_fine[:, 5] - base_state_fine[:, 7]
    linear_power = params.linear_damping * relative_velocity**2
    rotary_power = params.rotary_damping * relative_angular_velocity**2
    linear_average = float(np.mean(linear_power))
    rotary_average = float(np.mean(rotary_power))
    ratio = rotary_average / linear_average if linear_average > 0.0 else float("nan")
    print(
        f"PTO 平均耗散功率：直线={linear_average:.3e} W，"
        f"旋转={rotary_average:.3e} W，旋转/直线={ratio:.3e}"
    )
    print("旋转 PTO 输出较小时，请重点复核角度单位、转动惯量转轴和旋转阻尼单位。")
    report = {
        "summary": summaries,
        "linear_average": linear_average,
        "rotary_average": rotary_average,
        "ratio": ratio,
    }
    return True, report


def pitch_stiffness_matrix(params: Parameters) -> np.ndarray:
    """返回当前 G0/G1 解释下的 2×2 纵摇刚度矩阵。"""
    float_gravity_stiffness, oscillator_gravity_stiffness = gravity_pitch_terms(
        params
    )
    return np.array(
        [
            [
                params.hydrostatic_pitch_stiffness
                + params.pitch_spring
                + float_gravity_stiffness,
                -params.pitch_spring,
            ],
            [
                -params.pitch_spring,
                params.pitch_spring - oscillator_gravity_stiffness,
            ],
        ],
        dtype=float,
    )


def run_model_validity_check(
    params: Parameters, time: np.ndarray, state: np.ndarray
) -> tuple[bool, dict[str, float]]:
    """检查小角度假设与当前纵摇刚度矩阵的线性稳定性。"""
    _ = time
    maximum_float_angle = float(np.max(np.abs(state[:, 4])))
    maximum_oscillator_angle = float(np.max(np.abs(state[:, 6])))
    maximum_float_angle_deg = math.degrees(maximum_float_angle)
    maximum_oscillator_angle_deg = math.degrees(maximum_oscillator_angle)
    small_angle_passed = max(
        maximum_float_angle_deg, maximum_oscillator_angle_deg
    ) <= SMALL_ANGLE_LIMIT_DEG
    stiffness = pitch_stiffness_matrix(params)
    eigenvalues = np.linalg.eigvalsh(stiffness)
    stability_passed = bool(np.all(eigenvalues > 0.0))

    print("\n=== 检验四：小角度与纵摇稳定性 ===")
    print(
        f"当前模型={params.gravity_model.upper()}，"
        f"a={params.install_height_a:.6f} m，"
        f"r0={params.oscillator_orbit_radius:.6f} m"
    )
    print(
        f"max|theta1|={maximum_float_angle:.6e} rad，"
        f"{maximum_float_angle_deg:.4f}°；"
        f"max|theta2|={maximum_oscillator_angle:.6e} rad，"
        f"{maximum_oscillator_angle_deg:.4f}°"
    )
    print(
        f"小角度审核阈值={SMALL_ANGLE_LIMIT_DEG:.2f}°，"
        f"判定={'通过' if small_angle_passed else '需关注'}"
    )
    print(
        "纵摇刚度矩阵特征值="
        + ", ".join(f"{value:.6e}" for value in eigenvalues)
        + f"，判定={'通过' if stability_passed else '不通过'}"
    )
    passed = small_angle_passed and stability_passed
    print(f"模型适用性总判定：{'通过' if passed else '需关注'}")
    return passed, {
        "maximum_float_angle": maximum_float_angle,
        "maximum_oscillator_angle": maximum_oscillator_angle,
        "limit": math.radians(SMALL_ANGLE_LIMIT_DEG),
        "limit_deg": SMALL_ANGLE_LIMIT_DEG,
    }


def run_install_height_sensitivity(
    params: Parameters,
) -> tuple[bool, list[dict[str, float]]]:
    """按最终建模文档的尺度化取值扫描安装高度 a。"""
    rows: list[dict[str, float]] = []
    passed = True
    print("\n=== 补充检验：安装高度 a 敏感性 ===")
    print("a (m) | r0 (m) | J2O (kg*m^2) | Kg (N*m) | max|theta1| | max|theta2|")
    for install_height in INSTALL_HEIGHT_SAMPLES:
        scan_params = replace(
            params,
            install_height_a=install_height,
            gravity_model="G1",
        )
        eigenvalues = np.linalg.eigvalsh(pitch_stiffness_matrix(scan_params))
        minimum_eigenvalue = float(np.min(eigenvalues))
        time, state = solve_model(scan_params, step=OUTPUT_STEP)
        row = {
            "a": float(install_height),
            "r0": scan_params.oscillator_orbit_radius,
            "J2O": scan_params.effective_oscillator_pitch_inertia,
            "Kg": scan_params.gravity_pitch_stiffness,
            "max_theta1": float(np.max(np.abs(state[:, 4]))),
            "max_theta2": float(np.max(np.abs(state[:, 6]))),
            "rms_theta1": float(np.sqrt(np.mean(state[:, 4] ** 2))),
            "rms_theta2": float(np.sqrt(np.mean(state[:, 6] ** 2))),
            "minimum_eigenvalue": minimum_eigenvalue,
        }
        rows.append(row)
        case_passed = bool(
            minimum_eigenvalue > 0.0 and np.all(np.isfinite(state))
        )
        passed = passed and case_passed
        print(
            f"{row['a']:.3f} | {row['r0']:.6f} | {row['J2O']:.6f} | "
            f"{row['Kg']:.6f} | {row['max_theta1']:.6e} | "
            f"{row['max_theta2']:.6e}"
        )
    print(f"a 敏感性扫描判定：{'通过' if passed else '需关注'}")
    return passed, rows


def run_gravity_model_comparison(
    params: Parameters,
) -> tuple[bool, dict[str, Any]]:
    """比较最终修正版的 G1 与不显式补充重力的 G0。"""
    g0_params = replace(params, gravity_model="G0")
    g1_params = replace(params, gravity_model="G1")
    time, g0_state = solve_model(g0_params, step=OUTPUT_STEP)
    g1_time, g1_state = solve_model(g1_params, step=OUTPUT_STEP)
    if not np.allclose(time, g1_time, rtol=0.0, atol=1e-12):
        raise ValueError("G0 与 G1 的输出时间网格不一致")

    rows = [
        (
            "G0",
            0.0,
            float(np.max(np.abs(g0_state[:, 4]))),
            float(np.max(np.abs(g0_state[:, 6]))),
            float(np.sqrt(np.mean(g0_state[:, 4] ** 2))),
            float(np.sqrt(np.mean(g0_state[:, 6] ** 2))),
        ),
        (
            "G1",
            params.gravity_pitch_stiffness,
            float(np.max(np.abs(g1_state[:, 4]))),
            float(np.max(np.abs(g1_state[:, 6]))),
            float(np.sqrt(np.mean(g1_state[:, 4] ** 2))),
            float(np.sqrt(np.mean(g1_state[:, 6] ** 2))),
        ),
    ]
    maximum_float_difference = float(np.max(np.abs(g1_state[:, 4] - g0_state[:, 4])))
    maximum_oscillator_difference = float(np.max(np.abs(g1_state[:, 6] - g0_state[:, 6])))
    passed = bool(
        np.all(np.isfinite(g0_state))
        and np.all(np.isfinite(g1_state))
    )
    print("\n=== 补充检验：G0/G1 重力解释对比 ===")
    print("模型 | Kg (N*m) | max|theta1| | max|theta2| | RMS theta1 | RMS theta2")
    for row in rows:
        print(
            f"{row[0]} | {row[1]:.6e} | {row[2]:.6e} | {row[3]:.6e} | "
            f"{row[4]:.6e} | {row[5]:.6e}"
        )
    print(
        f"G0/G1 最大纵摇角差异：浮子={maximum_float_difference:.6e} rad，"
        f"振子={maximum_oscillator_difference:.6e} rad；"
        f"判定={'通过' if passed else '需关注'}"
    )
    return passed, {
        "rows": rows,
        "time": time,
        "g0_state": g0_state,
        "g1_state": g1_state,
        "maximum_float_difference": maximum_float_difference,
        "maximum_oscillator_difference": maximum_oscillator_difference,
    }


def rows_from_solution(
    time: np.ndarray, state: np.ndarray
) -> list[tuple[float, ...]]:
    """整理为 result3.xlsx 的列顺序。"""
    rows = []
    for current_time, values in zip(time, state):
        x_float, v_float, x_oscillator, v_oscillator = values[:4]
        theta_float, omega_float, theta_oscillator, omega_oscillator = values[4:]
        rows.append(
            (
                float(current_time),
                float(x_float),
                float(v_float),
                float(theta_float),
                float(omega_float),
                float(x_oscillator),
                float(v_oscillator),
                float(theta_oscillator),
                float(omega_oscillator),
            )
        )
    return rows


def write_result(path: Path, rows: list[tuple[float, ...]]) -> None:
    """写入题目提供的 result3 模板，并清理模板中的旧数据。"""
    if path.exists():
        workbook = load_workbook(path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet1"
        worksheet.merge_cells("A1:A2")
        worksheet.merge_cells("B1:E1")
        worksheet.merge_cells("F1:I1")

    headers = (
        "时间（秒）",
        "浮子",
        "振子",
    )
    worksheet["A1"] = headers[0]
    worksheet["B1"] = headers[1]
    worksheet["F1"] = headers[2]
    subheaders = (
        "垂荡位移（米）",
        "垂荡速度（米/秒）",
        "纵摇角位移（弧度）",
        "纵摇角速度（弧度/秒）",
    )
    for column, value in zip(("B", "C", "D", "E"), subheaders):
        worksheet[f"{column}2"] = value
    for column, value in zip(("F", "G", "H", "I"), subheaders):
        worksheet[f"{column}2"] = value

    # 与第一问示例表一致：表头不使用填充色、白色字体或加粗装饰。
    plain_header_style = [
        {
            "font": copy(worksheet.cell(3, column).font),
            "fill": copy(worksheet.cell(3, column).fill),
            "border": copy(worksheet.cell(3, column).border),
            "alignment": copy(worksheet.cell(3, column).alignment),
            "protection": copy(worksheet.cell(3, column).protection),
        }
        for column in range(1, 10)
    ]
    for row_number in (1, 2):
        for column_number in range(1, 10):
            cell = worksheet.cell(row_number, column_number)
            style = plain_header_style[column_number - 1]
            cell.font = style["font"]
            cell.fill = style["fill"]
            cell.border = style["border"]
            cell.alignment = style["alignment"]
            cell.protection = style["protection"]

    template_max_row = worksheet.max_row
    template_styles = [
        {
            "font": copy(worksheet.cell(3, column).font),
            "fill": copy(worksheet.cell(3, column).fill),
            "border": copy(worksheet.cell(3, column).border),
            "alignment": copy(worksheet.cell(3, column).alignment),
            "protection": copy(worksheet.cell(3, column).protection),
        }
        for column in range(1, 10)
    ]

    for row_number, values in enumerate(rows, start=3):
        for column_number, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column_number, value=value)
            if row_number > template_max_row:
                style = template_styles[column_number - 1]
                cell.font = style["font"]
                cell.fill = style["fill"]
                cell.border = style["border"]
                cell.alignment = style["alignment"]
                cell.protection = style["protection"]
            cell.number_format = "0.000000"

    for row_number in range(3 + len(rows), worksheet.max_row + 1):
        for column_number in range(1, 10):
            worksheet.cell(row=row_number, column=column_number).value = None

    # 第一问示例表不设置实线网格边框，表格仅保留 Excel 默认网格线。
    for row_number in range(1, worksheet.max_row + 1):
        for column_number in range(1, 10):
            worksheet.cell(row_number, column_number).border = Border()

    worksheet.freeze_panes = "A3"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def output_directories() -> tuple[Path, Path]:
    """返回第三问结果表格和图像的输出目录。"""
    root = project_root()
    return root / "结果表格" / "第三问表格", root / "结果图像" / "第三问图像"


def write_q3_workbook(output_path: Path) -> Path:
    """按示例工作簿格式输出单工作表的第三问完整结果。"""
    source_path = project_root() / "data" / "result3.xlsx"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    copyfile(source_path, output_path)
    return output_path


def write_validation_sheet(
    workbook: Workbook,
    name: str,
    headers: tuple[str, ...],
    rows: list[tuple[Any, ...]],
) -> None:
    """按示例工作簿的无填充样式写入一个验证明细表。"""
    worksheet = workbook.create_sheet(name)
    for column, header in enumerate(headers, start=1):
        worksheet.cell(row=1, column=column, value=header)
    for row_number, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            if isinstance(value, (int, float)):
                cell.number_format = "0.000000E+00"
    for column in range(1, len(headers) + 1):
        values = [worksheet.cell(row=row, column=column).value for row in range(1, worksheet.max_row + 1)]
        width = max(len(str(value)) if value is not None else 0 for value in values) + 2
        worksheet.column_dimensions[chr(64 + column)].width = min(max(width, 12), 32)
    worksheet.freeze_panes = "A2"


def write_validation_workbook(
    numerical_rows: list[tuple[str, float, str]],
    equation_report: dict[str, float | bool],
    sensitivity_report: dict[str, Any],
    small_angle_report: dict[str, float],
    a_rows: list[dict[str, float]],
    gravity_report: dict[str, Any],
) -> Path:
    """输出补充验证结果；主结果表仍保持单工作表示例格式。"""
    table_dir, _ = output_directories()
    table_dir.mkdir(parents=True, exist_ok=True)
    path = table_dir / "第三问_补充验证结果.xlsx"
    workbook = Workbook()
    del workbook[workbook.sheetnames[0]]

    write_validation_sheet(
        workbook,
        "数值交叉验证",
        ("比较对象", "最大相对误差或差异", "结论"),
        numerical_rows,
    )
    equation_rows = [
        (
            "动力学方程归一化残差最大值",
            equation_report["equation_max"],
            "通过"
            if equation_report["equation_max"] < ENERGY_RELATIVE_TOLERANCE
            else "需关注",
        ),
        (
            "动力学方程归一化残差 RMS",
            equation_report["equation_rms"],
            "通过"
            if equation_report["equation_rms"] < ENERGY_RELATIVE_TOLERANCE
            else "需关注",
        ),
        (
            "能量收支相对残差",
            equation_report["energy_relative"],
            "通过"
            if equation_report["energy_relative"] < ENERGY_RELATIVE_TOLERANCE
            else "需关注",
        ),
        (
            "能量导数差分残差",
            equation_report["differential_relative"],
            "通过"
            if equation_report["differential_relative"] < ENERGY_RELATIVE_TOLERANCE
            else "需关注",
        ),
        (
            "直线 PTO 最小耗散功率（W）",
            equation_report["linear_pto_min"],
            "通过" if equation_report["linear_pto_min"] >= -1e-10 else "需关注",
        ),
        (
            "旋转 PTO 最小耗散功率（W）",
            equation_report["rotary_pto_min"],
            "通过" if equation_report["rotary_pto_min"] >= -1e-10 else "需关注",
        ),
        (
            "浮子最大垂荡位移（m）",
            equation_report["maximum_displacement"],
            "通过",
        ),
        (
            "水线允许下界（m）",
            equation_report["waterline_lower"],
            "通过" if equation_report["waterline_passed"] else "需关注",
        ),
        (
            "水线允许上界（m）",
            equation_report["waterline_upper"],
            "通过" if equation_report["waterline_passed"] else "需关注",
        ),
    ]
    write_validation_sheet(
        workbook,
        "方程能量检验",
        ("指标", "实际结果", "结论"),
        equation_rows,
    )
    write_validation_sheet(
        workbook,
        "安装高度敏感性",
        (
            "a（m）",
            "r0（m）",
            "J2O（kg*m^2）",
            "Kg（N*m）",
            "max|theta1|（rad）",
            "max|theta2|（rad）",
            "RMS theta1（rad）",
            "RMS theta2（rad）",
            "纵摇刚度矩阵最小特征值",
        ),
        [
            (
                row["a"],
                row["r0"],
                row["J2O"],
                row["Kg"],
                row["max_theta1"],
                row["max_theta2"],
                row["rms_theta1"],
                row["rms_theta2"],
                row["minimum_eigenvalue"],
            )
            for row in a_rows
        ],
    )
    gravity_rows = [
        tuple(row)
        + ((None, None) if row[0] == "G0" else (
            gravity_report["maximum_float_difference"],
            gravity_report["maximum_oscillator_difference"],
        ))
        for row in gravity_report["rows"]
    ]
    write_validation_sheet(
        workbook,
        "G0G1对比",
        (
            "模型",
            "Kg（N*m）",
            "max|theta1|（rad）",
            "max|theta2|（rad）",
            "RMS theta1（rad）",
            "RMS theta2（rad）",
            "与G0浮子最大差异（rad）",
            "与G0振子最大差异（rad）",
        ),
        gravity_rows,
    )
    write_validation_sheet(
        workbook,
        "小角度检验",
        ("指标", "实际结果（rad）", "阈值（rad）", "结论"),
        [
            (
                "max|theta1|",
                small_angle_report["maximum_float_angle"],
                small_angle_report["limit"],
                "通过" if small_angle_report["maximum_float_angle"] < small_angle_report["limit"] else "需关注",
            ),
            (
                "max|theta2|",
                small_angle_report["maximum_oscillator_angle"],
                small_angle_report["limit"],
                "通过" if small_angle_report["maximum_oscillator_angle"] < small_angle_report["limit"] else "需关注",
            ),
        ],
    )
    write_validation_sheet(
        workbook,
        "关键参数敏感性",
        ("关键参数", "垂荡运动最大|S|", "纵摇运动最大|S|"),
        sensitivity_report["summary"],
    )
    workbook.save(path)
    workbook.close()
    return path


def save_figure(figure: Any, output_stem: Path) -> None:
    """同时保存高分辨率 PNG 和可编辑 SVG。"""
    figure.savefig(
        output_stem.with_suffix(".png"),
        dpi=300,
        facecolor="white",
        bbox_inches="tight",
        pad_inches=0.12,
    )
    figure.savefig(
        output_stem.with_suffix(".svg"),
        format="svg",
        facecolor="white",
        bbox_inches="tight",
        pad_inches=0.12,
    )


def write_q3_charts(
    params: Parameters,
    time: np.ndarray,
    state: np.ndarray,
    *,
    a_rows: list[dict[str, float]],
    gravity_report: dict[str, Any],
    sensitivity_report: dict[str, Any],
    small_angle_report: dict[str, float],
) -> Path:
    """生成第三问主响应和补充验证论文图。"""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    _, image_dir = output_directories()
    image_dir.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "Microsoft YaHei",
            "axes.unicode_minus": False,
            "font.size": 10,
            "axes.labelsize": 11,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "legend.fontsize": 9,
            "svg.fonttype": "none",
        }
    )

    # 图1：垂荡位移和速度。
    figure, axes = plt.subplots(2, 1, figsize=(10, 7.0), sharex=True, dpi=220)
    axes[0].plot(time, state[:, 0], color="#1F5A85", linewidth=1.5, label="浮子")
    axes[0].plot(time, state[:, 2], color="#D1495B", linewidth=1.3, label="振子")
    axes[0].set_ylabel("垂荡位移（m）")
    axes[0].legend(frameon=False, ncol=2, loc="upper right")
    axes[1].plot(time, state[:, 1], color="#1F5A85", linewidth=1.5, label="浮子")
    axes[1].plot(time, state[:, 3], color="#D1495B", linewidth=1.3, label="振子")
    for axis in axes:
        axis.set_xlabel("时间（秒）")
        axis.tick_params(axis="x", labelbottom=True)
    axes[1].set_ylabel("垂荡速度（m/s）")
    for axis in axes:
        axis.grid(True, color="#D9E2EA", linewidth=0.7, alpha=0.8)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
    save_figure(figure, image_dir / "第三问_图1_垂荡响应")
    plt.close(figure)

    # 图2：纵摇角和角速度。
    figure, axes = plt.subplots(2, 1, figsize=(10, 7.0), sharex=True, dpi=220)
    axes[0].plot(time, state[:, 4], color="#1F5A85", linewidth=1.5, label="浮子")
    axes[0].plot(time, state[:, 6], color="#D1495B", linewidth=1.3, label="振子")
    axes[0].set_ylabel("纵摇角（rad）")
    axes[0].legend(frameon=False, ncol=2, loc="upper right")
    axes[1].plot(time, state[:, 5], color="#1F5A85", linewidth=1.5, label="浮子")
    axes[1].plot(time, state[:, 7], color="#D1495B", linewidth=1.3, label="振子")
    for axis in axes:
        axis.set_xlabel("时间（秒）")
        axis.tick_params(axis="x", labelbottom=True)
    axes[1].set_ylabel("角速度（rad/s）")
    for axis in axes:
        axis.grid(True, color="#D9E2EA", linewidth=0.7, alpha=0.8)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
    save_figure(figure, image_dir / "第三问_图2_纵摇响应")
    plt.close(figure)

    # 图3：PTO 相对运动与耗散功率。
    relative_velocity = state[:, 3] - state[:, 1]
    relative_angular_velocity = state[:, 7] - state[:, 5]
    linear_power = params.linear_damping * relative_velocity**2
    rotary_power = params.rotary_damping * relative_angular_velocity**2
    figure, axes = plt.subplots(2, 1, figsize=(10, 7.0), sharex=True, dpi=220)
    axes[0].plot(time, relative_velocity, color="#2A9D8F", linewidth=1.4, label="相对垂荡速度")
    angular_axis = axes[0].twinx()
    angular_axis.plot(
        time,
        relative_angular_velocity,
        color="#F4A261",
        linewidth=1.4,
        label="相对角速度",
    )
    axes[0].set_ylabel("相对垂荡速度（m/s）", color="#2A9D8F")
    angular_axis.set_ylabel("相对角速度（rad/s）", color="#F4A261")
    handles, labels = axes[0].get_legend_handles_labels()
    angular_handles, angular_labels = angular_axis.get_legend_handles_labels()
    axes[0].legend(
        handles + angular_handles,
        labels + angular_labels,
        frameon=False,
        ncol=2,
        loc="upper right",
    )
    axes[1].plot(time, linear_power, color="#1F5A85", linewidth=1.4, label="直线 PTO")
    power_axis = axes[1].twinx()
    power_axis.plot(
        time,
        rotary_power,
        color="#D1495B",
        linewidth=1.4,
        label="旋转 PTO",
    )
    for axis in axes:
        axis.set_xlabel("时间（秒）")
        axis.tick_params(axis="x", labelbottom=True)
    axes[1].set_ylabel("直线 PTO 耗散功率（W）", color="#1F5A85")
    power_axis.set_ylabel("旋转 PTO 耗散功率（W）", color="#D1495B")
    handles, labels = axes[1].get_legend_handles_labels()
    power_handles, power_labels = power_axis.get_legend_handles_labels()
    axes[1].legend(
        handles + power_handles,
        labels + power_labels,
        frameon=False,
        ncol=2,
        loc="upper right",
    )
    for axis in axes:
        axis.grid(True, color="#D9E2EA", linewidth=0.7, alpha=0.8)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
    power_axis.spines["top"].set_visible(False)
    angular_axis.spines["top"].set_visible(False)
    save_figure(figure, image_dir / "第三问_图3_PTO相对运动与耗散功率")
    plt.close(figure)

    # 图4：安装高度 a 敏感性。
    a_values = np.array([row["a"] for row in a_rows], dtype=float)
    maximum_float_angles = np.array(
        [row["max_theta1"] for row in a_rows], dtype=float
    )
    maximum_oscillator_angles = np.array(
        [row["max_theta2"] for row in a_rows], dtype=float
    )
    figure, axis = plt.subplots(1, 1, figsize=(10.5, 5.8), dpi=220)
    axis.plot(
        a_values,
        maximum_float_angles,
        marker="o",
        color="#1F5A85",
        linewidth=1.6,
        label="浮子最大纵摇角",
    )
    axis.plot(
        a_values,
        maximum_oscillator_angles,
        marker="s",
        color="#D1495B",
        linewidth=1.6,
        label="振子最大纵摇角",
    )
    axis.set_xlabel("安装高度 a（米）")
    axis.set_ylabel("最大绝对纵摇角（弧度）")
    axis.ticklabel_format(
        axis="y",
        style="sci",
        scilimits=(0, 0),
        useMathText=True,
    )
    axis.legend(
        frameon=False,
        ncol=1,
        loc="upper left",
        bbox_to_anchor=(1.02, 1.0),
        borderaxespad=0.0,
        labelspacing=0.6,
    )
    axis.grid(True, color="#D9E2EA", linewidth=0.7, alpha=0.8)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    save_figure(figure, image_dir / "第三问_图4_安装高度敏感性")
    plt.close(figure)

    # 图5：G0 与 G1 的纵摇解释对比。
    gravity_time = gravity_report["time"]
    g0_state = gravity_report["g0_state"]
    g1_state = gravity_report["g1_state"]
    figure, axes = plt.subplots(2, 1, figsize=(10, 7.0), sharex=True, dpi=220)
    axes[0].plot(gravity_time, g0_state[:, 4], color="#7A7A7A", linewidth=1.4, label="G0")
    axes[0].plot(gravity_time, g1_state[:, 4], color="#1F5A85", linewidth=1.4, label="G1")
    axes[0].set_ylabel("浮子纵摇角（rad）")
    axes[0].legend(frameon=False, ncol=2, loc="upper right")
    axes[1].plot(gravity_time, g0_state[:, 6], color="#7A7A7A", linewidth=1.4, label="G0")
    axes[1].plot(gravity_time, g1_state[:, 6], color="#D1495B", linewidth=1.4, label="G1")
    axes[1].set_ylabel("振子纵摇角（rad）")
    for axis in axes:
        axis.set_xlabel("时间（秒）")
        axis.tick_params(axis="x", labelbottom=True)
        axis.grid(True, color="#D9E2EA", linewidth=0.7, alpha=0.8)
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
    axes[1].legend(frameon=False, ncol=2, loc="upper right")
    save_figure(figure, image_dir / "第三问_图5_G0与G1纵摇对比")
    plt.close(figure)

    # 图6：小角度适用性。
    angle_values = [
        small_angle_report["maximum_float_angle"],
        small_angle_report["maximum_oscillator_angle"],
    ]
    figure, axis = plt.subplots(1, 1, figsize=(8.5, 5.8), dpi=220)
    axis.bar(
        ("浮子", "振子"),
        angle_values,
        color=("#1F5A85", "#D1495B"),
        width=0.55,
        label="最大绝对纵摇角",
    )
    axis.axhline(
        small_angle_report["limit"],
        color="#333333",
        linestyle="--",
        linewidth=1.2,
        label="小角度阈值 5°",
    )
    axis.set_xlabel("对象")
    axis.set_ylabel("最大绝对纵摇角（弧度）")
    axis.ticklabel_format(
        axis="y",
        style="sci",
        scilimits=(0, 0),
        useMathText=True,
    )
    axis.legend(
        frameon=False,
        loc="upper right",
        bbox_to_anchor=(1.0, 0.78),
        borderaxespad=0.0,
    )
    axis.grid(True, axis="y", color="#D9E2EA", linewidth=0.7, alpha=0.8)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    save_figure(figure, image_dir / "第三问_图6_小角度适用性检验")
    plt.close(figure)

    # 图7：按指定样式绘制中心对称的关键参数灵敏度龙卷风图。
    sensitivity_summary = sensitivity_report["summary"]
    sensitivity_by_label = {
        item[0]: (float(item[1]), float(item[2]))
        for item in sensitivity_summary
    }
    parameter_order = (
        (r"$m_1$", "浮子质量"),
        (r"$m_2$", "振子质量"),
        (r"$j_a$", "附加转动惯量"),
        (r"$J_1$", "浮子纵摇转动惯量"),
        (r"$m_a$", "附加质量"),
        (r"$k$", "直线弹簧刚度"),
        (r"$k_\theta$", "扭转弹簧刚度"),
        (r"$L$", "波浪激励力矩"),
        (r"$f$", "波浪激励力"),
        (r"$c$", "直线阻尼系数"),
        (r"$c_\theta$", "旋转阻尼系数"),
    )
    symbols = [item[0] for item in parameter_order]
    labels = [item[1] for item in parameter_order]
    heave_values = np.array(
        [sensitivity_by_label[label][0] for label in labels], dtype=float
    )
    pitch_values = np.array(
        [sensitivity_by_label[label][1] for label in labels], dtype=float
    )
    y_values = np.arange(len(labels), dtype=float)
    maximum_score = max(float(np.max(heave_values)), float(np.max(pitch_values)))
    x_limit = max(1.0, math.ceil(maximum_score * 1.25 * 2.0) / 2.0)
    tick_values = np.linspace(-x_limit, x_limit, 9)

    def format_tornado_value(value: float) -> str:
        if abs(value) < 1e-12:
            return "0"
        exponent = int(math.floor(math.log10(abs(value))))
        mantissa = value / (10.0**exponent)
        return rf"${mantissa:.2f}\times10^{{{exponent}}}$"

    def format_tornado_tick(value: float) -> str:
        magnitude = abs(value)
        if magnitude < 1e-12:
            return "0"
        if math.isclose(magnitude, round(magnitude), abs_tol=1e-9):
            return f"{int(round(magnitude))}"
        return f"{magnitude:.2f}".rstrip("0").rstrip(".")

    from matplotlib.patches import Patch

    figure = plt.figure(figsize=(14.0, 9.6), dpi=220)
    grid = figure.add_gridspec(
        nrows=2,
        ncols=2,
        height_ratios=(0.14, 0.86),
        width_ratios=(0.88, 0.12),
        hspace=0.0,
        wspace=0.0,
    )
    header_axis = figure.add_subplot(grid[0, 0])
    header_symbol_axis = figure.add_subplot(grid[0, 1])
    axis = figure.add_subplot(grid[1, 0])
    symbol_axis = figure.add_subplot(grid[1, 1], sharey=axis)
    header_axis.axis("off")
    header_symbol_axis.axis("off")
    heave_color = "#4C78A8"
    pitch_color = "#F58518"
    axis.barh(
        y_values,
        -heave_values,
        height=0.66,
        color=heave_color,
        alpha=0.9,
        label="垂荡运动（左）",
    )
    axis.barh(
        y_values,
        pitch_values,
        height=0.66,
        color=pitch_color,
        alpha=0.9,
        label="纵摇运动（右）",
    )
    axis.axvline(0.0, color="#222222", linewidth=1.25, zorder=3)
    axis.set_xlim(-x_limit, x_limit)
    axis.set_xticks(tick_values)
    axis.set_xticklabels([format_tornado_tick(value) for value in tick_values])
    axis.set_xlabel("最大绝对灵敏度系数 |S|")
    axis.set_yticks(y_values)
    axis.set_yticklabels([])
    axis.invert_yaxis()
    axis.grid(True, axis="x", color="#D9E2EA", linewidth=0.7, alpha=0.8)
    axis.set_axisbelow(True)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.spines["left"].set_visible(False)
    axis.tick_params(axis="y", length=0)
    header_axis.legend(
        handles=[
            Patch(facecolor=heave_color, edgecolor="none"),
            Patch(facecolor=pitch_color, edgecolor="none"),
        ],
        labels=["垂荡运动（左）", "纵摇运动（右）"],
        frameon=False,
        ncol=2,
        loc="center right",
        bbox_to_anchor=(0.98, 0.60),
        borderaxespad=0.0,
    )

    value_pad = x_limit * 0.025
    for y_value, heave_value, pitch_value in zip(
        y_values, heave_values, pitch_values
    ):
        axis.text(
            -heave_value - value_pad if heave_value > 0 else -value_pad,
            y_value,
            format_tornado_value(heave_value),
            ha="right",
            va="center",
            color=heave_color,
        )
        axis.text(
            pitch_value + value_pad if pitch_value > 0 else value_pad,
            y_value,
            format_tornado_value(pitch_value),
            ha="left",
            va="center",
            color=pitch_color if pitch_value > 0 else "#666666",
        )

    symbol_axis.set_ylim(axis.get_ylim())
    symbol_axis.set_xlim(0.0, 1.0)
    symbol_axis.set_xticks([])
    symbol_axis.set_yticks([])
    symbol_axis.set_frame_on(False)
    for y_value, symbol in zip(y_values, symbols):
        symbol_axis.text(
            0.5,
            y_value,
            symbol,
            transform=symbol_axis.get_yaxis_transform(),
            ha="center",
            va="center",
            bbox={
                "boxstyle": "round,pad=0.25",
                "facecolor": "white",
                "edgecolor": "#444444",
                "linewidth": 1.0,
            },
        )
    header_symbol_axis.text(
        0.5,
        0.45,
        "参数符号",
        transform=header_symbol_axis.transAxes,
        ha="center",
        va="center",
        fontsize=11,
        fontweight="bold",
    )
    figure.subplots_adjust(left=0.08, right=0.98, bottom=0.10, top=0.98)
    save_figure(figure, image_dir / "第三问_图7_关键参数灵敏度龙卷风图")
    plt.close(figure)
    return image_dir


def print_key_results(
    params: Parameters, time: np.ndarray, state: np.ndarray
) -> None:
    """打印论文要求的五个时刻结果。"""
    print("\n=== 问题 3 计算结果 ===")
    print(f"波浪频率：{params.wave_frequency:.4f} s^-1")
    print(f"波浪周期：{params.period:.8f} s")
    print(f"有效浮子质量：{params.effective_float_mass:.6f} kg")
    print(
        "有效浮子纵摇转动惯量："
        f"{params.effective_float_pitch_inertia:.6f} kg*m^2"
    )
    print(
        "静平衡弹簧长度："
        f"{params.equilibrium_spring_length:.7f} m；"
        "振子质心到铰轴距离 r0："
        f"{params.oscillator_orbit_radius:.7f} m"
    )
    print(
        "G1 重力修正 Kg："
        f"{params.gravity_pitch_stiffness:.7f} N*m；"
        "振子关于铰轴转动惯量 J2O："
        f"{params.effective_oscillator_pitch_inertia:.7f} kg*m^2"
    )
    print(
        "\n时间(s) | 浮子 x v theta omega | 振子 x v theta omega"
    )
    for target in PAPER_TIMES:
        index = int(round(target / OUTPUT_STEP))
        if index >= len(time) or not math.isclose(time[index], target, abs_tol=1e-9):
            raise ValueError(f"积分结果中没有时刻 {target:g} s")
        values = state[index]
        print(
            f"{target:7.1f} | "
            f"{values[0]: .8f} {values[1]: .8f} "
            f"{values[4]: .8f} {values[5]: .8f} | "
            f"{values[2]: .8f} {values[3]: .8f} "
            f"{values[6]: .8f} {values[7]: .8f}"
        )


def main() -> None:
    params = read_parameters()
    time, state = solve_model(params)
    output_path = project_root() / "data" / "result3.xlsx"
    write_result(output_path, rows_from_solution(time, state))
    print_key_results(params, time, state)
    numerical_passed, numerical_rows = run_numerical_cross_validation(params)
    equation_energy_passed, equation_report = run_equation_and_energy_check(params)
    sensitivity_completed, sensitivity_report = run_parameter_sensitivity(params)
    validity_passed, small_angle_report = run_model_validity_check(
        params, time, state
    )
    a_sensitivity_passed, a_rows = run_install_height_sensitivity(params)
    gravity_comparison_passed, gravity_report = run_gravity_model_comparison(params)
    result_workbook_path = write_q3_workbook(
        output_directories()[0] / "第三问_结果表格.xlsx"
    )
    validation_workbook_path = write_validation_workbook(
        numerical_rows,
        equation_report,
        sensitivity_report,
        small_angle_report,
        a_rows,
        gravity_report,
    )
    image_dir = write_q3_charts(
        params,
        time,
        state,
        a_rows=a_rows,
        gravity_report=gravity_report,
        sensitivity_report=sensitivity_report,
        small_angle_report=small_angle_report,
    )
    all_passed = all(
        (
            numerical_passed,
            equation_energy_passed,
            sensitivity_completed,
            validity_passed,
            a_sensitivity_passed,
            gravity_comparison_passed,
        )
    )
    print(
        "\n=== 第三问全部检验总判定："
        f"{'通过' if all_passed else '需关注'} ==="
    )
    print(f"\n结果文件：{output_path}")
    print(f"结果表格：{result_workbook_path}")
    print(f"补充验证表格：{validation_workbook_path}")
    print(f"结果图像：{image_dir}")
    print(f"输出点数：{len(time)}（时间范围 0 至 {time[-1]:.1f} s）")


if __name__ == "__main__":
    main()
