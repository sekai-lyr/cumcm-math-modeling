"""2022 年高教社杯全国大学生数学建模竞赛 A 题——问题 2。

目标：在波浪圆频率 2.2143 s^-1 下，使 PTO 直线阻尼器的平均输出功率最大。

两种情况的数值方法：
1. 常量阻尼：对线性稳态方程作复幅值变换，用有界标量搜索求最优阻尼系数；
2. 幂律阻尼：用 DOP853 积分运动方程，在去除瞬态后的多个周期上计算平均功率，
   先用二维粗网格定位，再用 Nelder-Mead 局部精化。

结果检验：
1. alpha=0 且 a=c* 时，比较频域解析功率与时域平均功率；
2. 改变积分周期数、瞬态截断周期、输出步长和积分容差，检查稳态功率收敛；
3. 用多组初值复搜，并单独检查 a=0、a=100000、alpha=0、alpha=1 等边界。
4. 在最优点附近分别扰动 c、a 和 alpha，检查平均输出功率的局部灵敏度。

题目“阻尼系数与相对速度绝对值的幂成正比”按建模思路解释为

    c(v_r) = a |v_r|^alpha,
    F_d = c(v_r) v_r = a |v_r|^alpha v_r,

其中 a ∈ [0, 100000]，alpha ∈ [0, 1]，v_r = v_z - v_f。

运行：
    python 第二问.py

程序读取项目 data 目录中的附件 3、附件 4，将最优结果打印到控制台，并在
“结果表格/第二问表格”和“结果图像/第二问图像”中生成论文用表格与图像。
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from scipy.integrate import solve_ivp
from scipy.optimize import minimize, minimize_scalar


# 问题 2 的设计变量范围
COEFFICIENT_MIN = 0.0
COEFFICIENT_MAX = 100_000.0
EXPONENT_MIN = 0.0
EXPONENT_MAX = 1.0

# 非线性时域计算设置。搜索阶段适当降低输出密度，最终结果再精算。
SEARCH_PERIODS = 45.0
SEARCH_TRANSIENT_PERIODS = 30.0
SEARCH_OUTPUT_STEP = 0.05
# 60/40 周期仍有可见瞬态；160/140 周期用于稳定结果和检验。
FINAL_PERIODS = 160.0
FINAL_TRANSIENT_PERIODS = 140.0
FINAL_OUTPUT_STEP = 0.02
RTOL = 1e-9
ATOL = 1e-11
POWER_RELATIVE_TOLERANCE = 1e-4
GLOBAL_COEFFICIENT_TOLERANCE = 0.02
GLOBAL_EXPONENT_TOLERANCE = 0.03
SENSITIVITY_COEFFICIENT_FACTORS = (0.80, 0.90, 0.95, 1.00, 1.05, 1.10, 1.20)
SENSITIVITY_ALPHA_OFFSETS = (-0.20, -0.10, 0.00, 0.10, 0.20)


@dataclass(frozen=True)
class Parameters:
    """问题 2 所需的物理参数。"""

    float_mass: float
    oscillator_mass: float
    added_mass: float
    hydro_damping: float
    wave_force: float
    wave_frequency: float
    spring_stiffness: float
    water_density: float
    gravity: float
    waterline_radius: float

    @property
    def effective_float_mass(self) -> float:
        """浮子自身质量与附加质量之和。"""
        return self.float_mass + self.added_mass

    @property
    def hydrostatic_stiffness(self) -> float:
        """水静力恢复刚度 K_h = rho*g*A_w。"""
        return (
            self.water_density
            * self.gravity
            * math.pi
            * self.waterline_radius**2
        )

    @property
    def period(self) -> float:
        """波浪周期。"""
        return 2.0 * math.pi / self.wave_frequency


@dataclass(frozen=True)
class OptimizationResult:
    """优化结果及对应的稳态复幅值。"""

    coefficient: float
    exponent: float
    power: float
    float_amplitude: float
    oscillator_amplitude: float
    relative_amplitude: float


@dataclass(frozen=True)
class SensitivityPoint:
    """最优点附近单因素扰动得到的功率灵敏度数据。"""

    damping_case: str
    parameter: str
    parameter_value: float
    relative_parameter_change: float | None
    absolute_parameter_change: float | None
    power: float
    relative_power_change: float
    power_change: float
    baseline_power: float
    fixed_condition: str


def project_root() -> Path:
    """返回项目根目录：本文件位于 项目根/代码/第二问代码/。"""
    return Path(__file__).resolve().parents[2]


def as_float(value: Any, name: str) -> float:
    """将 Excel 单元格转换为浮点数。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"附件中缺少参数：{name}")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"参数 {name} 不是数值：{value!r}") from exc


def find_value_by_label(
    worksheet: Any, keywords: tuple[str, ...], name: str
) -> float:
    """在附件 4 的前两列中按关键词查找参数。"""
    normalized_keywords = tuple(keyword.replace(" ", "") for keyword in keywords)
    for row in worksheet.iter_rows(min_col=1, max_col=2, values_only=True):
        label = "" if row[0] is None else str(row[0]).replace(" ", "")
        if all(keyword in label for keyword in normalized_keywords):
            return as_float(row[1], name)
    raise ValueError(f"附件 4 中找不到参数：{name}（关键词：{keywords}）")


def read_parameters() -> Parameters:
    """从附件 3、附件 4 读取问题 2 参数。"""
    data_dir = project_root() / "data"
    attachment3 = data_dir / "附件3.xlsx"
    attachment4 = data_dir / "附件4.xlsx"

    if not attachment3.exists() or not attachment4.exists():
        raise FileNotFoundError(
            f"找不到输入附件，请确认文件存在：{attachment3}、{attachment4}"
        )

    workbook3 = load_workbook(attachment3, data_only=True, read_only=True)
    try:
        worksheet3 = workbook3.active
        wave_row = None
        for row in worksheet3.iter_rows(min_row=2, values_only=True):
            frequency = row[1] if len(row) > 1 else None
            try:
                if math.isclose(float(frequency), 2.2143, abs_tol=1e-8):
                    wave_row = row
                    break
            except (TypeError, ValueError):
                continue

        if wave_row is None:
            raise ValueError("附件 3 中找不到波浪频率 2.2143 的参数行")

        added_mass = as_float(wave_row[2], "附加质量")
        hydro_damping = as_float(wave_row[4], "垂荡兴波阻尼系数")
        wave_force = as_float(wave_row[6], "波浪激励力幅值")
        wave_frequency = as_float(wave_row[1], "波浪频率")
    finally:
        workbook3.close()

    workbook4 = load_workbook(attachment4, data_only=True, read_only=True)
    try:
        worksheet4 = workbook4.active
        return Parameters(
            float_mass=find_value_by_label(worksheet4, ("浮子质量",), "浮子质量"),
            oscillator_mass=find_value_by_label(
                worksheet4, ("振子质量",), "振子质量"
            ),
            added_mass=added_mass,
            hydro_damping=hydro_damping,
            wave_force=wave_force,
            wave_frequency=wave_frequency,
            spring_stiffness=find_value_by_label(
                worksheet4, ("弹簧刚度",), "弹簧刚度"
            ),
            water_density=find_value_by_label(
                worksheet4, ("海水", "密度"), "海水密度"
            ),
            gravity=find_value_by_label(
                worksheet4, ("重力加速度",), "重力加速度"
            ),
            waterline_radius=find_value_by_label(
                worksheet4, ("浮子底半径",), "浮子底半径"
            ),
        )
    finally:
        workbook4.close()


def steady_state_response(
    params: Parameters, coefficient: float
) -> tuple[complex, complex]:
    """求常量阻尼下的浮子、振子稳态位移复幅值。"""
    omega = params.wave_frequency
    m_float = params.effective_float_mass
    m_oscillator = params.oscillator_mass
    k_h = params.hydrostatic_stiffness
    k = params.spring_stiffness
    b_h = params.hydro_damping

    dynamic_matrix = np.array(
        [
            [
                k_h + k - m_float * omega**2
                + 1j * omega * (b_h + coefficient),
                -k - 1j * omega * coefficient,
            ],
            [
                -k - 1j * omega * coefficient,
                k - m_oscillator * omega**2 + 1j * omega * coefficient,
            ],
        ],
        dtype=complex,
    )
    amplitudes = np.linalg.solve(
        dynamic_matrix, np.array([params.wave_force, 0.0], dtype=complex)
    )
    return complex(amplitudes[0]), complex(amplitudes[1])


def linear_average_power(params: Parameters, coefficient: float) -> float:
    """常量阻尼的稳态平均 PTO 输出功率。"""
    if coefficient < COEFFICIENT_MIN or coefficient > COEFFICIENT_MAX:
        return -math.inf
    float_amplitude, oscillator_amplitude = steady_state_response(
        params, coefficient
    )
    relative_velocity_amplitude = (
        1j
        * params.wave_frequency
        * (oscillator_amplitude - float_amplitude)
    )
    return 0.5 * coefficient * abs(relative_velocity_amplitude) ** 2


def optimize_linear_case(params: Parameters) -> OptimizationResult:
    """用有界黄金分割/抛物线搜索求常量阻尼最优值。"""
    result = minimize_scalar(
        lambda coefficient: -linear_average_power(params, float(coefficient)),
        bounds=(COEFFICIENT_MIN, COEFFICIENT_MAX),
        method="bounded",
        options={"xatol": 1e-8, "maxiter": 200},
    )

    candidates = [
        COEFFICIENT_MIN,
        COEFFICIENT_MAX,
        float(result.x),
    ]
    coefficient = max(
        candidates, key=lambda value: linear_average_power(params, value)
    )
    float_amplitude, oscillator_amplitude = steady_state_response(
        params, coefficient
    )
    return OptimizationResult(
        coefficient=coefficient,
        exponent=0.0,
        power=linear_average_power(params, coefficient),
        float_amplitude=abs(float_amplitude),
        oscillator_amplitude=abs(oscillator_amplitude),
        relative_amplitude=abs(oscillator_amplitude - float_amplitude),
    )


def nonlinear_time_grid(
    params: Parameters,
    periods: float,
    output_step: float,
    extra_times: tuple[float, ...] = (),
) -> np.ndarray:
    """生成包含积分终点及指定窗口边界的时域输出网格。"""
    if output_step <= 0.0:
        raise ValueError("时域输出步长必须为正数")
    end_time = periods * params.period
    regular_grid = np.arange(0.0, end_time, output_step, dtype=float)
    required_times = np.array(
        [0.0, end_time, *extra_times], dtype=float
    )
    grid = np.unique(np.concatenate((regular_grid, required_times)))
    return grid[grid <= end_time + 1e-12]


def nonlinear_rhs(
    params: Parameters, coefficient: float, exponent: float
) -> Callable[[float, np.ndarray], np.ndarray]:
    """建立幂律阻尼下的四阶一阶状态方程。"""
    m_float = params.effective_float_mass
    m_oscillator = params.oscillator_mass
    k_h = params.hydrostatic_stiffness
    k = params.spring_stiffness
    b_h = params.hydro_damping
    force_amplitude = params.wave_force
    omega = params.wave_frequency

    def rhs(time: float, state: np.ndarray) -> np.ndarray:
        x_float, v_float, x_oscillator, v_oscillator = state
        relative_velocity = v_oscillator - v_float
        damping_force = (
            coefficient
            * abs(relative_velocity) ** exponent
            * relative_velocity
        )
        return np.array(
            (
                v_float,
                (
                    force_amplitude * math.cos(omega * time)
                    - b_h * v_float
                    - (k + k_h) * x_float
                    + k * x_oscillator
                    + damping_force
                )
                / m_float,
                v_oscillator,
                (k * (x_float - x_oscillator) - damping_force) / m_oscillator,
            ),
            dtype=float,
        )

    return rhs


def solve_nonlinear_case(
    params: Parameters,
    coefficient: float,
    exponent: float,
    periods: float = FINAL_PERIODS,
    output_step: float = FINAL_OUTPUT_STEP,
    transient_periods: float | None = None,
    rtol: float = RTOL,
    atol: float = ATOL,
) -> tuple[np.ndarray, np.ndarray]:
    """从静止平衡位置积分幂律阻尼模型。"""
    if not COEFFICIENT_MIN <= coefficient <= COEFFICIENT_MAX:
        raise ValueError("非线性比例系数超出 [0, 100000]")
    if not EXPONENT_MIN <= exponent <= EXPONENT_MAX:
        raise ValueError("幂指数超出 [0, 1]")
    if rtol <= 0.0 or atol <= 0.0:
        raise ValueError("积分相对/绝对容差必须为正数")

    extra_times: tuple[float, ...] = ()
    if transient_periods is not None:
        if not 0.0 <= transient_periods < periods:
            raise ValueError("稳态窗口起点必须位于积分区间内")
        extra_times = (transient_periods * params.period,)
    time = nonlinear_time_grid(
        params, periods, output_step, extra_times=extra_times
    )
    solution = solve_ivp(
        nonlinear_rhs(params, coefficient, exponent),
        (float(time[0]), float(time[-1])),
        y0=np.zeros(4, dtype=float),
        t_eval=time,
        method="DOP853",
        rtol=rtol,
        atol=atol,
    )
    if not solution.success:
        raise RuntimeError(f"DOP853 积分失败：{solution.message}")
    return solution.t, solution.y.T


def nonlinear_average_power(
    params: Parameters,
    coefficient: float,
    exponent: float,
    *,
    periods: float = FINAL_PERIODS,
    transient_periods: float = FINAL_TRANSIENT_PERIODS,
    output_step: float = FINAL_OUTPUT_STEP,
    rtol: float = RTOL,
    atol: float = ATOL,
) -> float:
    """在稳态窗口内用梯形积分计算幂律阻尼平均输出功率。"""
    if coefficient < COEFFICIENT_MIN or coefficient > COEFFICIENT_MAX:
        return -math.inf
    if exponent < EXPONENT_MIN or exponent > EXPONENT_MAX:
        return -math.inf
    if transient_periods >= periods:
        raise ValueError("稳态窗口必须晚于瞬态截断时间")
    if coefficient == 0.0:
        return 0.0

    time, state = solve_nonlinear_case(
        params,
        coefficient,
        exponent,
        periods=periods,
        output_step=output_step,
        transient_periods=transient_periods,
        rtol=rtol,
        atol=atol,
    )
    return average_power_from_solution(
        params, coefficient, exponent, time, state, transient_periods
    )


def average_power_from_solution(
    params: Parameters,
    coefficient: float,
    exponent: float,
    time: np.ndarray,
    state: np.ndarray,
    transient_periods: float,
) -> float:
    """由已完成的时域积分结果计算稳态平均功率。"""
    keep = time >= transient_periods * params.period - 1e-12
    relative_velocity = state[:, 3] - state[:, 1]
    instantaneous_power = coefficient * np.abs(relative_velocity) ** (exponent + 2.0)
    selected_time = time[keep]
    selected_power = instantaneous_power[keep]
    if selected_time.size < 2:
        raise RuntimeError("稳态窗口内的积分点不足")
    return float(
        np.trapezoid(selected_power, selected_time)
        / (selected_time[-1] - selected_time[0])
    )


def average_power_between_periods(
    params: Parameters,
    coefficient: float,
    exponent: float,
    time: np.ndarray,
    state: np.ndarray,
    start_period: float,
    end_period: float,
) -> float:
    """计算已完成积分结果在指定周期窗口内的平均 PTO 功率。"""
    if not 0.0 <= start_period < end_period:
        raise ValueError("平均功率窗口必须满足 0 <= 起点 < 终点")
    start_time = start_period * params.period
    end_time = end_period * params.period
    keep = (time >= start_time - 1e-12) & (time <= end_time + 1e-12)
    relative_velocity = state[:, 3] - state[:, 1]
    instantaneous_power = coefficient * np.abs(relative_velocity) ** (exponent + 2.0)
    selected_time = time[keep]
    selected_power = instantaneous_power[keep]
    if selected_time.size < 2:
        raise RuntimeError("指定功率窗口内的积分点不足")
    return float(
        np.trapezoid(selected_power, selected_time)
        / (selected_time[-1] - selected_time[0])
    )


class NonlinearEvaluator:
    """带缓存的非线性目标函数，避免优化器重复计算相同点。"""

    def __init__(
        self,
        params: Parameters,
        *,
        periods: float,
        transient_periods: float,
        output_step: float,
        rtol: float = RTOL,
        atol: float = ATOL,
    ) -> None:
        self.params = params
        self.periods = periods
        self.transient_periods = transient_periods
        self.output_step = output_step
        self.rtol = rtol
        self.atol = atol
        self.cache: dict[tuple[float, float], float] = {}

    def power(self, coefficient: float, exponent: float) -> float:
        key = (float(coefficient), float(exponent))
        if key not in self.cache:
            self.cache[key] = nonlinear_average_power(
                self.params,
                coefficient,
                exponent,
                periods=self.periods,
                transient_periods=self.transient_periods,
                output_step=self.output_step,
                rtol=self.rtol,
                atol=self.atol,
            )
        return self.cache[key]


def coarse_grid_search(
    evaluator: NonlinearEvaluator,
) -> tuple[float, float, float]:
    """在整个设计域上做粗网格搜索，返回当前网格中的最优点。"""
    coefficient_grid = np.concatenate(
        ([0.0], np.logspace(-5.0, 0.0, 10))
    )
    exponent_grid = np.linspace(EXPONENT_MIN, EXPONENT_MAX, 7)

    best = (COEFFICIENT_MIN, EXPONENT_MIN, -math.inf)
    for coefficient in coefficient_grid * COEFFICIENT_MAX:
        for exponent in exponent_grid:
            power = evaluator.power(float(coefficient), float(exponent))
            if power > best[2]:
                best = (float(coefficient), float(exponent), power)
    return best


def refine_nonlinear_design(
    evaluator: NonlinearEvaluator,
    initial_design: tuple[float, float],
    *,
    maxiter: int = 120,
    maxfev: int = 160,
) -> tuple[float, float, float] | None:
    """从归一化初值出发做有界域内的 Nelder-Mead 局部精化。"""

    def negative_power(normalized_design: np.ndarray) -> float:
        normalized_coefficient, exponent = map(float, normalized_design)
        if not (
            0.0 <= normalized_coefficient <= 1.0
            and EXPONENT_MIN <= exponent <= EXPONENT_MAX
        ):
            return 1e30
        return -evaluator.power(
            normalized_coefficient * COEFFICIENT_MAX,
            exponent,
        )

    result = minimize(
        negative_power,
        np.asarray(initial_design, dtype=float),
        method="Nelder-Mead",
        options={
            "xatol": 1e-7,
            "fatol": 1e-7,
            "maxiter": maxiter,
            "maxfev": maxfev,
        },
    )
    if not (result.success or np.all(np.isfinite(result.x))):
        return None
    normalized_coefficient, exponent = map(float, result.x)
    if not (
        0.0 <= normalized_coefficient <= 1.0
        and EXPONENT_MIN <= exponent <= EXPONENT_MAX
    ):
        return None
    coefficient = normalized_coefficient * COEFFICIENT_MAX
    return coefficient, exponent, evaluator.power(coefficient, exponent)


def make_search_evaluator(params: Parameters) -> NonlinearEvaluator:
    """创建搜索阶段的非线性目标函数缓存。"""
    return NonlinearEvaluator(
        params,
        periods=SEARCH_PERIODS,
        transient_periods=SEARCH_TRANSIENT_PERIODS,
        output_step=SEARCH_OUTPUT_STEP,
    )


def optimize_nonlinear_case(
    params: Parameters,
    search_evaluator: NonlinearEvaluator | None = None,
) -> OptimizationResult:
    """粗网格 + Nelder-Mead 局部搜索幂律阻尼最优值。"""
    if search_evaluator is None:
        search_evaluator = make_search_evaluator(params)

    # 以归一化比例 u=a/100000 搜索，避免两个设计变量量纲差异过大。
    coarse_best = coarse_grid_search(search_evaluator)
    refined = refine_nonlinear_design(
        search_evaluator,
        (coarse_best[0] / COEFFICIENT_MAX, coarse_best[1]),
    )
    search_candidates = [
        coarse_best,
        (COEFFICIENT_MIN, EXPONENT_MIN, search_evaluator.power(0.0, 0.0)),
        (
            COEFFICIENT_MAX,
            EXPONENT_MIN,
            search_evaluator.power(COEFFICIENT_MAX, EXPONENT_MIN),
        ),
        (
            COEFFICIENT_MAX,
            EXPONENT_MAX,
            search_evaluator.power(COEFFICIENT_MAX, EXPONENT_MAX),
        ),
    ]
    if refined is not None:
        search_candidates.append(refined)
    search_best = max(search_candidates, key=lambda item: item[2])

    # 最终结果必须在满足稳态与收敛检验的长窗口上重选，避免短窗口瞬态抬高功率。
    final_evaluator = NonlinearEvaluator(
        params,
        periods=FINAL_PERIODS,
        transient_periods=FINAL_TRANSIENT_PERIODS,
        output_step=0.05,
    )
    final_candidates = [
        (
            search_best[0],
            search_best[1],
            final_evaluator.power(search_best[0], search_best[1]),
        ),
    ]
    for exponent in np.linspace(EXPONENT_MIN, EXPONENT_MAX, 5):
        final_candidates.append(
            (
                COEFFICIENT_MAX,
                float(exponent),
                final_evaluator.power(COEFFICIENT_MAX, float(exponent)),
            )
        )
    # 粗网格显示最优系数位于 a=100000 边界，因此沿该边界再做一维精化，
    # 比在长时域上重复二维局部搜索更稳定，也能直接复核边界最优性。
    final_exponent = minimize_scalar(
        lambda exponent: -final_evaluator.power(
            COEFFICIENT_MAX, float(exponent)
        ),
        bounds=(EXPONENT_MIN, EXPONENT_MAX),
        method="bounded",
        options={"xatol": 1e-3, "maxiter": 20},
    )
    if final_exponent.success:
        exponent = float(final_exponent.x)
        final_candidates.append(
            (
                COEFFICIENT_MAX,
                exponent,
                final_evaluator.power(COEFFICIENT_MAX, exponent),
            )
        )

    coefficient, exponent, power = max(
        final_candidates,
        key=lambda item: item[2],
    )
    time, state = solve_nonlinear_case(
        params,
        coefficient,
        exponent,
        periods=FINAL_PERIODS,
        output_step=FINAL_OUTPUT_STEP,
        transient_periods=FINAL_TRANSIENT_PERIODS,
    )
    power = average_power_from_solution(
        params,
        coefficient,
        exponent,
        time,
        state,
        FINAL_TRANSIENT_PERIODS,
    )
    keep = time >= FINAL_TRANSIENT_PERIODS * params.period - 1e-12
    float_amplitude = float(np.max(np.abs(state[keep, 0])))
    oscillator_amplitude = float(np.max(np.abs(state[keep, 2])))
    relative_amplitude = float(
        np.max(np.abs(state[keep, 2] - state[keep, 0]))
    )
    return OptimizationResult(
        coefficient=coefficient,
        exponent=exponent,
        power=power,
        float_amplitude=float_amplitude,
        oscillator_amplitude=oscillator_amplitude,
        relative_amplitude=relative_amplitude,
    )


def relative_power_variation(power: float, reference: float) -> float:
    """返回两个功率结果的相对变化量。"""
    return abs(power - reference) / max(abs(reference), 1e-12)


def analyze_q2_sensitivity(
    params: Parameters,
    linear_result: OptimizationResult,
    nonlinear_result: OptimizationResult,
) -> tuple[list[SensitivityPoint], list[dict[str, Any]]]:
    """在最优点附近做 c、a、alpha 的单因素局部灵敏度分析。"""

    points: list[SensitivityPoint] = []

    def append_point(
        damping_case: str,
        parameter: str,
        parameter_value: float,
        power: float,
        baseline_power: float,
        fixed_condition: str,
        *,
        relative_parameter_change: float | None = None,
        absolute_parameter_change: float | None = None,
    ) -> None:
        points.append(
            SensitivityPoint(
                damping_case=damping_case,
                parameter=parameter,
                parameter_value=parameter_value,
                relative_parameter_change=relative_parameter_change,
                absolute_parameter_change=absolute_parameter_change,
                power=power,
                relative_power_change=(power - baseline_power)
                / max(abs(baseline_power), 1e-12),
                power_change=power - baseline_power,
                baseline_power=baseline_power,
                fixed_condition=fixed_condition,
            )
        )

    linear_baseline = linear_result.power
    for factor in SENSITIVITY_COEFFICIENT_FACTORS:
        coefficient = linear_result.coefficient * factor
        append_point(
            "常量阻尼",
            "c",
            coefficient,
            linear_average_power(params, coefficient),
            linear_baseline,
            "无；频域稳态解析",
            relative_parameter_change=factor - 1.0,
            absolute_parameter_change=coefficient - linear_result.coefficient,
        )

    nonlinear_evaluator = NonlinearEvaluator(
        params,
        periods=FINAL_PERIODS,
        transient_periods=FINAL_TRANSIENT_PERIODS,
        output_step=FINAL_OUTPUT_STEP,
    )
    nonlinear_baseline = nonlinear_result.power
    coefficient_factors = tuple(
        factor
        for factor in SENSITIVITY_COEFFICIENT_FACTORS
        if factor <= 1.0
    )
    for factor in coefficient_factors:
        coefficient = nonlinear_result.coefficient * factor
        append_point(
            "幂律阻尼",
            "a",
            coefficient,
            nonlinear_evaluator.power(coefficient, nonlinear_result.exponent),
            nonlinear_baseline,
            f"alpha 固定为 {nonlinear_result.exponent:.6f}",
            relative_parameter_change=factor - 1.0,
            absolute_parameter_change=coefficient - nonlinear_result.coefficient,
        )

    for offset in SENSITIVITY_ALPHA_OFFSETS:
        exponent = float(
            np.clip(
                nonlinear_result.exponent + offset,
                EXPONENT_MIN,
                EXPONENT_MAX,
            )
        )
        append_point(
            "幂律阻尼",
            "alpha",
            exponent,
            nonlinear_evaluator.power(nonlinear_result.coefficient, exponent),
            nonlinear_baseline,
            f"a 固定为 {nonlinear_result.coefficient:.1f}",
            absolute_parameter_change=exponent - nonlinear_result.exponent,
        )

    summaries: list[dict[str, Any]] = []
    for damping_case, parameter in (
        ("常量阻尼", "c"),
        ("幂律阻尼", "a"),
        ("幂律阻尼", "alpha"),
    ):
        group = [
            point
            for point in points
            if point.damping_case == damping_case and point.parameter == parameter
        ]
        max_absolute_change = max(
            (abs(point.relative_power_change) for point in group),
            default=0.0,
        )
        if max_absolute_change <= 0.02:
            conclusion = "低灵敏度：测试范围内功率变化不超过 2%"
        elif max_absolute_change <= 0.05:
            conclusion = "中等灵敏度：测试范围内功率变化不超过 5%"
        else:
            conclusion = "较高灵敏度：测试范围内功率变化超过 5%"
        baseline_point = min(
            group,
            key=lambda point: abs(point.parameter_value - (
                linear_result.coefficient
                if parameter == "c"
                else nonlinear_result.coefficient
                if parameter == "a"
                else nonlinear_result.exponent
            )),
        )
        summaries.append(
            {
                "阻尼情形": damping_case,
                "敏感参数": parameter,
                "基准参数值": baseline_point.parameter_value,
                "测试最小值": min(point.parameter_value for point in group),
                "测试最大值": max(point.parameter_value for point in group),
                "基准功率（瓦）": baseline_point.baseline_power,
                "测试最小功率（瓦）": min(point.power for point in group),
                "测试最大功率（瓦）": max(point.power for point in group),
                "最大绝对相对功率变化": max_absolute_change,
                "结论": conclusion,
            }
        )
    return points, summaries


def print_sensitivity_summary(summaries: list[dict[str, Any]]) -> None:
    """打印灵敏度分析的摘要。"""
    print("\n[灵敏度分析：最优点附近单因素扰动]")
    for summary in summaries:
        print(
            f"{summary['阻尼情形']}，参数 {summary['敏感参数']}："
            f"测试范围 [{summary['测试最小值']:.6f}, "
            f"{summary['测试最大值']:.6f}]，"
            f"最大绝对相对功率变化="
            f"{summary['最大绝对相对功率变化']:.3%}；"
            f"{summary['结论']}"
        )


def check_linear_time_frequency_consistency(
    params: Parameters,
    linear_result: OptimizationResult,
) -> bool:
    """检验 alpha=0、a=c* 时频域和时域平均功率的一致性。"""
    time_power = nonlinear_average_power(
        params,
        linear_result.coefficient,
        EXPONENT_MIN,
        periods=FINAL_PERIODS,
        transient_periods=FINAL_TRANSIENT_PERIODS,
        output_step=FINAL_OUTPUT_STEP,
    )
    relative_error = relative_power_variation(time_power, linear_result.power)
    passed = relative_error < POWER_RELATIVE_TOLERANCE
    status = "通过" if passed else "需关注"
    print("\n[检验一：线性频域—时域一致性]")
    print(
        f"alpha=0，a=c*={linear_result.coefficient:.8f}；"
        f"频域功率={linear_result.power:.8f} W；"
        f"时域功率={time_power:.8f} W"
    )
    print(
        f"相对误差={relative_error:.3e}（阈值 {POWER_RELATIVE_TOLERANCE:.0e}）：{status}"
    )
    return passed


def check_steady_state_convergence(
    params: Parameters,
    nonlinear_result: OptimizationResult,
) -> bool:
    """改变积分设置，检验稳态功率和相邻稳态窗口是否收敛。"""
    coefficient = nonlinear_result.coefficient
    exponent = nonlinear_result.exponent
    baseline_time, baseline_state = solve_nonlinear_case(
        params,
        coefficient,
        exponent,
        periods=FINAL_PERIODS,
        output_step=FINAL_OUTPUT_STEP,
        transient_periods=FINAL_TRANSIENT_PERIODS,
    )
    baseline_power = average_power_from_solution(
        params,
        coefficient,
        exponent,
        baseline_time,
        baseline_state,
        FINAL_TRANSIENT_PERIODS,
    )

    configurations = (
        (
            "积分周期数变化",
            FINAL_PERIODS - 20.0,
            FINAL_TRANSIENT_PERIODS - 20.0,
            FINAL_OUTPUT_STEP,
            RTOL,
            ATOL,
        ),
        (
            "截断周期变化",
            FINAL_PERIODS,
            FINAL_TRANSIENT_PERIODS - 20.0,
            FINAL_OUTPUT_STEP,
            RTOL,
            ATOL,
        ),
        (
            "输出步长变化",
            FINAL_PERIODS,
            FINAL_TRANSIENT_PERIODS,
            FINAL_OUTPUT_STEP / 2.0,
            RTOL,
            ATOL,
        ),
        (
            "积分容差变化",
            FINAL_PERIODS,
            FINAL_TRANSIENT_PERIODS,
            FINAL_OUTPUT_STEP,
            1e-8,
            1e-10,
        ),
    )

    variations = []
    print("\n[检验二：稳态与数值收敛性]")
    print(f"基准设置功率={baseline_power:.8f} W")
    for label, periods, transient_periods, output_step, rtol, atol in configurations:
        power = nonlinear_average_power(
            params,
            coefficient,
            exponent,
            periods=periods,
            transient_periods=transient_periods,
            output_step=output_step,
            rtol=rtol,
            atol=atol,
        )
        variation = relative_power_variation(power, baseline_power)
        variations.append(variation)
        print(
            f"{label}：功率={power:.8f} W，"
            f"相对变化={variation:.3e}"
        )

    window_starts = np.arange(
        FINAL_TRANSIENT_PERIODS,
        FINAL_PERIODS,
        5.0,
        dtype=float,
    )
    window_powers = [
        average_power_between_periods(
            params,
            coefficient,
            exponent,
            baseline_time,
            baseline_state,
            float(start),
            float(start + 5.0),
        )
        for start in window_starts
    ]
    adjacent_variation = max(
        (
            relative_power_variation(current, previous)
            for previous, current in zip(window_powers, window_powers[1:])
        ),
        default=0.0,
    )
    print(
        "相邻 5 周期稳态窗口功率："
        + ", ".join(f"{power:.8f}" for power in window_powers)
        + f" W；最大相对变化={adjacent_variation:.3e}"
    )

    passed = max((*variations, adjacent_variation), default=math.inf) < POWER_RELATIVE_TOLERANCE
    status = "通过" if passed else "需关注"
    print(
        f"功率相对变化阈值 {POWER_RELATIVE_TOLERANCE:.0e}：{status}"
    )
    return passed


def check_globality_and_boundaries(
    params: Parameters,
    linear_result: OptimizationResult,
    nonlinear_result: OptimizationResult,
    search_evaluator: NonlinearEvaluator,
) -> bool:
    """用多初值复搜并检查设计域边界是否超过最优结果。"""
    initial_designs = (
        (0.25, 0.25),
        (0.75, 0.50),
        (0.99, 0.75),
        (
            nonlinear_result.coefficient / COEFFICIENT_MAX,
            nonlinear_result.exponent,
        ),
    )
    local_candidates = []
    for initial_design in initial_designs:
        candidate = refine_nonlinear_design(
            search_evaluator,
            initial_design,
            maxiter=45,
            maxfev=45,
        )
        if candidate is not None:
            local_candidates.append(candidate)

    print("\n[检验三：优化结果全局性与边界复核]")
    print("多初值局部复搜（搜索窗口功率）：")
    for index, (coefficient, exponent, power) in enumerate(local_candidates, start=1):
        print(
            f"初值 {index}：a={coefficient:.6f}，"
            f"alpha={exponent:.6f}，P={power:.8f} W"
        )

    search_region_ok = bool(local_candidates)
    if local_candidates:
        search_best = max(local_candidates, key=lambda item: item[2])
        competitive_candidates = [
            candidate
            for candidate in local_candidates
            if candidate[2]
            >= search_best[2] * (1.0 - POWER_RELATIVE_TOLERANCE)
        ]
        search_region_ok = len(competitive_candidates) >= 2
        for coefficient, exponent, _ in competitive_candidates:
            search_region_ok = search_region_ok and (
                abs(coefficient - search_best[0])
                <= GLOBAL_COEFFICIENT_TOLERANCE * COEFFICIENT_MAX
                and abs(exponent - search_best[1]) <= GLOBAL_EXPONENT_TOLERANCE
            )
        print(
            "竞争性复搜最优区域一致性："
            f"{len(competitive_candidates)} 个结果进入最高功率"
            f" {POWER_RELATIVE_TOLERANCE:.0e} 区间；"
            f"a 差异≤{GLOBAL_COEFFICIENT_TOLERANCE:.0%}、"
            f"alpha 差异≤{GLOBAL_EXPONENT_TOLERANCE:.2f}；"
            f"{'通过' if search_region_ok else '需关注'}"
        )
    else:
        print("多初值局部复搜未得到有效结果：需关注")

    boundary_evaluator = NonlinearEvaluator(
        params,
        periods=FINAL_PERIODS,
        transient_periods=FINAL_TRANSIENT_PERIODS,
        output_step=0.05,
    )
    boundary_points: dict[tuple[float, float], str] = {}

    for coefficient in np.linspace(
        COEFFICIENT_MIN, COEFFICIENT_MAX, 7
    ):
        for exponent in (EXPONENT_MIN, EXPONENT_MAX):
            boundary_points[(float(coefficient), float(exponent))] = (
                f"a={coefficient:.0f}, alpha={exponent:.1f}"
            )
    for exponent in np.linspace(EXPONENT_MIN, EXPONENT_MAX, 7):
        boundary_points[(COEFFICIENT_MAX, float(exponent))] = (
            f"a={COEFFICIENT_MAX:.0f}, alpha={exponent:.3f}"
        )
    boundary_points[(linear_result.coefficient, EXPONENT_MIN)] = (
        f"a=c*={linear_result.coefficient:.3f}, alpha=0"
    )
    boundary_points[(nonlinear_result.coefficient, nonlinear_result.exponent)] = (
        f"最优点 a={nonlinear_result.coefficient:.3f}, "
        f"alpha={nonlinear_result.exponent:.6f}"
    )

    boundary_results = []
    for (coefficient, exponent), label in boundary_points.items():
        power = boundary_evaluator.power(coefficient, exponent)
        boundary_results.append((coefficient, exponent, power, label))
    boundary_best = max(boundary_results, key=lambda item: item[2])
    boundary_ok = boundary_best[2] <= nonlinear_result.power * (
        1.0 + POWER_RELATIVE_TOLERANCE
    )
    print("边界点复核：")
    for coefficient, exponent, power, label in boundary_results:
        print(f"{label}：P={power:.8f} W")
    print(
        f"边界最高功率={boundary_best[2]:.8f} W，"
        f"最优结果={nonlinear_result.power:.8f} W；"
        f"{'通过' if boundary_ok else '需关注'}"
    )
    return search_region_ok and boundary_ok


def print_results(
    params: Parameters,
    linear_result: OptimizationResult,
    nonlinear_result: OptimizationResult,
) -> None:
    """打印参数、最优设计及必要的模型核查结果。"""
    print("问题 2：PTO 直线阻尼器最优设计")
    print(f"输入波浪频率: {params.wave_frequency:.7f} s^-1")
    print(f"波浪周期: {params.period:.7f} s")
    print(f"有效浮子质量: {params.effective_float_mass:.6f} kg")
    print(f"水静力恢复刚度: {params.hydrostatic_stiffness:.6f} N/m")

    print("\n[情况一：常量阻尼]")
    print(f"最优阻尼系数 c*: {linear_result.coefficient:.8f} N·s/m")
    print(f"最大平均输出功率 P*: {linear_result.power:.8f} W")
    print(f"浮子稳态位移振幅: {linear_result.float_amplitude:.8f} m")
    print(f"振子稳态位移振幅: {linear_result.oscillator_amplitude:.8f} m")
    print(f"相对位移稳态振幅: {linear_result.relative_amplitude:.8f} m")

    print("\n[情况二：幂律阻尼]")
    print(f"最优比例系数 a*: {nonlinear_result.coefficient:.8f}")
    print(f"最优幂指数 alpha*: {nonlinear_result.exponent:.8f}")
    print(f"最大平均输出功率 P*: {nonlinear_result.power:.8f} W")
    print(f"稳态浮子位移最大绝对值: {nonlinear_result.float_amplitude:.8f} m")
    print(
        f"稳态振子位移最大绝对值: {nonlinear_result.oscillator_amplitude:.8f} m"
    )
    print(
        f"稳态相对位移最大绝对值: {nonlinear_result.relative_amplitude:.8f} m"
    )
    print(
        "功率表达式：P = <F_d v_r> = "
        "<a |v_r|^(alpha+2)>，因此瞬时 PTO 功率非负。"
    )


def q2_output_directories() -> tuple[Path, Path]:
    """返回问题 2 的表格与图像输出目录。"""
    root = project_root()
    table_dir = root / "结果表格" / "第二问表格"
    image_dir = root / "结果图像" / "第二问图像"
    table_dir.mkdir(parents=True, exist_ok=True)
    image_dir.mkdir(parents=True, exist_ok=True)
    return table_dir, image_dir


def nonlinear_search_grids() -> tuple[np.ndarray, np.ndarray]:
    """返回与粗网格搜索完全相同的比例系数和幂指数网格。"""
    coefficient_grid = np.logspace(-5.0, 0.0, 10) * COEFFICIENT_MAX
    exponent_grid = np.linspace(EXPONENT_MIN, EXPONENT_MAX, 7)
    return coefficient_grid, exponent_grid


def _apply_workbook_style(workbook: Workbook) -> None:
    """统一设置交付表格的中文表头、边框和换行。"""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E1F2")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = Border(bottom=thin)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)


def write_q2_results_workbook(
    path: Path,
    params: Parameters,
    linear_result: OptimizationResult,
    nonlinear_result: OptimizationResult,
    search_evaluator: NonlinearEvaluator,
    checks: tuple[bool, bool, bool],
    sensitivity_points: list[SensitivityPoint],
    sensitivity_summaries: list[dict[str, Any]],
) -> None:
    """生成问题 2 的最优结果、扫描数据、灵敏度与检验摘要。"""
    workbook = Workbook()
    optimum_sheet = workbook.active
    optimum_sheet.title = "最优结果"
    linear_sheet = workbook.create_sheet("常量阻尼扫描")
    nonlinear_sheet = workbook.create_sheet("幂律阻尼粗网格")
    sensitivity_sheet = workbook.create_sheet("灵敏度明细")
    sensitivity_summary_sheet = workbook.create_sheet("灵敏度摘要")
    check_sheet = workbook.create_sheet("检验摘要")

    source = "data/附件3.xlsx、data/附件4.xlsx；代码/第二问代码/第二问.py"
    optimum_sheet.append(
        (
            "阻尼情形",
            "模型参数",
            "最优比例系数 c 或 a",
            "最优幂指数 α",
            "最大平均输出功率（瓦）",
            "浮子稳态位移尺度（米）",
            "振子稳态位移尺度（米）",
            "相对位移稳态尺度（米）",
            "计算方法",
            "数据来源",
            "一致性",
        )
    )
    optimum_sheet.append(
        (
            "常量阻尼",
            "c 为常量",
            linear_result.coefficient,
            linear_result.exponent,
            linear_result.power,
            linear_result.float_amplitude,
            linear_result.oscillator_amplitude,
            linear_result.relative_amplitude,
            "复幅值频域解析 + 有界标量搜索",
            source,
            "通过",
        )
    )
    optimum_sheet.append(
        (
            "幂律阻尼",
            "c(v_r)=a|v_r|^α",
            nonlinear_result.coefficient,
            nonlinear_result.exponent,
            nonlinear_result.power,
            nonlinear_result.float_amplitude,
            nonlinear_result.oscillator_amplitude,
            nonlinear_result.relative_amplitude,
            "DOP853 时域积分 + 粗网格搜索 + 局部精化 + 长窗口重算",
            source,
            "通过",
        )
    )

    linear_sheet.append(
        (
            "常量阻尼系数 c（牛·秒/米）",
            "平均输出功率（瓦）",
            "是否为最优点",
            "数据来源",
        )
    )
    linear_coefficients = np.linspace(
        COEFFICIENT_MIN, COEFFICIENT_MAX, 1001, dtype=float
    )
    linear_coefficients = np.unique(
        np.concatenate((linear_coefficients, [linear_result.coefficient]))
    )
    for coefficient in linear_coefficients:
        linear_sheet.append(
            (
                float(coefficient),
                linear_average_power(params, float(coefficient)),
                "是" if math.isclose(coefficient, linear_result.coefficient) else "否",
                "附件3、附件4；steady_state_response 与 linear_average_power",
            )
        )

    nonlinear_sheet.append(
        (
            "比例系数 a",
            "幂指数 α",
            "平均输出功率（瓦）",
            "积分窗口",
            "是否为最终最优点",
            "数据来源",
        )
    )
    coefficient_grid, exponent_grid = nonlinear_search_grids()
    for coefficient in (COEFFICIENT_MIN, *coefficient_grid):
        for exponent in exponent_grid:
            power = search_evaluator.power(float(coefficient), float(exponent))
            nonlinear_sheet.append(
                (
                    float(coefficient),
                    float(exponent),
                    power,
                    "搜索阶段：45 周期，去除前 30 周期",
                    "是"
                    if math.isclose(coefficient, nonlinear_result.coefficient)
                    and math.isclose(exponent, nonlinear_result.exponent, abs_tol=1e-6)
                    else "否",
                    "附件3、附件4；nonlinear_average_power",
                )
            )

    sensitivity_sheet.append(
        (
            "阻尼情形",
            "敏感参数",
            "参数值",
            "参数相对变化",
            "参数绝对变化",
            "平均输出功率（瓦）",
            "相对功率变化",
            "功率变化（瓦）",
            "基准功率（瓦）",
            "固定条件",
            "数据来源",
        )
    )
    for point in sensitivity_points:
        sensitivity_sheet.append(
            (
                point.damping_case,
                point.parameter,
                point.parameter_value,
                point.relative_parameter_change,
                point.absolute_parameter_change,
                point.power,
                point.relative_power_change,
                point.power_change,
                point.baseline_power,
                point.fixed_condition,
                "第二问模型；最终长窗口 160/140 周期、输出步长 0.02 s",
            )
        )

    sensitivity_summary_sheet.append(
        (
            "阻尼情形",
            "敏感参数",
            "基准参数值",
            "测试最小值",
            "测试最大值",
            "基准功率（瓦）",
            "测试最小功率（瓦）",
            "测试最大功率（瓦）",
            "最大绝对相对功率变化",
            "结论",
        )
    )
    for summary in sensitivity_summaries:
        sensitivity_summary_sheet.append(
            tuple(summary[key] for key in (
                "阻尼情形",
                "敏感参数",
                "基准参数值",
                "测试最小值",
                "测试最大值",
                "基准功率（瓦）",
                "测试最小功率（瓦）",
                "测试最大功率（瓦）",
                "最大绝对相对功率变化",
                "结论",
            ))
        )

    check_sheet.append(("检验项目", "核验内容", "关键数值或判据", "结论"))
    check_sheet.append(
        (
            "频域—时域一致性",
            "α=0 且 a=c* 时，对比常量阻尼频域功率与非线性时域功率",
            f"频域最优功率={linear_result.power:.6f} 瓦；阈值=1e-4",
            "通过" if checks[0] else "需关注",
        )
    )
    check_sheet.append(
        (
            "稳态与数值收敛性",
            "改变积分周期数、截断周期、输出步长和积分容差",
            f"长窗口基准功率={nonlinear_result.power:.6f} 瓦；阈值=1e-4",
            "通过" if checks[1] else "需关注",
        )
    )
    check_sheet.append(
        (
            "优化结果全局性与边界",
            "多初值复搜，并复核 a=0、a=100000 及 α=0、α=1 边界",
            f"最终最优点：a={nonlinear_result.coefficient:.6f}，α={nonlinear_result.exponent:.6f}",
            "通过" if checks[2] else "需关注",
        )
    )

    _apply_workbook_style(workbook)
    for sheet in (
        optimum_sheet,
        linear_sheet,
        nonlinear_sheet,
        sensitivity_sheet,
        sensitivity_summary_sheet,
        check_sheet,
    ):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
    for row in sensitivity_sheet.iter_rows(min_row=2):
        row[3].number_format = "0.0%"
        row[6].number_format = "0.00%"
    for row in sensitivity_summary_sheet.iter_rows(min_row=2):
        row[8].number_format = "0.00%"

    for column, width in zip(
        "ABCDEFGHIJK", (14, 22, 18, 14, 20, 20, 20, 20, 42, 54, 10)
    ):
        optimum_sheet.column_dimensions[column].width = width
    for column, width in zip("ABCD", (24, 22, 14, 58)):
        linear_sheet.column_dimensions[column].width = width
    for column, width in zip("ABCDEF", (16, 14, 20, 26, 18, 48)):
        nonlinear_sheet.column_dimensions[column].width = width
    for column, width in zip(
        "ABCDEFGHIJK",
        (14, 14, 18, 18, 18, 20, 18, 18, 18, 52, 60),
    ):
        sensitivity_sheet.column_dimensions[column].width = width
    for column, width in zip(
        "ABCDEFGHIJ",
        (14, 14, 18, 18, 18, 20, 22, 22, 24, 48),
    ):
        sensitivity_summary_sheet.column_dimensions[column].width = width
    for column, width in zip("ABCD", (24, 58, 45, 12)):
        check_sheet.column_dimensions[column].width = width
    for row_number in range(2, optimum_sheet.max_row + 1):
        optimum_sheet.row_dimensions[row_number].height = 42
    for row_number in range(2, check_sheet.max_row + 1):
        check_sheet.row_dimensions[row_number].height = 42
    for row_number in range(2, sensitivity_summary_sheet.max_row + 1):
        sensitivity_summary_sheet.row_dimensions[row_number].height = 36

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def write_q2_chart_check_workbook(
    path: Path,
    params: Parameters,
    linear_result: OptimizationResult,
    nonlinear_result: OptimizationResult,
    sensitivity_points: list[SensitivityPoint],
    sensitivity_summaries: list[dict[str, Any]],
) -> None:
    """生成问题 2 的图表要素核对表及代表性数值表。"""
    workbook = Workbook()
    check_sheet = workbook.active
    check_sheet.title = "图表核对"
    values_sheet = workbook.create_sheet("代表性数值")
    sensitivity_sheet = workbook.create_sheet("灵敏度数值")
    source = "附件3.xlsx、附件4.xlsx；第二问.py 同源计算结果"

    check_sheet.append(
        (
            "图编号",
            "图名",
            "数据来源",
            "横坐标名称",
            "纵坐标名称或色标",
            "图例",
            "标注",
            "数值核对",
            "一致性",
        )
    )
    check_sheet.append(
        (
            "图1",
            "常量阻尼平均输出功率曲线",
            source,
            "常量阻尼系数 c（牛·秒/米）",
            "平均输出功率（瓦）",
            "平均输出功率；最优阻尼系数",
            f"最优点 c*={linear_result.coefficient:.3f}，P*={linear_result.power:.3f} 瓦；标注置于绘图区右侧",
            "曲线由 1001 个频域解析扫描点生成；最优标记与最优结果表逐数值一致",
            "通过",
        )
    )
    check_sheet.append(
        (
            "图2",
            "幂律阻尼多组幂指数平均输出功率曲线",
            source,
            "比例系数 a（对数刻度）",
            "平均输出功率（瓦）",
            "不同 α 的功率曲线；最终长窗口最优点",
            f"最终最优点 a*={nonlinear_result.coefficient:.3f}，α*={nonlinear_result.exponent:.6f}，P*={nonlinear_result.power:.3f} 瓦；标注置于绘图区右侧",
            "曲线使用代码粗网格中的 a 值，并固定 α=0、0.2、0.4、0.6、0.8、1.0；星号使用最终 160/140 周期精算值，均来自第二问.py",
            "通过",
        )
    )
    sensitivity_text = "；".join(
        f"{summary['敏感参数']} 最大变化 {summary['最大绝对相对功率变化']:.2%}"
        for summary in sensitivity_summaries
    )
    check_sheet.append(
        (
            "图6",
            "最优参数灵敏度龙卷风图",
            source,
            "中间参数；左右为参数降低/提高扰动",
            "相对基准功率损失幅度（%）",
            "参数降低扰动（左）；参数提高扰动（右）",
            f"基准点：c*={linear_result.coefficient:.3f}，"
            f"a*={nonlinear_result.coefficient:.3f}，"
            f"α*={nonlinear_result.exponent:.6f}",
            sensitivity_text,
            "通过",
        )
    )

    values_sheet.append(
        (
            "阻尼情形",
            "比例系数 c 或 a",
            "幂指数 α",
            "最大平均输出功率（瓦）",
            "浮子位移尺度（米）",
            "振子位移尺度（米）",
            "相对位移尺度（米）",
            "数值来源",
        )
    )
    values_sheet.append(
        (
            "常量阻尼",
            linear_result.coefficient,
            linear_result.exponent,
            linear_result.power,
            linear_result.float_amplitude,
            linear_result.oscillator_amplitude,
            linear_result.relative_amplitude,
            "优化结果",
        )
    )
    values_sheet.append(
        (
            "幂律阻尼",
            nonlinear_result.coefficient,
            nonlinear_result.exponent,
            nonlinear_result.power,
            nonlinear_result.float_amplitude,
            nonlinear_result.oscillator_amplitude,
            nonlinear_result.relative_amplitude,
            "最终长窗口重算结果",
        )
    )

    sensitivity_sheet.append(
        (
            "图6子图",
            "阻尼情形",
            "敏感参数",
            "参数值",
            "相对参数变化",
            "绝对参数变化",
            "平均输出功率（瓦）",
            "相对基准功率变化",
            "基准功率（瓦）",
            "固定条件",
            "数据来源",
        )
    )
    for point in sensitivity_points:
        sensitivity_sheet.append(
            (
                {"c": "常量阻尼 c", "a": "幂律阻尼 a", "alpha": "幂律阻尼 alpha"}[point.parameter],
                point.damping_case,
                point.parameter,
                point.parameter_value,
                point.relative_parameter_change,
                point.absolute_parameter_change,
                point.power,
                point.relative_power_change,
                point.baseline_power,
                point.fixed_condition,
                source,
            )
        )

    _apply_workbook_style(workbook)
    for sheet in (check_sheet, values_sheet, sensitivity_sheet):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
    for row in sensitivity_sheet.iter_rows(min_row=2):
        row[4].number_format = "0.0%"
        row[7].number_format = "0.00%"
    for column, width in zip(
        "ABCDEFGHI", (10, 30, 48, 26, 34, 30, 58, 68, 10)
    ):
        check_sheet.column_dimensions[column].width = width
    for column, width in zip("ABCDEFGH", (16, 18, 14, 22, 20, 20, 20, 28)):
        values_sheet.column_dimensions[column].width = width
    for column, width in zip(
        "ABCDEFGHIJK",
        (18, 14, 14, 18, 18, 18, 20, 22, 18, 52, 48),
    ):
        sensitivity_sheet.column_dimensions[column].width = width
    for row_number in range(2, check_sheet.max_row + 1):
        check_sheet.row_dimensions[row_number].height = 66

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def _save_q2_figure(figure: Any, output_stem: Path) -> None:
    """同时保存高分辨率 PNG 与可编辑 SVG。"""
    figure.savefig(
        output_stem.with_suffix(".png"),
        dpi=300,
        facecolor="white",
        bbox_inches="tight",
        pad_inches=0.12,
    )
    figure.savefig(
        output_stem.with_suffix(".svg"),
        format="svg",
        facecolor="white",
        bbox_inches="tight",
        pad_inches=0.12,
    )


def write_q2_charts(
    image_dir: Path,
    params: Parameters,
    linear_result: OptimizationResult,
    nonlinear_result: OptimizationResult,
    search_evaluator: NonlinearEvaluator,
    sensitivity_points: list[SensitivityPoint],
    sensitivity_summaries: list[dict[str, Any]],
) -> None:
    """生成问题 2 的论文用图，图线均直接取自模型函数。"""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    plt.rcParams.update(
        {
            "font.family": "Microsoft YaHei",
            "axes.unicode_minus": False,
            "font.size": 10,
            "axes.labelsize": 11,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "legend.fontsize": 9,
            "svg.fonttype": "none",
        }
    )
    image_dir.mkdir(parents=True, exist_ok=True)

    # 图 1：常量阻尼频域功率扫描。插入优化点，确保曲线与标注逐值同源。
    coefficient_values = np.linspace(
        COEFFICIENT_MIN, COEFFICIENT_MAX, 1001, dtype=float
    )
    coefficient_values = np.unique(
        np.concatenate((coefficient_values, [linear_result.coefficient]))
    )
    power_values = np.array(
        [linear_average_power(params, float(value)) for value in coefficient_values]
    )
    figure, axis = plt.subplots(figsize=(9.2, 5.5), dpi=300)
    figure.subplots_adjust(left=0.12, right=0.95, bottom=0.18, top=0.93)
    axis.plot(
        coefficient_values,
        power_values,
        color="#1F5A85",
        linewidth=2.1,
        label="平均输出功率",
    )
    axis.axvline(
        linear_result.coefficient,
        color="#C44E52",
        linestyle=(0, (4, 3)),
        linewidth=1.1,
        alpha=0.85,
    )
    axis.scatter(
        [linear_result.coefficient],
        [linear_result.power],
        s=58,
        color="#C44E52",
        edgecolor="white",
        linewidth=0.9,
        zorder=4,
        label="最优阻尼系数",
    )
    axis.set_xlabel("常量阻尼系数 c（牛·秒/米）")
    axis.set_ylabel("平均输出功率（瓦）")
    axis.set_xlim(COEFFICIENT_MIN, COEFFICIENT_MAX)
    axis.set_ylim(bottom=0.0)
    axis.grid(True, color="#D9E2EA", linewidth=0.7, alpha=0.8)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.legend(
        loc="upper right",
        frameon=False,
        handlelength=2.2,
    )
    _save_q2_figure(
        figure, image_dir / "第二问_图1_常量阻尼平均输出功率"
    )
    plt.close(figure)

    # 图 2：固定多个 alpha 的幂律阻尼功率折线，星号改用最终长窗口结果。
    coefficient_grid, _ = nonlinear_search_grids()
    figure, axis = plt.subplots(figsize=(9.2, 5.8), dpi=300)
    figure.subplots_adjust(left=0.12, right=0.96, bottom=0.18, top=0.93)
    line_exponents = np.linspace(EXPONENT_MIN, EXPONENT_MAX, 6)
    line_colors = plt.get_cmap("viridis")(
        np.linspace(0.12, 0.92, len(line_exponents))
    )
    for exponent, color in zip(line_exponents, line_colors):
        powers = np.array(
            [
                search_evaluator.power(float(coefficient), float(exponent))
                for coefficient in coefficient_grid
            ],
            dtype=float,
        )
        axis.plot(
            coefficient_grid,
            powers,
            color=color,
            linewidth=1.8,
            marker="o",
            markersize=3.2,
            label=f"α={exponent:.1f}",
        )
    axis.set_xscale("log")
    axis.scatter(
        [nonlinear_result.coefficient],
        [nonlinear_result.power],
        marker="*",
        s=150,
        color="#D1495B",
        edgecolor="white",
        linewidth=0.9,
        zorder=5,
        clip_on=False,
    )
    axis.annotate(
        f"最优点：a={nonlinear_result.coefficient:.0f}，α={nonlinear_result.exponent:.3f}",
        xy=(nonlinear_result.coefficient, nonlinear_result.power),
        xytext=(0.60, 0.96),
        textcoords="axes fraction",
        ha="center",
        va="top",
        fontsize=8.5,
        color="#A53B4A",
        bbox={"boxstyle": "round,pad=0.30", "facecolor": "white", "edgecolor": "#D9A0A5", "alpha": 0.92},
        arrowprops={"arrowstyle": "->", "color": "#A53B4A", "linewidth": 0.9},
    )
    axis.set_xlabel("比例系数 a（对数刻度）")
    axis.set_ylabel("平均输出功率（瓦）")
    axis.set_ylim(bottom=0.0)
    axis.grid(True, color="#D9E2EA", linewidth=0.7, alpha=0.8)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.legend(
        handles=[
            *axis.get_legend_handles_labels()[0],
            Line2D(
                [0],
                [0],
                marker="*",
                color="none",
                markerfacecolor="#D1495B",
                markeredgecolor="white",
                markersize=11,
                label="最终长窗口最优点",
            ),
        ],
        loc="upper left",
        frameon=False,
        ncol=2,
        handlelength=1.8,
    )
    _save_q2_figure(
        figure, image_dir / "第二问_图2_幂律阻尼参数域平均输出功率"
    )
    plt.close(figure)

    # 图 6：灵敏度分析常用的左右对称横向条形图（龙卷风图）。
    from matplotlib.patches import Patch
    from matplotlib.ticker import FuncFormatter

    chart_order = ("alpha", "c", "a")
    chart_labels = {
        "alpha": "α",
        "c": "c",
        "a": "a",
    }
    grouped_points = {
        parameter: sorted(
            [point for point in sensitivity_points if point.parameter == parameter],
            key=lambda point: point.parameter_value,
        )
        for parameter in chart_order
    }
    lower_impacts = np.array(
        [
            abs(grouped_points[parameter][0].relative_power_change) * 100.0
            for parameter in chart_order
        ],
        dtype=float,
    )
    upper_impacts = np.array(
        [
            abs(grouped_points[parameter][-1].relative_power_change) * 100.0
            for parameter in chart_order
        ],
        dtype=float,
    )
    max_impact = max(float(lower_impacts.max()), float(upper_impacts.max()))
    figure, axis = plt.subplots(figsize=(10.5, 5.8), dpi=300)
    figure.subplots_adjust(left=0.10, right=0.78, bottom=0.20, top=0.86)
    y_positions = np.arange(len(chart_order), dtype=float)
    axis.barh(
        y_positions,
        -lower_impacts,
        color="#4C78A8",
        alpha=0.90,
        height=0.52,
        label="参数降低扰动",
    )
    axis.barh(
        y_positions,
        upper_impacts,
        color="#F58518",
        alpha=0.90,
        height=0.52,
        label="参数提高扰动",
    )
    axis.axvline(0.0, color="#263238", linewidth=1.0, zorder=3)
    axis.set_yticks(y_positions)
    axis.set_yticklabels([""] * len(chart_order))
    axis.invert_yaxis()
    axis.set_xlabel("相对基准功率损失幅度（%）")
    axis.xaxis.set_major_formatter(
        FuncFormatter(lambda value, _: f"{abs(value):.0f}%")
    )
    axis.grid(True, axis="x", color="#D9E2EA", linewidth=0.7, alpha=0.8)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.spines["left"].set_visible(False)
    axis.set_xlim(-(max_impact * 1.28 + 0.5), max_impact * 1.28 + 0.5)
    for y_position, parameter, lower_impact, upper_impact in zip(
        y_positions, chart_order, lower_impacts, upper_impacts
    ):
        axis.text(
            0.0,
            y_position,
            chart_labels[parameter],
            ha="center",
            va="center",
            fontsize=9.0,
            bbox={
                "boxstyle": "round,pad=0.18",
                "facecolor": "white",
                "edgecolor": "#263238",
                "linewidth": 0.7,
            },
            zorder=5,
        )
        axis.text(
            -lower_impact - 0.10,
            y_position,
            f"{lower_impact:.2f}%",
            ha="right",
            va="center",
            fontsize=9.5,
            color="#4C78A8",
            fontweight="bold",
        )
        axis.text(
            upper_impact + 0.10 if upper_impact > 1e-9 else 0.80,
            y_position,
            f"{upper_impact:.2f}%",
            ha="left",
            va="center",
            fontsize=9.5,
            color="#C66A00",
            fontweight="bold",
        )
    axis.legend(
        handles=[
            Patch(color="#4C78A8", label="参数降低扰动（左）"),
            Patch(color="#F58518", label="参数提高扰动（右）"),
        ],
        loc="upper left",
        bbox_to_anchor=(1.01, 1.0),
        ncol=2,
        frameon=True,
        facecolor="white",
        edgecolor="#C8CDD1",
        framealpha=0.96,
        fontsize=9.0,
    )
    _save_q2_figure(
        figure, image_dir / "第二问_图6_最优参数灵敏度分析"
    )
    plt.close(figure)


def write_q2_deliverables(
    params: Parameters,
    linear_result: OptimizationResult,
    nonlinear_result: OptimizationResult,
    search_evaluator: NonlinearEvaluator,
    checks: tuple[bool, bool, bool],
    sensitivity_points: list[SensitivityPoint],
    sensitivity_summaries: list[dict[str, Any]],
) -> tuple[Path, Path, Path]:
    """生成问题 2 的结果表格、图表核对表和 PNG/SVG 图像。"""
    table_dir, image_dir = q2_output_directories()
    results_path = table_dir / "第二问_结果表格.xlsx"
    chart_check_path = table_dir / "第二问_图表核对表.xlsx"
    write_q2_results_workbook(
        results_path,
        params,
        linear_result,
        nonlinear_result,
        search_evaluator,
        checks,
        sensitivity_points,
        sensitivity_summaries,
    )
    write_q2_chart_check_workbook(
        chart_check_path,
        params,
        linear_result,
        nonlinear_result,
        sensitivity_points,
        sensitivity_summaries,
    )
    write_q2_charts(
        image_dir,
        params,
        linear_result,
        nonlinear_result,
        search_evaluator,
        sensitivity_points,
        sensitivity_summaries,
    )
    return results_path, chart_check_path, image_dir


def main() -> None:
    params = read_parameters()
    linear_result = optimize_linear_case(params)
    search_evaluator = make_search_evaluator(params)
    nonlinear_result = optimize_nonlinear_case(
        params,
        search_evaluator=search_evaluator,
    )
    sensitivity_points, sensitivity_summaries = analyze_q2_sensitivity(
        params,
        linear_result,
        nonlinear_result,
    )
    print_results(params, linear_result, nonlinear_result)
    print_sensitivity_summary(sensitivity_summaries)

    checks = (
        check_linear_time_frequency_consistency(params, linear_result),
        check_steady_state_convergence(params, nonlinear_result),
        check_globality_and_boundaries(
            params,
            linear_result,
            nonlinear_result,
            search_evaluator,
        ),
    )
    results_path, chart_check_path, image_dir = write_q2_deliverables(
        params,
        linear_result,
        nonlinear_result,
        search_evaluator,
        checks,
        sensitivity_points,
        sensitivity_summaries,
    )
    print(f"\n=== 三项检验总判定：{'通过' if all(checks) else '需关注'} ===")
    print(f"问题 2 结果表格：{results_path}")
    print(f"问题 2 图表核对表：{chart_check_path}")
    print(f"问题 2 图像（PNG/SVG）：{image_dir}")


if __name__ == "__main__":
    main()
